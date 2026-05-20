"""
Tests for the footer_scanner module.

Tests the refactored helper functions:
- find_column_id_by_index: Maps column index → column ID
- find_total_label_cell: Finds the TOTAL label in a worksheet
- get_cell_merge_colspan: Gets colspan of a merged cell
- find_pallet_count_column: Finds pallet count pattern on footer row
- scan_footer: Full orchestration test
"""

import unittest
import logging
from unittest.mock import MagicMock, PropertyMock
from dataclasses import dataclass
from typing import List, Optional


# --- Mock ColumnInfo to avoid importing the full scanner ---
@dataclass
class MockColumnInfo:
    """Lightweight mock matching ColumnInfo's interface."""
    id: str
    col_index: int
    colspan: int = 1


# Import the functions under test
from core.blueprint_generator.utils.footer_scanner import (
    find_column_id_by_index,
    scan_footer,
    FooterInfo,
)
from core.blueprint_generator.utils.content_extractor import (
    find_total_label_cell,
    get_cell_merge_colspan,
    find_pallet_count_column,
    find_footer_hs_code,
    extract_static_column_values,
)


class TestFindColumnIdByIndex(unittest.TestCase):
    """Tests for find_column_id_by_index()."""

    def setUp(self):
        """Set up common column definitions for tests."""
        self.columns = [
            MockColumnInfo(id="col_po", col_index=1, colspan=1),
            MockColumnInfo(id="col_desc", col_index=2, colspan=2),  # Spans cols 2-3
            MockColumnInfo(id="col_qty", col_index=4, colspan=1),
            MockColumnInfo(id="col_amount", col_index=5, colspan=1),
        ]

    def test_exact_match(self):
        """Column index matches col_index exactly."""
        self.assertEqual(find_column_id_by_index(1, self.columns), "col_po")
        self.assertEqual(find_column_id_by_index(4, self.columns), "col_qty")
        self.assertEqual(find_column_id_by_index(5, self.columns), "col_amount")

    def test_colspan_range_match(self):
        """Column index falls within a colspan range."""
        # col_desc spans cols 2-3 (col_index=2, colspan=2)
        self.assertEqual(find_column_id_by_index(2, self.columns), "col_desc")
        self.assertEqual(find_column_id_by_index(3, self.columns), "col_desc")

    def test_no_match_returns_none(self):
        """Column index not covered by any column returns None."""
        self.assertIsNone(find_column_id_by_index(10, self.columns))
        self.assertIsNone(find_column_id_by_index(99, self.columns))

    def test_empty_columns_list(self):
        """Empty columns list returns None."""
        self.assertIsNone(find_column_id_by_index(1, []))


class TestFindTotalLabelCell(unittest.TestCase):
    """Tests for find_total_label_cell()."""

    def setUp(self):
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL：", "TOTAL OF："]
            }
        }

    def _make_worksheet(self, cell_values: dict, max_column: int = 10):
        """
        Create a mock worksheet with specific cell values.
        
        Args:
            cell_values: Dict of {(row, col): value}
            max_column: Max column count for scanning
        """
        ws = MagicMock()
        ws.max_column = max_column

        def mock_cell(row, column):
            cell = MagicMock()
            cell.value = cell_values.get((row, column))
            cell.row = row
            cell.column = column
            return cell

        ws.cell = mock_cell
        return ws

    def test_finds_total(self):
        """Detects 'TOTAL' keyword."""
        ws = self._make_worksheet({(5, 2): "TOTAL"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertEqual(cell.row, 5)
        self.assertEqual(cell.column, 2)
        self.assertTrue(is_exact)

    def test_finds_total_of_colon(self):
        """Detects 'TOTAL OF:' keyword."""
        ws = self._make_worksheet({(8, 3): "TOTAL OF:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertEqual(cell.row, 8)
        self.assertTrue(is_exact)

    def test_finds_total_fullwidth_colon(self):
        """Detects 'TOTAL：' with fullwidth colon (common in CJK templates)."""
        ws = self._make_worksheet({(3, 1): "TOTAL："})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertTrue(is_exact)

    def test_case_insensitive(self):
        """Detection is case-insensitive."""
        ws = self._make_worksheet({(4, 1): "total of:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertTrue(is_exact)

    def test_no_match_returns_none(self):
        """Returns None if no TOTAL keyword found."""
        ws = self._make_worksheet({(1, 1): "Subtotal", (2, 1): "Grand Summary"})
        custom_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL OF:"]
            }
        }
        cell = find_total_label_cell(ws, 1, 10, custom_config)
        self.assertIsNone(cell)

    def test_ignores_partial_match(self):
        """Does NOT match partial strings like 'Total Net Weight'."""
        ws = self._make_worksheet({(1, 1): "Total Net Weight"})
        custom_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL OF:"]
            }
        }
        cell = find_total_label_cell(ws, 1, 10, custom_config)
        # "TOTAL NET WEIGHT" is not in the keywords set and doesn't startswith "TOTAL OF"
        self.assertIsNone(cell)

    def test_returns_first_match(self):
        """Returns the first TOTAL found when scanning top-down."""
        ws = self._make_worksheet({(3, 1): "TOTAL:", (7, 1): "TOTAL OF:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertEqual(cell.row, 3)


class TestGetCellMergeColspan(unittest.TestCase):
    """Tests for get_cell_merge_colspan()."""

    def _make_worksheet_with_merges(self, merge_ranges):
        """
        Create a mock worksheet with merged cell ranges.
        
        Args:
            merge_ranges: List of (min_row, min_col, max_row, max_col) tuples
        """
        ws = MagicMock()
        ranges = []
        for min_r, min_c, max_r, max_c in merge_ranges:
            r = MagicMock()
            r.min_row, r.min_col = min_r, min_c
            r.max_row, r.max_col = max_r, max_c
            ranges.append(r)
        ws.merged_cells.ranges = ranges
        return ws

    def test_merged_cell_returns_colspan(self):
        """Cell in a merged range returns correct colspan."""
        ws = self._make_worksheet_with_merges([(5, 2, 5, 4)])  # Cols 2-4 merged
        cell = MagicMock()
        cell.row, cell.column = 5, 2
        self.assertEqual(get_cell_merge_colspan(ws, cell), 3)

    def test_unmerged_cell_returns_1(self):
        """Cell not in any merged range returns 1."""
        ws = self._make_worksheet_with_merges([(5, 2, 5, 4)])
        cell = MagicMock()
        cell.row, cell.column = 3, 1
        self.assertEqual(get_cell_merge_colspan(ws, cell), 1)

    def test_no_merges_returns_1(self):
        """Worksheet with no merges returns 1."""
        ws = self._make_worksheet_with_merges([])
        cell = MagicMock()
        cell.row, cell.column = 1, 1
        self.assertEqual(get_cell_merge_colspan(ws, cell), 1)


class TestFindPalletCountColumn(unittest.TestCase):
    """Tests for find_pallet_count_column()."""

    def setUp(self):
        self.logger = logging.getLogger("test")
        self.columns = [
            MockColumnInfo(id="col_po", col_index=1),
            MockColumnInfo(id="col_desc", col_index=2),
            MockColumnInfo(id="col_pallet_count", col_index=3),
            MockColumnInfo(id="col_qty", col_index=4),
        ]

    def _make_worksheet(self, row_values: dict, max_column: int = 10):
        """
        Create a mock worksheet with values on a single row.
        
        Args:
            row_values: Dict of {col: value} for the footer row
        """
        ws = MagicMock()
        ws.max_column = max_column

        def mock_cell(row, column):
            cell = MagicMock()
            cell.value = row_values.get(column)
            cell.row = row
            cell.column = column
            return cell

        ws.cell = mock_cell
        return ws

    def test_finds_pallet_pattern(self):
        """Finds '25 PALLETS' and maps to correct column ID."""
        ws = self._make_worksheet({3: "25 PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_single_pallet(self):
        """Finds '1 PALLET' (singular)."""
        ws = self._make_worksheet({3: "1 PALLET"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_no_space_pattern(self):
        """Finds '12PALLETS' (no space)."""
        ws = self._make_worksheet({3: "12PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_no_pallet_returns_none(self):
        """Returns None when no pallet pattern found on the row."""
        ws = self._make_worksheet({1: "TOTAL OF:", 2: "LEATHER"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertIsNone(result)

    def test_unmapped_column_returns_none(self):
        """Pallet found in a column not covered by any ColumnInfo returns None."""
        ws = self._make_worksheet({8: "10 PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertIsNone(result)

    def test_finds_formula_based_pallet(self):
        """Finds pallet count from formula like '=SUM(D21:D25) & " PALLETS"'."""
        ws = self._make_worksheet({3: '=SUM(D21:D25) & " PALLETS"'})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_formula_singular_pallet(self):
        """Finds pallet count from formula like '=A1 & " PALLET"'."""
        ws = self._make_worksheet({3: '=A1 & " PALLET"'})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")


class TestFindHsCode(unittest.TestCase):
    """Tests for find_footer_hs_code()."""

    def _make_worksheet(self, cell_values: dict, max_column: int = 10):
        """
        Create a mock worksheet with specific cell values.
        """
        ws = MagicMock()
        ws.max_column = max_column

        def mock_cell(row, column):
            cell = MagicMock()
            cell.value = cell_values.get((row, column))
            cell.row = row
            cell.column = column
            return cell

        ws.cell = mock_cell
        return ws
        
    def _make_worksheet_with_merges(self, cell_values: dict, merge_ranges: list, max_column: int = 10):
        """
        Create a mock worksheet with specific cell values and merges.
        """
        ws = self._make_worksheet(cell_values, max_column)
        
        ranges = []
        for min_r, min_c, max_r, max_c in merge_ranges:
            r = MagicMock()
            r.min_row, r.min_col = min_r, min_c
            r.max_row, r.max_col = max_r, max_c
            ranges.append(r)
        ws.merged_cells.ranges = ranges
        return ws

    def test_finds_hs_code_exact(self):
        """Detects exact 'HS.CODE: XXXX'."""
        ws = self._make_worksheet({(5, 2): "HS.CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS.CODE: 4107.12.00")
        self.assertEqual(colspan, 1)

    def test_finds_hs_code_space(self):
        """Detects 'HS CODE: XXXX'."""
        ws = self._make_worksheet({(6, 3): "HS CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS CODE: 4107.12.00")

    def test_finds_hs_code_dash(self):
        """Detects 'HS-CODE: XXXX'."""
        ws = self._make_worksheet({(6, 3): "HS-CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS-CODE: 4107.12.00")

    def test_finds_hs_code_case_insensitive(self):
        """Detection is case-insensitive."""
        ws = self._make_worksheet({(4, 1): "hs code: 4107.12"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "hs code: 4107.12")
        
    def test_returns_colspan_if_merged(self):
        """Returns the correct colspan if the cell is merged."""
        ws = self._make_worksheet_with_merges(
            {(5, 2): "HS.CODE: 4107.12.00"}, 
            [(5, 2, 5, 3)] # cols 2 to 3 are merged
        )
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS.CODE: 4107.12.00")
        self.assertEqual(colspan, 2)

    def test_no_match_returns_none(self):
        """Returns None if no HS code keyword is found."""
        ws = self._make_worksheet({(1, 1): "Total:", (2, 1): "Grand Summary"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertIsNone(val)
        self.assertEqual(colspan, 1)


class TestExcelLayoutScannerHsCode(unittest.TestCase):
    """Tests that ExcelLayoutScanner scans HS codes for different sheets separately."""

    def setUp(self):
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL：", "TOTAL OF："]
            }
        }

    def _make_mock_sheet(self, sheet_name: str, header_row: int, cell_values: dict):
        ws = MagicMock()
        ws.title = sheet_name
        ws.max_row = max(r for r, c in cell_values.keys())
        ws.max_column = max(c for r, c in cell_values.keys())
        
        ws.row_dimensions = {
            row: MagicMock(height=30.0) for row in range(1, ws.max_row + 2)
        }
        ws.sheet_format = MagicMock()
        ws.sheet_format.defaultRowHeight = 15.0
        ws.merged_cells = MagicMock()
        ws.merged_cells.ranges = []

        def mock_cell(row, column):
            cell = MagicMock()
            cell.value = cell_values.get((row, column))
            cell.row = row
            cell.column = column
            font = MagicMock()
            font.name = "Calibri"
            font.size = 11.0
            font.bold = False
            font.italic = False
            cell.font = font
            align = MagicMock()
            align.horizontal = "center"
            align.wrap_text = False
            cell.alignment = align
            return cell

        ws.cell = mock_cell
        return ws

    def test_scans_separate_hs_codes_per_sheet(self):
        from core.blueprint_generator.internal.scanner import ExcelLayoutScanner
        
        # Create a mock workbook with 2 sheets
        wb = MagicMock()
        wb.sheetnames = ["Invoice", "PL"]
        
        # Sheet 1: Invoice has HS Code: 4107.12.00, merged 2 cols (col 5 to 6)
        # Header row: 3
        # Footer row: 10
        sheet1_values = {
            (3, 1): "Mark & No",
            (3, 2): "P.O.",
            (3, 3): "Item No.",
            (3, 4): "Description",
            (3, 5): "Unit Price",
            (3, 6): "Amount",
            (4, 1): "DES: COW LEATHER",
            (4, 4): "Product A",
            (8, 5): "HS CODE: 4107.12.00",
            (10, 2): "TOTAL:"
        }
        ws1 = self._make_mock_sheet("Invoice", 3, sheet1_values)
        
        # Sheet 2: PL has HS Code: 4114.10.00, merged 3 cols (col 5 to 7)
        # Header row: 3
        # Footer row: 10
        sheet2_values = {
            (3, 1): "Mark & No",
            (3, 2): "P.O.",
            (3, 3): "Item No.",
            (3, 4): "Description",
            (3, 5): "Net Weight",
            (3, 6): "Gross Weight",
            (3, 7): "CBM",
            (4, 1): "DES: COW LEATHER",
            (4, 4): "Product A",
            (8, 5): "HS CODE: 4114.10.00",
            (10, 2): "TOTAL:"
        }
        ws2 = self._make_mock_sheet("PL", 3, sheet2_values)
        
        # Let's mock merges for HS codes
        range1 = MagicMock()
        range1.min_row, range1.min_col, range1.max_row, range1.max_col = 8, 5, 8, 6
        range1.bounds = (5, 8, 6, 8)  # min_col, min_row, max_col, max_row
        ws1.merged_cells.ranges = [range1]
 
        range2 = MagicMock()
        range2.min_row, range2.min_col, range2.max_row, range2.max_col = 8, 5, 8, 7
        range2.bounds = (5, 8, 7, 8)  # min_col, min_row, max_col, max_row
        ws2.merged_cells.ranges = [range2]
 
        def mock_getitem(name):
            if name == "Invoice":
                return ws1
            if name == "PL":
                return ws2
            raise KeyError(name)
        wb.__getitem__.side_effect = mock_getitem
 
        scanner = ExcelLayoutScanner()
        result = scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", self.mapping_config, workbook=wb)
        
        # Verify both sheets were scanned and have different HS codes/colspans
        self.assertEqual(len(result.sheets), 2)
        
        sheet1_analysis = next(s for s in result.sheets if s.name == "Invoice")
        sheet2_analysis = next(s for s in result.sheets if s.name == "PL")
        
        self.assertIsNotNone(sheet1_analysis.footer_info)
        self.assertEqual(sheet1_analysis.footer_info.hs_code_text, "HS CODE: 4107.12.00")
        self.assertEqual(sheet1_analysis.footer_info.hs_code_colspan, 2)
        
        self.assertIsNotNone(sheet2_analysis.footer_info)
        self.assertEqual(sheet2_analysis.footer_info.hs_code_text, "HS CODE: 4114.10.00")
        self.assertEqual(sheet2_analysis.footer_info.hs_code_colspan, 3)

    def test_missing_description_raises_error(self):
        from core.blueprint_generator.internal.scanner import ExcelLayoutScanner
        wb = MagicMock()
        wb.sheetnames = ["Invoice"]
        
        # Missing DES: in col_static
        sheet1_values = {
            (3, 1): "Mark & No",
            (3, 2): "P.O.",
            (3, 3): "Item No.",
            (3, 4): "Description",
            (3, 5): "Unit Price",
            (3, 6): "Amount",
            (4, 1): "VENDOR#: ABC",
            (4, 4): "Product A",
            (8, 5): "HS CODE: 4107.12.00",
            (10, 2): "TOTAL:"
        }
        ws1 = self._make_mock_sheet("Invoice", 3, sheet1_values)
        wb.__getitem__.return_value = ws1

        scanner = ExcelLayoutScanner()
        # Should raise ValueError because missing description fallback is not bypassed
        with self.assertRaises(ValueError) as ctx:
            scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", self.mapping_config, workbook=wb)
        self.assertIn("Missing Description Fallback", str(ctx.exception))

    def test_missing_description_ignored_by_config(self):
        from core.blueprint_generator.internal.scanner import ExcelLayoutScanner
        wb = MagicMock()
        wb.sheetnames = ["Invoice"]
        
        # Missing DES: in col_static
        sheet1_values = {
            (3, 1): "Mark & No",
            (3, 2): "P.O.",
            (3, 3): "Item No.",
            (3, 4): "Description",
            (3, 5): "Unit Price",
            (3, 6): "Amount",
            (4, 1): "VENDOR#: ABC",
            (4, 4): "Product A",
            (8, 5): "HS CODE: 4107.12.00",
            (10, 2): "TOTAL:"
        }
        ws1 = self._make_mock_sheet("Invoice", 3, sheet1_values)
        wb.__getitem__.return_value = ws1

        scanner = ExcelLayoutScanner()
        # Set ignore_missing_description to True in config mapping
        custom_mapping = self.mapping_config.copy()
        custom_mapping["ignore_missing_description"] = True
        
        # Should not raise ValueError
        result = scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", custom_mapping, workbook=wb)
        self.assertEqual(len(result.sheets), 1)
        self.assertTrue(any("Missing Description Fallback" in w for w in result.warnings))


class TestExtractStaticColumnValues(unittest.TestCase):
    """Tests for extract_static_column_values()."""

    def setUp(self):
        self.columns = [
            MockColumnInfo(id="col_static", col_index=1),
            MockColumnInfo(id="col_po", col_index=2),
        ]

    def _make_worksheet(self, cell_values: dict, max_row: int = 20, max_column: int = 5):
        ws = MagicMock()
        ws.max_row = max_row
        ws.max_column = max_column

        def mock_cell(row, column):
            cell = MagicMock()
            cell.value = cell_values.get((row, column))
            cell.row = row
            cell.column = column
            return cell

        ws.cell = mock_cell
        return ws

    def test_extract_simple_static_values(self):
        """Extracts normal non-empty static cells and ignores empty ones."""
        cell_values = {
            (4, 1): "VENDOR#: CLW",
            (5, 1): "CASE QTY: 123",
            (6, 1): "MADE IN CAMBODIA",
        }
        ws = self._make_worksheet(cell_values)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["VENDOR#: CLW", "CASE QTY: 123", "MADE IN CAMBODIA"])

    def test_preserves_description_cells_as_is(self):
        """Preserves description cells exactly as they are without converting to placeholders."""
        cell_values = {
            (4, 1): "VENDOR#: ABC",
            (5, 1): "DES: COW LEATHER",
            (6, 1): "DES. : {{placeholder}}",
            (7, 1): "MADE IN CAMBODIA",
        }
        ws = self._make_worksheet(cell_values)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, [
            "VENDOR#: ABC",
            "DES: COW LEATHER",
            "DES. : {{placeholder}}",
            "MADE IN CAMBODIA"
        ])

    def test_stops_at_consecutive_empty_rows(self):
        """Scanning stops if 3 consecutive cells are empty."""
        cell_values = {
            (4, 1): "Line 1",
            (5, 1): "Line 2",
            (9, 1): "Line 3",  # 3 empty rows (6, 7, 8) in between
        }
        ws = self._make_worksheet(cell_values)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["Line 1", "Line 2"])

    def test_skips_vertically_merged_header_cells(self):
        """If the header row is vertically merged (e.g. spanning rows 3-4), scanning starts at row 5."""
        cell_values = {
            (3, 1): "STATIC HEADER", # merged header top cell
            (4, 1): "STATIC HEADER", # merged header cell value safe
            (5, 1): "Line 1",        # actual static content starts here
            (6, 1): "Line 2",
        }
        ws = self._make_worksheet(cell_values)
        
        # Mock vertically merged header spanning row 3 to 4, column 1 to 1
        mock_range = MagicMock()
        mock_range.min_row = 3
        mock_range.max_row = 4
        mock_range.min_col = 1
        mock_range.max_col = 1
        ws.merged_cells.ranges = [mock_range]
        
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["Line 1", "Line 2"])

    def test_stops_strictly_before_footer_row(self):
        """Scanning stops strictly before the specified footer/total row."""
        cell_values = {
            (4, 1): "Line 1",
            (5, 1): "Line 2",
            (6, 1): "TOTAL:",  # This is the footer row
            (7, 1): "Junk Line 3",  # This is beyond the footer row
        }
        ws = self._make_worksheet(cell_values)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns, footer_row=6)
        self.assertEqual(lines, ["Line 1", "Line 2"])


class TestExcelTemplateSanitizer(unittest.TestCase):
    """Tests for ExcelTemplateSanitizer sheet deletion and leak prevention."""

    def test_sanitize_keeps_at_least_one_sheet_and_clears_it(self):
        from core.blueprint_generator.internal.sanitizer import ExcelTemplateSanitizer
        from core.blueprint_generator.internal.scanner import TemplateAnalysisResult, SheetAnalysis
        import openpyxl

        # Create a real workbook with 2 sheets
        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Secret Customer Data 1"
        
        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Secret Customer Data 2"
        ws2.merge_cells("B1:C2")
        ws2["B1"] = "Merged Secret Data"

        # Mock the analysis so both sheets are analyzed (mapped)
        sheet_analysis1 = MagicMock(spec=SheetAnalysis)
        sheet_analysis1.name = "Sheet1"
        
        sheet_analysis2 = MagicMock(spec=SheetAnalysis)
        sheet_analysis2.name = "Sheet2"

        analysis = MagicMock(spec=TemplateAnalysisResult)
        analysis.customer_code = "TEST"
        analysis.sheets = [sheet_analysis1, sheet_analysis2]

        sanitizer = ExcelTemplateSanitizer()
        
        # Mock _clean_sheet to return dummy empty dict for each sheet
        sanitizer._clean_sheet = MagicMock(return_value={})

        # Sanitize the workbook
        cleaned_wb, layout = sanitizer.sanitize_template(wb, analysis)

        # It must only have 1 sheet remaining because both were mapped, but we cannot delete the last sheet
        self.assertEqual(len(cleaned_wb.sheetnames), 1)
        remaining_sheet_name = cleaned_wb.sheetnames[0]
        # The remaining sheet must have its secret data cleared!
        remaining_ws = cleaned_wb[remaining_sheet_name]
        self.assertIsNone(remaining_ws["A1"].value)
        self.assertIsNone(remaining_ws["B1"].value)


if __name__ == '__main__':
    unittest.main()
