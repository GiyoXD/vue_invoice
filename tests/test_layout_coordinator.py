import pytest
from openpyxl import Workbook
from core.invoice_generator.models.layout import SheetLayoutState
from core.models.cell import UnitRow, UnitCell, TemplateMerge

def test_layout_coordinator_binding():
    wb = Workbook()
    ws = wb.active
    state = SheetLayoutState()
    
    mapping = {'col_desc': 1, 'col_qty': 2}
    state.bind(ws, mapping, style_registry=None)
    
    assert state.ws == ws
    assert state.resolve_column('col_desc') == 1
    assert state.resolve_column('col_qty') == 2
    assert state.resolve_column(3) == 3

def test_layout_coordinator_merge_conflicts():
    wb = Workbook()
    ws = wb.active
    state = SheetLayoutState()
    state.bind(ws, {'col_a': 1, 'col_b': 2, 'col_c': 3}, style_registry=None)
    
    # Merge A1:B1 (row 1, cols 1-2)
    success = state.merge_cells(start_row=1, start_col_id_or_idx='col_a', end_row=1, end_col_id_or_idx='col_b')
    assert success is True
    assert (1, 1) in state.merged_cells
    assert (1, 2) in state.merged_cells
    assert (1, 2) in state.occupied_cells
    assert (1, 1) not in state.occupied_cells  # Anchor not swallowed
    
    # Try overlapping merge B1:C1 (blocked because B1 is in merged_cells)
    success2 = state.merge_cells(start_row=1, start_col_id_or_idx='col_b', end_row=1, end_col_id_or_idx='col_c')
    assert success2 is False

def test_layout_coordinator_write_cell_occupied():
    wb = Workbook()
    ws = wb.active
    state = SheetLayoutState()
    state.bind(ws, {'col_a': 1, 'col_b': 2}, style_registry=None)
    
    state.merge_cells(start_row=1, start_col_id_or_idx='col_a', end_row=1, end_col_id_or_idx='col_b')
    
    # Write to anchor (A1) -> allowed
    state.write_cell(row=1, col_id_or_idx='col_a', value='Anchor')
    assert ws.cell(row=1, column=1).value == 'Anchor'
    
    # Write to swallowed cell (B1) -> skipped
    state.write_cell(row=1, col_id_or_idx='col_b', value='Swallowed')
    assert ws.cell(row=1, column=2).value is None

def test_layout_coordinator_write_row_models():
    wb = Workbook()
    ws = wb.active
    state = SheetLayoutState()
    state.bind(ws, {'col_a': 1, 'col_b': 2}, style_registry=None)
    
    # Create row model with merged cell A1:B1
    merge_obj = TemplateMerge(min_col=1, max_col=2, row_span=1, value='Header')
    cell_a = UnitCell(col_index=1, value='Header', merge=merge_obj)
    cell_b = UnitCell(col_index=2, value=None)
    row_model = UnitRow(relative_index=0, height=15.0, cells=[cell_a, cell_b])
    
    state.write_row_models([row_model], start_row=1)
    
    # Verify merge registered in coordinator sets
    assert (1, 1) in state.merged_cells
    assert (1, 2) in state.merged_cells
    assert (1, 2) in state.occupied_cells
    assert ws.cell(row=1, column=1).value == 'Header'

def test_get_last_cell():
    wb = Workbook()
    ws = wb.active
    state = SheetLayoutState()
    state.bind(ws, {'col_a': 1, 'col_b': 5}, style_registry=None)
    
    state.advance_to(15)
    result = state.get_last_cell()
    assert result == (14, 'E')


def test_layout_builder_get_last_cell():
    from unittest.mock import MagicMock
    from core.invoice_generator.builders.layout_builder import LayoutBuilder
    
    builder = MagicMock(spec=LayoutBuilder)
    builder.layout_state = SheetLayoutState()
    builder.layout_state.advance_to(20)
    
    res = LayoutBuilder.get_last_cell(builder, col=" n ")
    assert res == (19, 'N')


def test_get_last_cell_explicit_write():
    from unittest.mock import MagicMock
    from core.invoice_generator.builders.layout_builder import LayoutBuilder

    wb = Workbook()
    ws = wb.active
    layout_state = SheetLayoutState()
    layout_state.bind(ws, {}, style_registry=None)

    ws.cell(row=50, column=10, value="test_end")

    res = layout_state.get_last_cell()
    assert res == (50, 'J')

    builder = MagicMock(spec=LayoutBuilder)
    builder.layout_state = layout_state
    builder_res = LayoutBuilder.get_last_cell(builder)
    assert builder_res == (50, 'J')



