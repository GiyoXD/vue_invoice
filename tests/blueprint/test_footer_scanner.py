"""
Tests for the footer_scanner module.
"""

import unittest
import logging
from unittest.mock import MagicMock
from dataclasses import dataclass
from typing import Optional

# --- Mock ColumnInfo to avoid importing the full scanner ---
@dataclass
class MockColumnInfo:
    """Lightweight mock matching ColumnInfo's interface."""
    id: str
    col_index: int
    colspan: int = 1


from core.blueprint_generator.internal.scanner.tabular_scanner import (
    find_column_id_by_index,
)
from core.blueprint_generator.utils.content_extractor import (
    find_total_label_cell,
    get_cell_merge_colspan,
    find_pallet_count_column,
    find_footer_hs_code,
    extract_static_column_values,
)


def make_mock_worksheet(
    cell_values: dict = None, 
    max_row: Optional[int] = None, 
    max_column: Optional[int] = None, 
    merge_ranges: list = None,
    sheet_name: str = "Sheet"
) -> MagicMock:
    """Helper to create a unified mock openpyxl worksheet."""
    ws = MagicMock()
    ws.title = sheet_name
    
    cells = cell_values or {}
    ws.max_row = max_row or (max(r for r, _ in cells.keys()) if cells and isinstance(next(iter(cells.keys())), tuple) else 1)
    ws.max_column = max_column or (max(c for _, c in cells.keys()) if cells and isinstance(next(iter(cells.keys())), tuple) else (max(cells.keys()) if cells else 1))
    
    ws.row_dimensions = {
        row: MagicMock(height=30.0) for row in range(1, ws.max_row + 2)
    }
    ws.sheet_format = MagicMock()
    ws.sheet_format.defaultRowHeight = 15.0
    
    ranges = []
    for range_item in (merge_ranges or []):
        r = MagicMock()
        if isinstance(range_item, tuple) and len(range_item) == 4:
            min_r, min_c, max_r, max_c = range_item
            r.min_row, r.min_col = min_r, min_c
            r.max_row, r.max_col = max_r, max_c
            r.bounds = (min_c, min_r, max_c, max_r)
        ranges.append(r)
    ws.merged_cells = MagicMock()
    ws.merged_cells.ranges = ranges

    def mock_cell(row, column):
        cell = MagicMock()
        if (row, column) in cells:
            cell.value = cells[(row, column)]
        elif column in cells and isinstance(next(iter(cells.keys())), int):
            cell.value = cells[column]
        else:
            cell.value = None
            
        cell.row = row
        cell.column = column
        
        font = MagicMock()
        font.name = "Calibri"
        font.size = 11.0
        font.bold = False
        font.italic = False
        cell.font = font
        
        align = MagicMock()
        align.horizontal = "center"
        align.wrap_text = False
        cell.alignment = align
        return cell

    ws.cell = mock_cell
    return ws


class TestFindColumnIdByIndex(unittest.TestCase):
    """Tests for find_column_id_by_index()."""

    def setUp(self):
        self.columns = [
            MockColumnInfo(id="col_po", col_index=1, colspan=1),
            MockColumnInfo(id="col_desc", col_index=2, colspan=2),  # Spans cols 2-3
            MockColumnInfo(id="col_qty", col_index=4, colspan=1),
            MockColumnInfo(id="col_amount", col_index=5, colspan=1),
        ]

    def test_exact_match(self):
        self.assertEqual(find_column_id_by_index(1, self.columns), "col_po")
        self.assertEqual(find_column_id_by_index(4, self.columns), "col_qty")
        self.assertEqual(find_column_id_by_index(5, self.columns), "col_amount")

    def test_colspan_range_match(self):
        self.assertEqual(find_column_id_by_index(2, self.columns), "col_desc")
        self.assertEqual(find_column_id_by_index(3, self.columns), "col_desc")

    def test_no_match_returns_none(self):
        self.assertIsNone(find_column_id_by_index(10, self.columns))

    def test_empty_columns_list(self):
        self.assertIsNone(find_column_id_by_index(1, []))


class TestFindTotalLabelCell(unittest.TestCase):
    """Tests for find_total_label_cell()."""

    def setUp(self):
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL：", "TOTAL OF："]
            }
        }

    def test_finds_total(self):
        ws = make_mock_worksheet({(5, 2): "TOTAL"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertEqual(cell.row, 5)
        self.assertEqual(cell.column, 2)
        self.assertTrue(is_exact)

    def test_finds_total_of_colon(self):
        ws = make_mock_worksheet({(8, 3): "TOTAL OF:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertEqual(cell.row, 8)
        self.assertTrue(is_exact)

    def test_finds_total_fullwidth_colon(self):
        ws = make_mock_worksheet({(3, 1): "TOTAL："})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertTrue(is_exact)

    def test_case_insensitive(self):
        ws = make_mock_worksheet({(4, 1): "total of:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertIsNotNone(cell)
        self.assertTrue(is_exact)

    def test_no_match_returns_none(self):
        ws = make_mock_worksheet({(1, 1): "Subtotal", (2, 1): "Grand Summary"})
        custom_config = {"footer_label_mappings": {"keywords": ["TOTAL OF:"]}}
        cell = find_total_label_cell(ws, 1, 10, custom_config)
        self.assertIsNone(cell)

    def test_ignores_partial_match(self):
        ws = make_mock_worksheet({(1, 1): "Total Net Weight"})
        custom_config = {"footer_label_mappings": {"keywords": ["TOTAL OF:"]}}
        cell = find_total_label_cell(ws, 1, 10, custom_config)
        self.assertIsNone(cell)

    def test_returns_first_match(self):
        ws = make_mock_worksheet({(3, 1): "TOTAL:", (7, 1): "TOTAL OF:"})
        cell, is_exact = find_total_label_cell(ws, 1, 10, self.mapping_config)
        self.assertEqual(cell.row, 3)


class TestGetCellMergeColspan(unittest.TestCase):
    """Tests for get_cell_merge_colspan()."""

    def test_merged_cell_returns_colspan(self):
        ws = make_mock_worksheet(merge_ranges=[(5, 2, 5, 4)])  # Cols 2-4 merged
        cell = MagicMock()
        cell.row, cell.column = 5, 2
        self.assertEqual(get_cell_merge_colspan(ws, cell), 3)

    def test_unmerged_cell_returns_1(self):
        ws = make_mock_worksheet(merge_ranges=[(5, 2, 5, 4)])
        cell = MagicMock()
        cell.row, cell.column = 3, 1
        self.assertEqual(get_cell_merge_colspan(ws, cell), 1)

    def test_no_merges_returns_1(self):
        ws = make_mock_worksheet(merge_ranges=[])
        cell = MagicMock()
        cell.row, cell.column = 1, 1
        self.assertEqual(get_cell_merge_colspan(ws, cell), 1)


class TestFindPalletCountColumn(unittest.TestCase):
    """Tests for find_pallet_count_column()."""

    def setUp(self):
        self.logger = logging.getLogger("test")
        self.columns = [
            MockColumnInfo(id="col_po", col_index=1),
            MockColumnInfo(id="col_desc", col_index=2),
            MockColumnInfo(id="col_pallet_count", col_index=3),
            MockColumnInfo(id="col_qty", col_index=4),
        ]

    def test_finds_pallet_pattern(self):
        ws = make_mock_worksheet({3: "25 PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_single_pallet(self):
        ws = make_mock_worksheet({3: "1 PALLET"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_no_space_pattern(self):
        ws = make_mock_worksheet({3: "12PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_no_pallet_returns_none(self):
        ws = make_mock_worksheet({1: "TOTAL OF:", 2: "LEATHER"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertIsNone(result)

    def test_unmapped_column_returns_none(self):
        ws = make_mock_worksheet({8: "10 PALLETS"})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertIsNone(result)

    def test_finds_formula_based_pallet(self):
        ws = make_mock_worksheet({3: '=SUM(D21:D25) & " PALLETS"'})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")

    def test_finds_formula_singular_pallet(self):
        ws = make_mock_worksheet({3: '=A1 & " PALLET"'})
        result = find_pallet_count_column(ws, 10, self.columns, find_column_id_by_index, self.logger)
        self.assertEqual(result, "col_pallet_count")


class TestFindHsCode(unittest.TestCase):
    """Tests for find_footer_hs_code()."""

    def test_finds_hs_code_exact(self):
        ws = make_mock_worksheet({(5, 2): "HS.CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS.CODE: 4107.12.00")
        self.assertEqual(colspan, 1)

    def test_finds_hs_code_space(self):
        ws = make_mock_worksheet({(6, 3): "HS CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS CODE: 4107.12.00")

    def test_finds_hs_code_dash(self):
        ws = make_mock_worksheet({(6, 3): "HS-CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS-CODE: 4107.12.00")

    def test_finds_hs_code_case_insensitive(self):
        ws = make_mock_worksheet({(4, 1): "hs code: 4107.12"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "hs code: 4107.12")
        
    def test_returns_colspan_if_merged(self):
        ws = make_mock_worksheet({(5, 2): "HS.CODE: 4107.12.00"}, merge_ranges=[(5, 2, 5, 3)])
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS.CODE: 4107.12.00")
        self.assertEqual(colspan, 2)

    def test_finds_hs_code_no_delimiter(self):
        ws = make_mock_worksheet({(5, 2): "HSCODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HSCODE: 4107.12.00")

    def test_finds_hs_code_underscore(self):
        ws = make_mock_worksheet({(5, 2): "HS_CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS_CODE: 4107.12.00")

    def test_finds_hs_code_spaces_between_letters(self):
        ws = make_mock_worksheet({(5, 2): "H S CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "H S CODE: 4107.12.00")

    def test_finds_hs_code_dot_colon(self):
        ws = make_mock_worksheet({(5, 2): "HS.CODE: 4107.12.00"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "HS.CODE: 4107.12.00")

    def test_finds_hs_code_substring(self):
        ws = make_mock_worksheet({(5, 2): "COMMODITY HSCODE: 4107"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertEqual(val, "COMMODITY HSCODE: 4107")

    def test_no_match_returns_none(self):
        ws = make_mock_worksheet({(1, 1): "Total:", (2, 1): "Grand Summary"})
        val, colspan, _ = find_footer_hs_code(ws, 1, 10)
        self.assertIsNone(val)
        self.assertEqual(colspan, 1)



class TestWorkbookManagerHsCode(unittest.TestCase):
    """Tests that WorkbookManager scans HS codes for different sheets separately."""

    def setUp(self):
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL：", "TOTAL OF："]
            }
        }

    def test_scans_separate_hs_codes_per_sheet(self):
        from core.blueprint_generator.internal.scanner import WorkbookManager
        
        wb = MagicMock()
        wb.sheetnames = ["Invoice", "PL"]
        
        sheet1_values = {
            (3, 1): "Mark & No", (3, 2): "P.O.", (3, 3): "Item No.", (3, 4): "Description", (3, 5): "Unit Price", (3, 6): "Amount",
            (4, 1): "DES: COW LEATHER", (4, 4): "Product A", (8, 5): "HS CODE: 4107.12.00", (10, 2): "TOTAL:"
        }
        ws1 = make_mock_worksheet(sheet1_values, merge_ranges=[(8, 5, 8, 6)], sheet_name="Invoice")
        
        sheet2_values = {
            (3, 1): "Mark & No", (3, 2): "P.O.", (3, 3): "Item No.", (3, 4): "Description", (3, 5): "Net Weight", (3, 6): "Gross Weight", (3, 7): "CBM",
            (4, 1): "DES: COW LEATHER", (4, 4): "Product A", (8, 5): "HS CODE: 4114.10.00", (10, 2): "TOTAL:"
        }
        ws2 = make_mock_worksheet(sheet2_values, merge_ranges=[(8, 5, 8, 7)], sheet_name="PL")
        
        def mock_getitem(name):
            if name == "Invoice": return ws1
            if name == "PL": return ws2
            raise KeyError(name)
        wb.__getitem__.side_effect = mock_getitem
 
        scanner = WorkbookManager()
        result = scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", self.mapping_config, workbook=wb)
        
        self.assertEqual(len(result.sheets), 2)
        sheet1_analysis = next(s for s in result.sheets if s.name == "Invoice")
        sheet2_analysis = next(s for s in result.sheets if s.name == "PL")
        
        self.assertIsNotNone(sheet1_analysis.footer_info)
        self.assertEqual(sheet1_analysis.footer_info.hs_code_text, "HS CODE: 4107.12.00")
        self.assertEqual(sheet1_analysis.footer_info.hs_code_colspan, 2)
        
        self.assertIsNotNone(sheet2_analysis.footer_info)
        self.assertEqual(sheet2_analysis.footer_info.hs_code_text, "HS CODE: 4114.10.00")
        self.assertEqual(sheet2_analysis.footer_info.hs_code_colspan, 3)

    def test_missing_description_raises_error(self):
        from core.blueprint_generator.internal.scanner import WorkbookManager
        wb = MagicMock()
        wb.sheetnames = ["Invoice"]
        
        sheet1_values = {
            (3, 1): "Mark & No", (3, 2): "P.O.", (3, 3): "Item No.", (3, 4): "Description", (3, 5): "Unit Price", (3, 6): "Amount",
            (4, 1): "VENDOR#: ABC", (4, 4): "Product A", (8, 5): "HS CODE: 4107.12.00", (10, 2): "TOTAL:"
        }
        ws1 = make_mock_worksheet(sheet1_values, sheet_name="Invoice")
        wb.__getitem__.return_value = ws1

        scanner = WorkbookManager()
        with self.assertRaises(ValueError) as ctx:
            scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", self.mapping_config, workbook=wb)
        self.assertIn("Missing Description Fallback", str(ctx.exception))

    def test_missing_description_ignored_by_config(self):
        from core.blueprint_generator.internal.scanner import WorkbookManager
        wb = MagicMock()
        wb.sheetnames = ["Invoice"]
        
        sheet1_values = {
            (3, 1): "Mark & No", (3, 2): "P.O.", (3, 3): "Item No.", (3, 4): "Description", (3, 5): "Unit Price", (3, 6): "Amount",
            (4, 1): "VENDOR#: ABC", (4, 4): "Product A", (8, 5): "HS CODE: 4107.12.00", (10, 2): "TOTAL:"
        }
        ws1 = make_mock_worksheet(sheet1_values, sheet_name="Invoice")
        wb.__getitem__.return_value = ws1

        scanner = WorkbookManager()
        custom_mapping = self.mapping_config.copy()
        custom_mapping["ignore_missing_description"] = True
        
        result = scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", custom_mapping, workbook=wb)
        self.assertEqual(len(result.sheets), 1)
        self.assertTrue(any("Missing Description Fallback" in w for w in result.warnings))


class TestExtractStaticColumnValues(unittest.TestCase):
    """Tests for extract_static_column_values()."""

    def setUp(self):
        self.columns = [
            MockColumnInfo(id="col_static", col_index=1),
            MockColumnInfo(id="col_po", col_index=2),
        ]

    def test_extract_simple_static_values(self):
        cell_values = {
            (4, 1): "VENDOR#: CLW",
            (5, 1): "CASE QTY: 123",
            (6, 1): "MADE IN CAMBODIA",
        }
        ws = make_mock_worksheet(cell_values, max_row=20, max_column=5)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["VENDOR#: CLW", "CASE QTY: 123", "MADE IN CAMBODIA"])

    def test_preserves_description_cells_as_is(self):
        cell_values = {
            (4, 1): "VENDOR#: ABC",
            (5, 1): "DES: COW LEATHER",
            (6, 1): "DES. : {{placeholder}}",
            (7, 1): "MADE IN CAMBODIA",
        }
        ws = make_mock_worksheet(cell_values, max_row=20, max_column=5)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["VENDOR#: ABC", "DES: COW LEATHER", "DES. : {{placeholder}}", "MADE IN CAMBODIA"])

    def test_stops_at_consecutive_empty_rows(self):
        cell_values = {
            (4, 1): "Line 1",
            (5, 1): "Line 2",
            (9, 1): "Line 3",
        }
        ws = make_mock_worksheet(cell_values, max_row=20, max_column=5)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["Line 1", "Line 2"])

    def test_skips_vertically_merged_header_cells(self):
        cell_values = {
            (3, 1): "STATIC HEADER",
            (4, 1): "STATIC HEADER",
            (5, 1): "Line 1",
            (6, 1): "Line 2",
        }
        ws = make_mock_worksheet(cell_values, max_row=20, max_column=5, merge_ranges=[(3, 1, 4, 1)])
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns)
        self.assertEqual(lines, ["Line 1", "Line 2"])

    def test_stops_strictly_before_footer_row(self):
        cell_values = {
            (4, 1): "Line 1",
            (5, 1): "Line 2",
            (6, 1): "TOTAL:",
            (7, 1): "Junk Line 3",
        }
        ws = make_mock_worksheet(cell_values, max_row=20, max_column=5)
        lines = extract_static_column_values(ws, header_row=3, columns=self.columns, footer_row=6)
        self.assertEqual(lines, ["Line 1", "Line 2"])


class TestExcelTemplateSanitizer(unittest.TestCase):
    """Tests for ExcelTemplateSanitizer sheet deletion and leak prevention."""

    def test_sanitize_keeps_at_least_one_sheet_and_clears_it(self):
        from core.blueprint_generator.internal.sanitizer import ExcelTemplateSanitizer
        from core.blueprint_generator.internal.scanner import TemplateAnalysisResult, SheetAnalysis
        import openpyxl

        wb = openpyxl.Workbook()
        ws1 = wb.active
        ws1.title = "Sheet1"
        ws1["A1"] = "Secret Customer Data 1"
        
        ws2 = wb.create_sheet("Sheet2")
        ws2["A1"] = "Secret Customer Data 2"
        ws2.merge_cells("B1:C2")
        ws2["B1"] = "Merged Secret Data"

        sanitizer = ExcelTemplateSanitizer()
        cleaned_wb = sanitizer.sanitize_template(wb, ["Sheet1", "Sheet2"])

        self.assertEqual(len(cleaned_wb.sheetnames), 1)
        remaining_sheet_name = cleaned_wb.sheetnames[0]
        remaining_ws = cleaned_wb[remaining_sheet_name]
        self.assertIsNone(remaining_ws["A1"].value)
        self.assertIsNone(remaining_ws["B1"].value)


class TestWorkbookManagerSheetClassification(unittest.TestCase):
    """Tests that WorkbookManager correctly classifies sheets including Summary Packing List."""

    def setUp(self):
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL：", "TOTAL OF："]
            },
            "aggregation_sheets": [
                "invoice", "contract", "inv", "commercial", "shipping", "bill", "summary_packing_list"
            ],
            "processed_tables_sheets": [
                "packing list", "packing", "pl", "detail", "content", "weight", "detail_packing_list"
            ]
        }

    def test_classifies_summary_packing_list_correctly(self):
        from core.blueprint_generator.internal.scanner import WorkbookManager
        
        wb = MagicMock()
        wb.sheetnames = ["Invoice", "PL", "Summary Packing List"]
        
        sheet_values = {
            (3, 1): "Mark & No", (3, 2): "P.O.", (3, 3): "Item No.", (3, 4): "Description", (3, 5): "Unit Price", (3, 6): "Amount",
            (4, 1): "DES: COW LEATHER", (4, 4): "Product A", (8, 5): "HS CODE: 4107.12.00", (10, 2): "TOTAL:"
        }
        
        ws_inv = make_mock_worksheet(sheet_values, sheet_name="Invoice")
        ws_pl = make_mock_worksheet(sheet_values, sheet_name="PL")
        ws_spl = make_mock_worksheet(sheet_values, sheet_name="Summary Packing List")
        
        def mock_getitem(name):
            if name == "Invoice": return ws_inv
            if name == "PL": return ws_pl
            if name == "Summary Packing List": return ws_spl
            raise KeyError(name)
        wb.__getitem__.side_effect = mock_getitem
        
        scanner = WorkbookManager()
        result = scanner.scan_template("tests/experiment_sample/shipping_list/JF25057.xlsx", self.mapping_config, workbook=wb)
        
        self.assertEqual(len(result.sheets), 3)
        sheet_inv = next(s for s in result.sheets if s.name == "Invoice")
        sheet_pl = next(s for s in result.sheets if s.name == "PL")
        sheet_spl = next(s for s in result.sheets if s.name == "Summary Packing List")
        
        self.assertEqual(sheet_inv.data_source, "aggregation")
        self.assertEqual(sheet_pl.data_source, "processed_tables_multi")
        self.assertEqual(sheet_spl.data_source, "summary_packing_list")
class TestBoundaryDetectorFooter(unittest.TestCase):
    """Tests for BoundaryDetector footer/boundary resolution logic."""

    def setUp(self):
        from core.blueprint_generator.internal.scanner.header_detector import BoundaryDetector
        self.detector = BoundaryDetector()
        self.mapping_config = {
            "footer_label_mappings": {
                "keywords": ["TOTAL", "TOTAL:", "TOTAL OF:", "TOTAL："]
            }
        }
        # A mock worksheet with header on row 3 and some columns
        self.base_cells = {
            (3, 1): "Item",
            (3, 2): "Qty",
            (3, 3): "Price",
            (3, 4): "Amount",
        }

    def test_detects_via_formula_adjacency(self):
        # Row 8 has adjacent formulas
        cells = self.base_cells.copy()
        cells.update({
            (4, 1): "A", (4, 2): 10, (4, 3): 5, (4, 4): 50,
            (8, 2): "=SUM(B4:B7)",
            (8, 3): "=SUM(C4:C7)",
        })
        ws = make_mock_worksheet(cells, max_row=10, max_column=5)
        
        self.detector.find_header_row = MagicMock(return_value=3)
        
        boundaries = self.detector.detect_boundaries(ws, mapping_config=self.mapping_config)
        self.assertIsNotNone(boundaries)
        self.assertEqual(boundaries.footer_row, 8)

    def test_detects_via_keyword_fallback(self):
        # No formula adjacency, but row 7 has "TOTAL:" keyword
        cells = self.base_cells.copy()
        cells.update({
            (4, 1): "A", (4, 2): 10, (4, 3): 5, (4, 4): 50,
            (7, 1): "TOTAL:",
        })
        ws = make_mock_worksheet(cells, max_row=10, max_column=5)
        
        self.detector.find_header_row = MagicMock(return_value=3)
        
        boundaries = self.detector.detect_boundaries(ws, mapping_config=self.mapping_config)
        self.assertIsNotNone(boundaries)
        self.assertEqual(boundaries.footer_row, 7)

    def test_detects_via_strict_bottom_up_total(self):
        # No formula adjacency, keyword not in mapping config, but row 9 has strict total "TOTAL AMOUNT:"
        cells = self.base_cells.copy()
        cells.update({
            (4, 1): "A", (4, 2): 10, (4, 3): 5, (4, 4): 50,
            (9, 1): "TOTAL AMOUNT:",
        })
        ws = make_mock_worksheet(cells, max_row=10, max_column=5)
        
        self.detector.find_header_row = MagicMock(return_value=3)
        
        # Pass empty mapping config so keyword match fails
        empty_config = {"footer_label_mappings": {"keywords": []}}
        boundaries = self.detector.detect_boundaries(ws, mapping_config=empty_config)
        self.assertIsNotNone(boundaries)
        self.assertEqual(boundaries.footer_row, 9)

    def test_detects_contiguous_footer_end_row(self):
        # Row 7 has TOTAL (detected as footer_row)
        # Row 8 has "COW LEATHER" and a numeric value (contiguous addon row)
        # Row 9 has "BUFFALO LEATHER" and a numeric value (contiguous addon row)
        # Row 10 is empty
        # Row 11 has a signature keyword (marks end)
        cells = self.base_cells.copy()
        cells.update({
            (4, 1): "A", (4, 2): 10, (4, 3): 5, (4, 4): 50,
            (7, 1): "TOTAL:", (7, 2): 10,
            (8, 1): "COW LEATHER", (8, 2): 5,
            (9, 1): "BUFFALO LEATHER", (9, 2): 5,
            (11, 1): "AUTHORIZED SIGNATURE",
        })
        ws = make_mock_worksheet(cells, max_row=12, max_column=5)
        
        self.detector.find_header_row = MagicMock(return_value=3)
        
        boundaries = self.detector.detect_boundaries(ws, mapping_config=self.mapping_config)
        self.assertIsNotNone(boundaries)
        self.assertEqual(boundaries.footer_row, 7)
        self.assertEqual(boundaries.footer_end_row, 9)


if __name__ == '__main__':
    unittest.main()
