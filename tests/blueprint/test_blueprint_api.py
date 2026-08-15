import io
import json
import pytest
import openpyxl
from io import BytesIO
from core.database.db_manager import Blueprint

def test_blueprint_scan_and_generate_flow(client, db):
    # 1. Create a minimal valid Excel template in memory that matches scanner expectations
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    # Headers at Row 3
    ws.cell(row=3, column=1, value="Mark & No")
    ws.cell(row=3, column=2, value="P.O. No.")
    ws.cell(row=3, column=3, value="Random Unmapped Column")  # Unknown header to trigger 'needs_mapping'
    ws.cell(row=3, column=4, value="Quantity")
    ws.cell(row=3, column=5, value="Unit Price")
    ws.cell(row=3, column=6, value="Amount")
    
    # Description Fallback in Column 1, Row 4 (part of static column values)
    ws.cell(row=4, column=1, value="DES: COW LEATHER")
    ws.cell(row=4, column=3, value="Product A")
    ws.cell(row=4, column=4, value=10)
    ws.cell(row=4, column=5, value=5.0)
    ws.cell(row=4, column=6, value=50.0)
    
    # HS Code in Row 8
    ws.cell(row=8, column=5, value="HS CODE: 4107.12.00")
    
    # Footer Row at Row 10
    ws.cell(row=10, column=2, value="TOTAL:")
    ws.cell(row=10, column=6, value=50.0)
    
    buf = BytesIO()
    wb.save(buf)
    xlsx_bytes = buf.getvalue()

    # 2. Upload the file to Step 1: /api/blueprint/scan
    response = client.post(
        "/api/blueprint/scan",
        files={"file": ("test_api_template.xlsx", xlsx_bytes, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    scan_data = response.json()
    
    assert scan_data["status"] == "needs_mapping"
    assert scan_data["file_token"] == "scan_test_api_template.xlsx"
    assert "Random Unmapped Column" in scan_data["unknown_headers"]

    # 3. Post user mappings to Step 2: /api/blueprint/generate
    generate_payload = {
        "file_token": scan_data["file_token"],
        "customer_code": "APITESTGEN",
        "locale": "KH",
        "mappings": {
            "P.O. No.": "col_po",
            "Random Unmapped Column": "col_desc",
            "Quantity": "col_qty_sf",
            "Unit Price": "col_unit_price",
            "Amount": "col_amount"
        },
        "footer_mappings": [],
        "pricing_mode": "standard"
    }
    
    response = client.post("/api/blueprint/generate", json=generate_payload)
    assert response.status_code == 200
    gen_data = response.json()
    assert gen_data["status"] == "success"
    assert "generated and saved to database" in gen_data["message"]

    # 4. Verify that the generated template is saved in the database and readable
    response = client.get("/api/template/view?customer_code=APITESTGEN&locale=KH")
    assert response.status_code == 200
    template_data = response.json()
    assert "template_layout" in template_data
    assert "Invoice" in template_data["template_layout"]

def test_unrecognized_sheets_detection(client, db):
    wb = openpyxl.Workbook()
    ws1 = wb.active
    ws1.title = "Invoice"
    ws1.cell(row=3, column=1, value="Mark & No")
    ws1.cell(row=3, column=2, value="P.O. No.")
    ws1.cell(row=3, column=3, value="Quantity")
    ws1.cell(row=3, column=4, value="Unit Price")
    ws1.cell(row=3, column=5, value="Amount")
    ws1.cell(row=4, column=1, value="DES: LEATHER")
    ws1.cell(row=8, column=5, value="HS CODE: 1234")
    ws1.cell(row=10, column=2, value="TOTAL:")
    
    ws2 = wb.create_sheet(title="UnknownSheet123")
    ws2.cell(row=1, column=1, value="Some Data")
    
    buf = BytesIO()
    wb.save(buf)
    
    response = client.post(
        "/api/template/analyze",
        files={"file": ("test_unrecognized.xlsx", buf.getvalue(), "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
    )
    assert response.status_code == 200
    data = response.json()
    assert "unrecognized_sheets" in data
    assert "UnknownSheet123" in data["unrecognized_sheets"]

def test_analyze_existing_template(client, db):
    from api.routers.upload import get_source_folder
    folder = get_source_folder(db)
    
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    ws.cell(row=3, column=1, value="Mark & No")
    ws.cell(row=3, column=2, value="P.O. No.")
    ws.cell(row=3, column=3, value="Quantity")
    ws.cell(row=3, column=4, value="Unit Price")
    ws.cell(row=3, column=5, value="Amount")
    ws.cell(row=4, column=1, value="DES: COW LEATHER")
    ws.cell(row=8, column=5, value="HS CODE: 4107.12.00")
    ws.cell(row=10, column=2, value="TOTAL:")
    ws.cell(row=10, column=5, value=50.0)
    
    dummy_filename = "test_existing_template.xlsx"
    file_path = folder / dummy_filename
    wb.save(file_path)
    
    try:
        # Test valid existing template analysis
        response = client.post(
            "/api/template/analyze-existing",
            json={"filename": dummy_filename, "ignore_missing_description": False}
        )
        assert response.status_code == 200
        data = response.json()
        assert "missing_headers" in data
        assert "missing_footers" in data
        assert "temp_filename" in data
        
        # Test 404 when file does not exist
        response_404 = client.post(
            "/api/template/analyze-existing",
            json={"filename": "non_existent_file.xlsx"}
        )
        assert response_404.status_code == 404
    finally:
        if file_path.exists():
            try:
                file_path.unlink()
            except Exception:
                pass


    
