import json
import pytest
import openpyxl
from pathlib import Path
from unittest.mock import MagicMock
from core.blueprint_generator.generator import BlueprintGenerator, BlueprintGenerationOptions
from core.blueprint_generator.internal.scanner import TemplateAnalysisResult, SheetAnalysis
from core.database.db_manager import Blueprint, BlueprintTemplate

def create_valid_test_xlsx(path: Path):
    """Creates a minimal valid Excel template on disk for generator tests."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    # Headers at Row 3
    ws.cell(row=3, column=1, value="Mark & No")
    ws.cell(row=3, column=2, value="P.O. No.")
    ws.cell(row=3, column=3, value="Quantity")
    ws.cell(row=3, column=4, value="Unit Price")
    ws.cell(row=3, column=5, value="Amount")
    
    # Static column value (Description Fallback) in Column 1, Row 4
    ws.cell(row=4, column=1, value="DES: MOCK LEATHER PRODUCT")
    ws.cell(row=4, column=2, value="PO-12345")
    ws.cell(row=4, column=3, value=10)
    ws.cell(row=4, column=4, value=15.0)
    ws.cell(row=4, column=5, value=150.0)
    
    # HS Code in Row 8
    ws.cell(row=8, column=4, value="HS CODE: 4107.12.00")
    
    # Footer Row at Row 10
    ws.cell(row=10, column=2, value="TOTAL:")
    ws.cell(row=10, column=5, value=150.0)
    
    wb.save(path)


def test_generator_load_mapping_config():
    generator = BlueprintGenerator()
    config = generator._load_mapping_config()
    assert isinstance(config, dict)
    assert "footer_label_mappings" in config
    assert "header_text_mappings" in config


def test_generator_build_table_info():
    generator = BlueprintGenerator()
    
    # Mock template analysis result with 2 sheets
    # PL sheet: has description fallback "MOCK DESC"
    # INV sheet: has description fallback "INV DESC"
    # Sourcing algorithm should prefer "Packing list" sheet description fallback
    
    sheet_pl = MagicMock(spec=SheetAnalysis)
    sheet_pl.name = "Packing list"
    sheet_pl.static_content_hints = {"description_fallback": "MOCK DESC"}
    sheet_pl.footer_info = MagicMock()
    sheet_pl.footer_info.hs_code_text = "HS 4107"
    
    sheet_inv = MagicMock(spec=SheetAnalysis)
    sheet_inv.name = "Invoice"
    sheet_inv.static_content_hints = {"description_fallback": "INV DESC"}
    sheet_inv.footer_info = MagicMock()
    sheet_inv.footer_info.hs_code_text = "HS 9999"
    
    analysis = MagicMock(spec=TemplateAnalysisResult)
    analysis.sheets = [sheet_inv, sheet_pl]
    
    table_info = generator._build_table_info(analysis)
    assert table_info["fallback_description"]["standard"] == "MOCK DESC"
    assert table_info["hs_code"] == "HS 4107"


def test_generator_preserve_user_overrides():
    generator = BlueprintGenerator()
    
    # Mock newly generated layout metadata
    layout_metadata = {
        "Invoice": {
            "header_rows": [
                {
                    "relative_index": 0,
                    "cells": [
                        {"col_index": 1, "value": "New Value A1"},
                        {"col_index": 2, "value": "New Value B2"}
                    ]
                }
            ],
            "footer_rows": [
                {
                    "relative_index": 0,
                    "cells": [
                        {"col_index": 1, "value": "New Footer 1"}
                    ]
                }
            ],
            "col_widths": {}
        }
    }
    
    # Mock old template config data
    old_data = {
        "notes": "User saved notes",
        "template_layout": {
            "Invoice": {
                "header_rows": [
                    {
                        "relative_index": 0,
                        "cells": [
                            {
                                "col_index": 1,
                                "value": {"default": "Old Plain", "standard": "Override Standard", "daf": "Override Daf"}
                            }
                        ]
                    }
                ],
                "footer_rows": [
                    {
                        "relative_index": 0,
                        "cells": [
                            {
                                "col_index": 1,
                                "value": {"default": "Old Footer", "standard": "Override Footer Standard"}
                            }
                        ]
                    }
                ],
                "col_widths": {}
            }
        }
    }
    
    preserved_notes = generator._preserve_user_overrides(None, layout_metadata, old_data=old_data)
    
    assert preserved_notes == "User saved notes"
    sheet = layout_metadata["Invoice"]
    
    # A1 override should merge the NEW default value with the OLD standard/daf overrides
    row0 = next(r for r in sheet["header_rows"] if r["relative_index"] == 0)
    cell_a1 = next(c for c in row0["cells"] if c["col_index"] == 1)
    assert isinstance(cell_a1["value"], dict)
    assert cell_a1["value"]["default"] == "New Value A1"
    assert cell_a1["value"]["standard"] == "Override Standard"
    assert cell_a1["value"]["daf"] == "Override Daf"
    
    # B2 was not overridden in old, so it remains a plain string
    cell_b2 = next(c for c in row0["cells"] if c["col_index"] == 2)
    assert cell_b2["value"] == "New Value B2"
    
    # Footer row relative_index=0 col_index=1 should merge the NEW default with the OLD overrides
    frow0 = next(r for r in sheet["footer_rows"] if r["relative_index"] == 0)
    fcell1 = next(c for c in frow0["cells"] if c["col_index"] == 1)
    assert isinstance(fcell1["value"], dict)
    assert fcell1["value"]["default"] == "New Footer 1"
    assert fcell1["value"]["standard"] == "Override Footer Standard"


def test_generator_analyze_and_generate_in_memory(tmp_path):
    template_file = tmp_path / "MOCK_template.xlsx"
    create_valid_test_xlsx(template_file)
    
    generator = BlueprintGenerator()
    
    # Test analyze
    json_result_str = generator.analyze(str(template_file))
    json_data = json.loads(json_result_str)
    assert "MOCK_template.xlsx" in json_data["file_path"]
    assert "sheets" in json_data
    assert any(sheet["sheet_name"] == "Invoice" for sheet in json_data["sheets"])
    
    # Test generate in_memory
    bundle, template_json_data, template_xlsx_bytes = generator.generate(
        template_path=str(template_file),
        options=BlueprintGenerationOptions(in_memory=True, pricing_mode="standard")
    )
    
    assert isinstance(bundle, dict)
    assert bundle["_meta"]["customer"] == "MOCK_TEMPLATE"
    assert bundle["_meta"]["pricing_mode"] == "standard"
    assert "table_info" in bundle
    
    assert isinstance(template_json_data, dict)
    assert "Invoice" in template_json_data["template_layout"]
    assert isinstance(template_xlsx_bytes, bytes)


def test_generator_generate_to_disk(tmp_path):
    template_file = tmp_path / "MOCK_template.xlsx"
    create_valid_test_xlsx(template_file)
    
    output_dir = tmp_path / "output_bundled"
    generator = BlueprintGenerator(output_base_dir=output_dir)
    
    # Test file-based generation
    config_file_path = generator.generate(
        template_path=str(template_file),
        options=BlueprintGenerationOptions(output_dir=str(output_dir), custom_prefix="CUSTOMCUST")
    )
    
    assert config_file_path is not None
    assert config_file_path.exists()
    
    # Check directory structure and file names
    config_dir = output_dir / "CUSTOMCUST"
    assert config_dir.exists()
    
    config_file = config_dir / "CUSTOMCUST_KH_config.json"
    template_file_json = config_dir / "CUSTOMCUST_KH_template.json"
    template_xlsx = config_dir / "CUSTOMCUST_KH.xlsx"
    
    assert config_file.exists()
    assert template_file_json.exists()
    assert template_xlsx.exists()
    
    # Verify content
    with open(config_file, 'r', encoding='utf-8') as f:
        config_data = json.load(f)
        assert config_data["_meta"]["customer"] == "CUSTOMCUST"
