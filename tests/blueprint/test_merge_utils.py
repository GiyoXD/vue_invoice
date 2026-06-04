import pytest
import openpyxl
from core.blueprint_generator.utils.merge_utils import (
    store_original_merges,
    find_and_restore_merges_heuristic,
    force_unmerge_from_row_down,
    MergeOffsetTracker,
    store_empty_merges_with_coordinates,
    restore_empty_merges_with_offset
)

def test_merge_offset_tracker():
    tracker = MergeOffsetTracker()
    
    # Test tracking deletes
    tracker.log_delete_rows(start_row=15, count=5, sheet_name="Invoice")
    # Row before delete is unaffected
    assert tracker.calculate_new_position(12, "Invoice") == 12
    # Row deleted is mapped to -1
    assert tracker.calculate_new_position(17, "Invoice") == -1
    # Row after delete is shifted up by 5
    assert tracker.calculate_new_position(22, "Invoice") == 17
    # Shift for different sheet is unaffected
    assert tracker.calculate_new_position(22, "PL") == 22

    # Test tracking inserts
    tracker.log_insert_rows(position=10, count=3, sheet_name="Invoice")
    # Row before insert is unaffected
    assert tracker.calculate_new_position(5, "Invoice") == 5
    # Row after insert is shifted down by 3
    assert tracker.calculate_new_position(11, "Invoice") == 14


def test_store_and_restore_merges_heuristic():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    # Create some merged ranges (min_row >= 10, max_row == min_row)
    ws.cell(row=12, column=2, value="MergedValue1")
    ws.merge_cells("B12:D12")  # Cols 2 to 4 (colspan = 3)
    ws.row_dimensions[12].height = 25.0
    
    ws.cell(row=15, column=3, value="MergedValue2")
    ws.merge_cells("C15:E15")  # Cols 3 to 5 (colspan = 3)
    
    # Store merges
    stored = store_original_merges(wb, ["Invoice"])
    assert "Invoice" in stored
    assert len(stored["Invoice"]) == 2
    
    # Verify stored info
    item1 = next(item for item in stored["Invoice"] if item[1] == "MergedValue1")
    assert item1[0] == 3  # colspan
    assert item1[2] == 25.0  # height
    
    # Unmerge them manually to simulate a generation/sanitization process
    ws.unmerge_cells("B12:D12")
    ws.unmerge_cells("C15:E15")
    assert len(ws.merged_cells.ranges) == 0
    
    # Restore merges heuristically
    find_and_restore_merges_heuristic(wb, stored, ["Invoice"])
    
    # Check that they are merged again
    merged_ranges = list(ws.merged_cells.ranges)
    assert len(merged_ranges) == 2
    
    bounds = [r.bounds for r in merged_ranges]
    assert (2, 12, 4, 12) in bounds  # B12:D12 restored
    assert (3, 15, 5, 15) in bounds  # C15:E15 restored
    assert ws.row_dimensions[12].height == 25.0


def test_force_unmerge_from_row_down():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    # Merges above and below start_row=15
    ws.merge_cells("B5:C5")
    ws.merge_cells("D15:E15")
    ws.merge_cells("F20:G20")
    
    assert len(ws.merged_cells.ranges) == 3
    
    force_unmerge_from_row_down(ws, 15)
    
    merged_ranges = list(ws.merged_cells.ranges)
    assert len(merged_ranges) == 1
    assert merged_ranges[0].bounds == (2, 5, 3, 5)


def test_store_and_restore_empty_merges_with_offset():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    
    # Create empty merge at row 20
    ws.merge_cells("C20:E20")  # Cols 3 to 5 (colspan = 3)
    ws.row_dimensions[20].height = 18.0
    
    # Store empty merges
    stored = store_empty_merges_with_coordinates(wb, ["Invoice"])
    assert "Invoice" in stored
    assert len(stored["Invoice"]) == 1
    assert stored["Invoice"][0]["original_row"] == 20
    assert stored["Invoice"][0]["span"] == 3
    assert stored["Invoice"][0]["height"] == 18.0
    
    # Simulate row deletion of 5 rows at row 10
    tracker = MergeOffsetTracker()
    tracker.log_delete_rows(start_row=10, count=5, sheet_name="Invoice")
    
    # Clear merges
    ws.unmerge_cells("C20:E20")
    assert len(ws.merged_cells.ranges) == 0
    
    # Restore empty merges using tracker
    restore_empty_merges_with_offset(wb, stored, tracker, ["Invoice"])
    
    # The new row index should be 20 - 5 = 15
    merged_ranges = list(ws.merged_cells.ranges)
    assert len(merged_ranges) == 1
    assert merged_ranges[0].bounds == (3, 15, 5, 15)  # C15:E15
    assert ws.row_dimensions[15].height == 18.0
