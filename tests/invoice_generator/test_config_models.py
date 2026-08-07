import json
from pathlib import Path
import pytest
from pydantic import ValidationError

from core.invoice_generator.models.config import ClientConfigBundle

def test_load_and_validate_bundle_config():
    # 1. Resolve path to JF_bundle_config.json
    config_path = Path(__file__).parent.parent.parent / "core" / "invoice_generator" / "JF_bundle_config.json"
    assert config_path.exists(), f"Config file not found at: {config_path}"

    # 2. Load JSON dict
    with open(config_path, "r", encoding="utf-8") as f:
        config_data = json.load(f)

    # 3. Parse via Pydantic model
    bundle = ClientConfigBundle.model_validate(config_data)

    # 4. Verify Meta
    assert bundle.meta.customer == "JF"
    assert bundle.meta.config_version == "2.2_strict_mode"
    assert bundle.meta.has_static_sheets is False

    # 5. Verify Processing
    assert bundle.processing.sheets == ["Invoice", "Contract", "Packing list"]
    assert bundle.processing.data_sources["Invoice"] == "aggregation"
    assert bundle.processing.data_sources["Packing list"] == "processed_tables_multi"

    # 6. Verify Data Preparation Hint (Typo alias check)
    assert bundle.meta.data_preparation_module_hint is not None
    assert bundle.meta.data_preparation_module_hint.priority == ["po"]
    assert bundle.meta.data_preparation_module_hint.numbers_per_group_by_po == 7

    # 7. Verify Layout configurations
    invoice_layout = bundle.layout_bundle.get("Invoice")
    assert invoice_layout is not None
    # We can check that the type is SheetLayoutModel and not GlobalDefaultsModel
    assert hasattr(invoice_layout, "structure")
    assert invoice_layout.structure.header_row == 21
    
    # Check column width mapping
    columns = invoice_layout.structure.columns
    assert len(columns) > 0
    assert columns[0].id == "col_static"
    assert columns[0].width == 24.71
    assert columns[1].id == "col_po"
    assert columns[1].header == "P.O. Nº"

def test_load_and_validate_master_config():
    config_path = Path(__file__).parent.parent.parent / "database" / "blueprints" / "mapper" / "master_config.json"
    assert config_path.exists(), f"Config file not found at: {config_path}"
    with open(config_path, "r", encoding="utf-8") as f:
        config_data = json.load(f)
    bundle = ClientConfigBundle.model_validate(config_data)
    assert bundle.meta.customer == "JF"
    assert bundle.meta.config_version == "2.2_strict_mode"

def test_validation_failure():
    # Verify that invalid structure raises ValidationError
    # processing values must be str or Dict[str, str], not int
    invalid_data = {
        "_meta": {
            "config_version": "2.2",
            "customer": "Test"
        },
        "processing": {
            "Sheet1": 12345
        },
        "styling_bundle": {},
        "layout_bundle": {}
    }
    with pytest.raises(ValidationError):
        ClientConfigBundle.model_validate(invalid_data)


def test_deep_copy_worksheet_copies_print_area():
    from core.invoice_generator.utils.workbook_utils import deep_copy_worksheet
    import openpyxl

    source_wb = openpyxl.Workbook()
    source_ws = source_wb.active
    source_ws.title = "StaticSheet"
    source_ws.print_area = "A1:G50"
    source_ws.page_setup.orientation = source_ws.ORIENTATION_LANDSCAPE
    source_ws.page_setup.paperSize = source_ws.PAPERSIZE_A4

    target_wb = openpyxl.Workbook()
    target_ws = target_wb.create_sheet("StaticSheet")

    deep_copy_worksheet(source_ws, target_ws)

    assert target_ws.print_area == source_ws.print_area
    assert target_ws.page_setup.orientation == source_ws.page_setup.orientation
    assert target_ws.page_setup.paperSize == source_ws.page_setup.paperSize




