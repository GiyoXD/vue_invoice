import pytest
from unittest.mock import patch
import openpyxl

from core.invoice_generator.utils.formula_formatter import shift_formula_coordinates
from core.blueprint_generator.internal.scanner.models import TemplateLayout
from core.invoice_generator.builders.json_template_builder import JsonTemplateStateBuilder
from core.models.cell import UnitRow, UnitCell


class TestShiftFormulaCoordinates:
    """Unit tests for shift_formula_coordinates function."""

    def test_zero_offset_returns_unchanged(self):
        """Zero offset should return formula string unchanged."""
        formula = "=SUM(D16:D20)"
        result = shift_formula_coordinates(
            formula,
            row_offset=0,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result == formula

        formula_header = "=D4 + E15"
        result_header = shift_formula_coordinates(
            formula_header,
            row_offset=0,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result_header == formula_header

    def test_header_references_not_shifted(self):
        """Header references (row <= header_end_row) remain unchanged while footer references shift."""
        # row 4 <= header_end_row (5): untouched
        # row 15 >= original_footer_start_row (15): shifted +20 -> 35
        formula = "=D4 + E15"
        result = shift_formula_coordinates(
            formula,
            row_offset=20,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=D4 + E35"

    def test_lower_template_references_shifted(self):
        """Lower template references (row >= orig_footer) are shifted by row_offset."""
        formula = "=SUM(D16:D20)"
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result == "=SUM(D26:D30)"

    def test_absolute_cell_references_shifted(self):
        """Absolute cell references (with $ signs) shift row coordinates while preserving dollar signs."""
        formula = "=$D$4 + $E$15"
        result = shift_formula_coordinates(
            formula,
            row_offset=20,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=$D$4 + $E$35"

        formula_range = "=SUM($D$16:$D$20)"
        result_range = shift_formula_coordinates(
            formula_range,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result_range == "=SUM($D$26:$D$30)"

    def test_mixed_reference_coordinates(self):
        """Mixed references ($ on column only or row only) are correctly preserved and shifted."""
        formula = "=$D4 + E$15"
        result = shift_formula_coordinates(
            formula,
            row_offset=20,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=$D4 + E$35"

    def test_non_formula_and_non_string_values_unchanged(self):
        """Non-formula strings, empty strings, numbers, None, and booleans should be returned unchanged."""
        assert shift_formula_coordinates("Total Amount", 20, 5, 15) == "Total Amount"
        assert shift_formula_coordinates("D4 + E15", 20, 5, 15) == "D4 + E15"
        assert shift_formula_coordinates("", 20, 5, 15) == ""
        assert shift_formula_coordinates(None, 20, 5, 15) is None
        assert shift_formula_coordinates(12345, 20, 5, 15) == 12345
        assert shift_formula_coordinates(99.9, 20, 5, 15) == 99.9
        assert shift_formula_coordinates(True, 20, 5, 15) is True

    def test_functions_with_numbers_not_shifted(self):
        """Function names with digits (e.g. LOG10) are guarded and not treated as cell references."""
        formula = "=LOG10(A15) + A15"
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=LOG10(A25) + A25"

        formula_header = "=LOG10(A5) + A15"
        result_header = shift_formula_coordinates(
            formula_header,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result_header == "=LOG10(A5) + A25"

    def test_functions_with_space_before_paren_not_shifted(self):
        """Function names followed by space and paren (e.g. LOG10 (A15)) are not treated as cell references."""
        formula = "=LOG10 (A15) + A15"
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=LOG10 (A25) + A25"

    def test_quoted_string_literals_masked_and_preserved(self):
        """String literals in quotes are masked so text inside quotes is never altered."""
        formula = '=IF(A15="D16", "D16", D16)'
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == '=IF(A25="D16", "D16", D26)'

    def test_sheet_qualifiers_not_shifted(self):
        """Sheet names ending with numbers (e.g. S1!) are external sheets and their cell coordinates must not be shifted."""
        formula = "=S1!A15"
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "=S1!A15"

    def test_quoted_sheet_reference_not_shifted(self):
        """Quoted sheet names ending with numbers (e.g. 'S1'!) have external cell coordinates frozen while local cells shift."""
        formula = "='S1'!A15 + A15"
        result = shift_formula_coordinates(
            formula,
            row_offset=10,
            header_end_row=5,
            original_footer_start_row=15
        )
        assert result == "='S1'!A15 + A25"

    def test_quoted_sheet_reference_with_spaces_and_tokens(self):
        """Quoted sheet names with spaces and coordinate-like tokens are masked and not shifted."""
        assert shift_formula_coordinates("='S1 Summary'!A15 + A15", 10, header_end_row=5, original_footer_start_row=15) == "='S1 Summary'!A15 + A25"
        assert shift_formula_coordinates("='[B10 File.xlsx]Sheet1'!A15 + A15", 10, header_end_row=5, original_footer_start_row=15) == "='[B10 File.xlsx]Sheet1'!A15 + A25"

    def test_external_sheet_references_never_shifted(self):
        """External sheet references (single cells, ranges, lookups, external workbooks) never shift."""
        # Single cell on external sheet
        assert shift_formula_coordinates("=deepsheet!A1 + A15", 10, header_end_row=5, original_footer_start_row=15) == "=deepsheet!A1 + A25"
        # Range on external sheet
        assert shift_formula_coordinates("=SUM(deepsheet!A1:B10) + SUM(A15:B20)", 10, header_end_row=5, original_footer_start_row=15) == "=SUM(deepsheet!A1:B10) + SUM(A25:B30)"
        # Quoted external sheet range
        assert shift_formula_coordinates("=SUM('deepsheet'!$A$1:$B$10) + A15", 10, header_end_row=5, original_footer_start_row=15) == "=SUM('deepsheet'!$A$1:$B$10) + A25"
        # Lookup on external sheet
        assert shift_formula_coordinates("=VLOOKUP(A15, 'deepsheet'!$A$1:$D$50, 2, FALSE)", 10, header_end_row=5, original_footer_start_row=15) == "=VLOOKUP(A25, 'deepsheet'!$A$1:$D$50, 2, FALSE)"
        # External workbook ref
        assert shift_formula_coordinates("='[Budget.xlsx]deepsheet'!A1 + A15", 10, header_end_row=5, original_footer_start_row=15) == "='[Budget.xlsx]deepsheet'!A1 + A25"

    def test_escaped_double_quotes_inside_strings(self):
        """String literals containing escaped double quotes preserve coordinate strings inside."""
        assert shift_formula_coordinates('="He said ""D16""" & D16', 10, header_end_row=5, original_footer_start_row=15) == '="He said ""D16""" & D26'

    def test_multi_range_comma_separated_formula(self):
        """Multi-range comma-separated formula arguments are all properly shifted."""
        assert shift_formula_coordinates("=SUM(A15:B20, C15:D20)", 5, header_end_row=5, original_footer_start_row=15) == "=SUM(A20:B25, C20:D25)"

    def test_negative_offsets_clamped_to_ref_error(self):
        """Negative offsets resulting in row <= 0 clamp reference to #REF!."""
        assert shift_formula_coordinates("=A5", -10) == "=#REF!"
        assert shift_formula_coordinates("=A5", -10, original_footer_start_row=5) == "=#REF!"
        assert shift_formula_coordinates("=A5", -5) == "=#REF!"
        assert shift_formula_coordinates("=A5", -4) == "=A1"
    def test_table_data_range_end_shifted_and_data_start_preserved(self):
        """Table data range end (original_footer - 1) shifts while data start stays frozen."""
        formula = "=SUM(D6:D15)"
        result = shift_formula_coordinates(
            formula,
            row_offset=40,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result == "=SUM(D6:D55)"

    def test_table_data_range_shrink_with_negative_offset(self):
        """Table data range shrinks when row_offset is negative."""
        formula = "=SUM(D6:D15)"
        result = shift_formula_coordinates(
            formula,
            row_offset=-8,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result == "=SUM(D6:D7)"

    def test_data_start_standalone_not_shifted(self):
        """Standalone reference to data start row (< orig_footer - 1) stays frozen."""
        formula = "=D6 + E16"
        result = shift_formula_coordinates(
            formula,
            row_offset=20,
            header_end_row=5,
            original_footer_start_row=16
        )
        assert result == "=D6 + E36"


class TestTemplateLayoutSerialization:
    """Unit tests for TemplateLayout serialization and deserialization."""

    def test_original_footer_start_row_roundtrip(self):
        """original_footer_start_row round-trips correctly through to_dict() and from_dict()."""
        layout = TemplateLayout(original_footer_start_row=15)
        layout_dict = layout.to_dict()

        assert "original_footer_start_row" in layout_dict
        assert layout_dict["original_footer_start_row"] == 15

        restored_layout = TemplateLayout.from_dict(layout_dict)
        assert restored_layout is not None
        assert restored_layout.original_footer_start_row == 15

    def test_original_footer_start_row_none_roundtrip(self):
        """When original_footer_start_row is None, it defaults and round-trips cleanly."""
        layout = TemplateLayout()
        assert layout.original_footer_start_row is None

        layout_dict = layout.to_dict()
        assert layout_dict.get("original_footer_start_row") is None

        restored_layout = TemplateLayout.from_dict(layout_dict)
        assert restored_layout is not None
        assert restored_layout.original_footer_start_row is None

    def test_template_layout_roundtrip_with_rows_and_footer_start(self):
        """Full TemplateLayout with header and footer rows preserves original_footer_start_row."""
        header_row = UnitRow(
            relative_index=0,
            cells=[UnitCell(col_index=1, value="Company Header")],
            height=20.0
        )
        footer_row = UnitRow(
            relative_index=0,
            cells=[UnitCell(col_index=4, value="=SUM(D16:D20)")],
            height=22.5
        )
        layout = TemplateLayout(
            header_rows=[header_row],
            footer_rows=[footer_row],
            col_widths={"A": 15.0, "D": 20.0},
            original_footer_start_row=16
        )

        layout_dict = layout.to_dict()
        restored = TemplateLayout.from_dict(layout_dict)

        assert restored is not None
        assert restored.original_footer_start_row == 16
        assert len(restored.header_rows) == 1
        assert restored.header_rows[0].cells[0].value == "Company Header"
        assert len(restored.footer_rows) == 1
        assert restored.footer_rows[0].cells[0].value == "=SUM(D16:D20)"
        assert restored.col_widths == {"A": 15.0, "D": 20.0}


class TestRestoreFooterRowOffset:
    """Tests for restore_footer row offset calculation and application."""

    def test_restore_footer_row_offset_calculation(self):
        """When orig_start=15 and footer_start_row=35, verifies row_offset=20."""
        layout_dict = {
            "header_rows": [
                {
                    "relative_index": 0,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Header Row 1"}]
                },
                {
                    "relative_index": 4,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Header Row 5"}]
                }
            ],
            "footer_rows": [
                {
                    "relative_index": 0,
                    "height": 25.0,
                    "cells": [
                        {"col_index": 1, "value": "Total"},
                        {"col_index": 5, "value": "=D4 + E15"}
                    ]
                }
            ],
            "original_footer_start_row": 15
        }

        builder = JsonTemplateStateBuilder(layout_dict)
        wb = openpyxl.Workbook()
        ws = wb.active

        with patch.object(builder, 'restore_rows', wraps=builder.restore_rows) as spy_restore_rows:
            builder.restore_footer(ws, footer_start_row=35)

            assert spy_restore_rows.called
            call_kwargs = spy_restore_rows.call_args.kwargs
            assert call_kwargs.get("row_offset") == 20
            assert call_kwargs.get("original_footer_start") == 15
            assert call_kwargs.get("header_end_row") == 5

    def test_restore_footer_applies_offset_to_worksheet_formulas(self):
        """Integration test: restore_footer writes shifted formulas to the worksheet cells."""
        layout_dict = {
            "header_rows": [
                {
                    "relative_index": 0,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Title"}]
                },
                {
                    "relative_index": 4,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Subheader"}]
                }
            ],
            "footer_rows": [
                {
                    "relative_index": 0,
                    "height": 25.0,
                    "cells": [
                        {"col_index": 1, "value": "Total Summary"},
                        {"col_index": 4, "value": "=$D$4 + $E$15"},
                        {"col_index": 5, "value": "=SUM(D16:D20)"}
                    ]
                }
            ],
            "original_footer_start_row": 15
        }

        builder = JsonTemplateStateBuilder(layout_dict)
        wb = openpyxl.Workbook()
        ws = wb.active

        # Restore footer at row 35 (orig_start=15 -> offset=20)
        builder.restore_footer(ws, footer_start_row=35)

        # Row 35 col 1: label unchanged
        assert ws.cell(row=35, column=1).value == "Total Summary"
        # Row 35 col 4: header ref D4 unchanged, footer ref E15 -> E35
        assert ws.cell(row=35, column=4).value == "=$D$4 + $E$35"
        # Row 35 col 5: footer range D16:D20 -> D36:D40 (shifted by 20)
        assert ws.cell(row=35, column=5).value == "=SUM(D36:D40)"

    def test_restore_footer_zero_offset_when_no_orig_start(self):
        """When original_footer_start_row is missing or None, row_offset is 0 and formulas remain unchanged."""
        layout_dict = {
            "header_rows": [
                {
                    "relative_index": 0,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Header"}]
                }
            ],
            "footer_rows": [
                {
                    "relative_index": 0,
                    "height": 25.0,
                    "cells": [
                        {"col_index": 1, "value": "=SUM(D16:D20)"}
                    ]
                }
            ]
            # No original_footer_start_row
        }

        builder = JsonTemplateStateBuilder(layout_dict)
        wb = openpyxl.Workbook()
        ws = wb.active

        with patch.object(builder, 'restore_rows', wraps=builder.restore_rows) as spy_restore_rows:
            builder.restore_footer(ws, footer_start_row=35)

            assert spy_restore_rows.called
            call_kwargs = spy_restore_rows.call_args.kwargs
            assert call_kwargs.get("row_offset") == 0

        # Formula remains unchanged because offset is 0
        assert ws.cell(row=35, column=1).value == "=SUM(D16:D20)"

    def test_restore_footer_zero_offset_when_positions_match(self):
        """When footer_start_row equals orig_start (e.g. 15), row_offset is 0."""
        layout_dict = {
            "header_rows": [
                {
                    "relative_index": 0,
                    "height": 20.0,
                    "cells": [{"col_index": 1, "value": "Header"}]
                }
            ],
            "footer_rows": [
                {
                    "relative_index": 0,
                    "height": 25.0,
                    "cells": [
                        {"col_index": 1, "value": "=SUM(D16:D20)"}
                    ]
                }
            ],
            "original_footer_start_row": 15
        }

        builder = JsonTemplateStateBuilder(layout_dict)
        wb = openpyxl.Workbook()
        ws = wb.active

        with patch.object(builder, 'restore_rows', wraps=builder.restore_rows) as spy_restore_rows:
            builder.restore_footer(ws, footer_start_row=15)

            assert spy_restore_rows.called
            assert spy_restore_rows.call_args.kwargs.get("row_offset") == 0

        assert ws.cell(row=15, column=1).value == "=SUM(D16:D20)"


class TestTemplateScannerFooterEndRow:
    """Tests for TemplateScanner using footer_end_row for original_footer_start_row and footer rows."""

    def test_scanner_uses_footer_end_row_when_present(self):
        """TemplateScanner.scan_static_content should use footer_end_row + 1 if available."""
        from core.blueprint_generator.internal.scanner.template_scanner import TemplateScanner
        from core.blueprint_generator.internal.scanner.models import ZoneBoundaries

        scanner = TemplateScanner()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=26, column=1, value="Footer Content")

        boundaries = ZoneBoundaries(
            header_row=5,
            data_start_row=6,
            footer_row=20,
            footer_end_row=25,
            max_col=10
        )
        layout = scanner.scan_static_content(ws, boundaries, [], "Sheet1")
        assert layout.original_footer_start_row == 26
        assert len(layout.footer_rows) == 1
        assert layout.footer_rows[0].relative_index == 0

    def test_capture_footer_rows_directly_relative_index_zero(self):
        """_capture_footer_rows directly without mocking verifies rel_r for the first footer row is 0."""
        from core.blueprint_generator.internal.scanner.template_scanner import TemplateScanner
        from core.blueprint_generator.internal.scanner.models import ZoneBoundaries

        scanner = TemplateScanner()
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=26, column=1, value="Footer Row 1 (Note)")
        ws.cell(row=27, column=1, value="Footer Row 2 (Signature)")

        boundaries = ZoneBoundaries(
            header_row=5,
            data_start_row=6,
            footer_row=20,
            footer_end_row=25,
            max_col=10
        )

        footer_rows = scanner._capture_footer_rows(ws, boundaries, safe_max_column=10, sheet_name="Sheet1")
        assert len(footer_rows) == 2
        assert footer_rows[0].relative_index == 0
        assert footer_rows[0].cells[0].value == "Footer Row 1 (Note)"
        assert footer_rows[1].relative_index == 1
        assert footer_rows[1].cells[0].value == "Footer Row 2 (Signature)"


