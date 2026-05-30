# tests/test_database_blueprints.py
import os
import json
import pytest
from pathlib import Path
from core.database.db_manager import init_db, SessionLocal, Blueprint, BlueprintTemplate
from core.invoice_generator.resolvers import InvoiceAssetResolver
from core.system_config import sys_config

@pytest.fixture(scope="module")
def setup_db():
    init_db()
    yield

def test_database_blueprint_resolution(setup_db):
    db = SessionLocal()
    
    # 1. Clean up any existing test records
    db.query(Blueprint).filter(Blueprint.customer_code == "DBTEST").delete()
    db.commit()
    
    # 2. Create mock configuration JSONs
    mock_config = {
        "_meta": {
            "config_version": "2.2_strict_mode",
            "customer": "DBTEST_KH",
            "description": "Mock Database Config for Testing"
        },
        "processing": {
            "sheets": ["Invoice"]
        }
    }
    
    mock_template_layout = {
        "fingerprint": {
            "source_file": "DBTEST_raw.xlsx"
        },
        "template_layout": {
            "Invoice": {
                "header_content": {}
            }
        }
    }
    
    # Simple valid minimal Excel file mock bytes
    # (Just some dummy binary bytes to simulate template XLSX)
    mock_xlsx_bytes = b"PK\x03\x04MockExcelTemplateBinaryDataBytes"
    
    # 3. Save to database
    blueprint = Blueprint(
        customer_code="DBTEST",
        locale="KH",
        description="Database Resolution Test",
        config_json=json.dumps(mock_config),
        template_json=json.dumps(mock_template_layout)
    )
    blueprint.template_binary = BlueprintTemplate(
        filename="DBTEST_KH.xlsx",
        xlsx_blob=mock_xlsx_bytes
    )
    db.add(blueprint)
    db.commit()
    
    try:
        # 4. Use InvoiceAssetResolver with default bundled directory to trigger database lookup
        resolver = InvoiceAssetResolver(
            base_config_dir=sys_config.bundled_dir,
            base_template_dir=sys_config.bundled_dir
        )
        
        # 5. Resolve assets
        assets = resolver.resolve_assets_for_input_file("DBTEST25058.json")
        
        # 6. Verify assertions
        assert assets is not None, "Failed to resolve assets from database."
        assert assets.config_data is not None, "config_data should be loaded directly from DB."
        assert assets.template_xlsx_bytes is not None, "template_xlsx_bytes should be loaded directly from DB."
        
        # Verify config content
        assert assets.config_data["_meta"]["customer"] == "DBTEST_KH"
        assert assets.config_data["_meta"]["description"] == "Mock Database Config for Testing"
            
        # Verify template binary content
        assert assets.template_xlsx_bytes == mock_xlsx_bytes
            
        # Verify variant resolution works
        variants = resolver.resolve_all_variants("DBTEST25058.json")
        assert len(variants) == 1
        assert variants[0]["suffix"] == "_KH"
        assert variants[0]["config_data"] is not None
        assert variants[0]["template_xlsx_bytes"] == mock_xlsx_bytes
        
    finally:
        # 7. Clean up DB record
        db.query(Blueprint).filter(Blueprint.customer_code == "DBTEST").delete()
        db.commit()
        db.close()
        
        # Clean up temp files
        temp_dir = sys_config.temp_uploads_dir / "runtime_blueprints" / "DBTEST_KH"
        if temp_dir.exists():
            import shutil
            shutil.rmtree(temp_dir)

def test_in_memory_blueprint_generation(setup_db):
    import openpyxl
    from core.orchestrator import Orchestrator
    
    # 1. Create a dummy workbook template
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Invoice"
    ws["A1"] = "Customer: Test Client"
    ws["A2"] = "Address: 123 Street"
    ws["A3"] = "P.O. No."
    ws["B3"] = "Description"
    ws["C3"] = "Quantity"
    ws["D3"] = "Unit Price"
    ws["E3"] = "Amount"
    
    ws["A4"] = "PO-001"
    ws["B4"] = "Product A"
    ws["C4"] = 10
    ws["D4"] = 5.0
    ws["E4"] = "=C4*D4"
    
    ws["A5"] = "Total"
    ws["E5"] = "=SUM(E4)"
    
    ws["A6"] = "HS.CODE"
    ws["B6"] = "1234.56.78"
    
    temp_excel_path = sys_config.temp_uploads_dir / "test_gen_template.xlsx"
    sys_config.temp_uploads_dir.mkdir(parents=True, exist_ok=True)
    wb.save(temp_excel_path)
    
    try:
        # 2. Run Orchestrator.generate_blueprint_bundle with in_memory=True
        orchestrator = Orchestrator()
        config_data, template_json_data, template_xlsx_bytes = orchestrator.generate_blueprint_bundle(
            template_path=temp_excel_path,
            custom_prefix="MEMTEST",
            runtime_mappings={
                "P.O. No.": "col_po",
                "Description": "col_desc",
                "Quantity": "col_qty_sf",
                "Unit Price": "col_unit_price",
                "Amount": "col_amount"
            },
            in_memory=True,
            ignore_missing_description=True
        )
        
        # 3. Assert outputs are dicts/bytes and not saved to disk
        assert isinstance(config_data, dict)
        assert isinstance(template_json_data, dict)
        assert isinstance(template_xlsx_bytes, bytes)
        
        # Verify it didn't write files to persistent bundled_dir
        config_file = sys_config.bundled_dir / "MEMTEST" / "MEMTEST_KH_config.json"
        assert not config_file.exists()
        
        # Verify content has expected data
        assert config_data["_meta"]["customer"] == "MEMTEST"
        assert "Invoice" in template_json_data["template_layout"]
        
    finally:
        if temp_excel_path.exists():
            temp_excel_path.unlink()
