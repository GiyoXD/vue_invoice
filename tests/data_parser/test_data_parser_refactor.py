import unittest
from decimal import Decimal
from core.data_parser import data_processor
from core.data_parser import sheet_parser
from core.data_parser.validation import validate_data, DataValidationError, validate_col_level_discount, validate_no_duplicate_amount_columns

class TestDataParserRefactor(unittest.TestCase):

    def setUp(self):
        self.mock_new_extract = [
            [
                {
                    "col_po": "A1", "col_item": "I1", "col_qty_sf": Decimal('100.5'),
                    "col_cbm": "1x2x3", "col_unit_price": Decimal('1.5'), "col_amount": Decimal('150.75')
                }
            ]
        ]
        
        self.mock_new_distribute = [
            [
                {"col_po": "A1", "col_item": "I1", "col_qty_sf": Decimal('50'), "col_amount": Decimal('100')},
                {"col_po": "A1", "col_item": "I2", "col_qty_sf": Decimal('30'), "col_amount": None},
                {"col_po": "A1", "col_item": "I3", "col_qty_sf": Decimal('20'), "col_amount": None}
            ]
        ]

    def test_cbm_column(self):
        processed = data_processor.process_cbm_column(list(self.mock_new_extract[0]))
        self.assertEqual(processed[0]['col_cbm'], Decimal('6.0000'))

    def test_distribute_values(self):
        processed = data_processor.distribute_values(
            list(self.mock_new_distribute[0]), 
            columns_to_distribute=['col_amount'], 
            basis_column='col_qty_sf'
        )
        self.assertEqual(processed[0]['col_amount'], Decimal('50.0000'))
        self.assertEqual(processed[1]['col_amount'], Decimal('30.0000'))
        self.assertEqual(processed[2]['col_amount'], Decimal('20.0000'))

    def test_standard_aggregation(self):
        global_map = {}
        data = list(self.mock_new_extract[0])
        data.append(data[0].copy())
        
        res = data_processor.aggregate_standard_by_po_item_price(data, global_map)
        expected_key = ("A1", "I1", Decimal('1.5'), None)
        self.assertIn(expected_key, res)
        self.assertEqual(res[expected_key]['col_qty_sf'], Decimal('201.0'))
        self.assertEqual(res[expected_key]['col_amount'], Decimal('301.50'))

    def test_cbm_pcs_proportion_orphaned_row(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('0'), "_row_num": 12},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('0'), "col_cbm": Decimal('1.5'), "_row_num": 13}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}
        
        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='cbm_proportion')
        self.assertIn("Row 12: Zero CBM (100pcs)", str(context.exception))

        # With ignore_cbm_warning=True, it should pass
        validate_data(data, "Table 1", column_mapping, phase='cbm_proportion', ignore_cbm_warning=True)

    def test_cbm_pcs_proportion_monotonicity_violation(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('1.0'), "_row_num": 14},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('50'), "col_cbm": Decimal('2.0'), "_row_num": 15}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}

        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='cbm_proportion')
        self.assertIn("Row 14 & 15: Qty/CBM wrong (100pcs=1.00, 50pcs=2.00, discrepancy 1.00 CBM)", str(context.exception))

        # With ignore_cbm_warning=True, it should pass
        validate_data(data, "Table 1", column_mapping, phase='cbm_proportion', ignore_cbm_warning=True)

    def test_cbm_pcs_proportion_monotonicity_tolerance(self):
        data_valid = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('1.95'), "_row_num": 14},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('50'), "col_cbm": Decimal('2.0'), "_row_num": 15}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}
        validate_data(data_valid, "Table 1", column_mapping, phase='cbm_proportion')

        data_invalid = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('1.93'), "_row_num": 14},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('50'), "col_cbm": Decimal('2.0'), "_row_num": 15}
        ]
        with self.assertRaises(DataValidationError) as context:
            validate_data(data_invalid, "Table 1", column_mapping, phase='cbm_proportion')
        self.assertIn("Qty/CBM wrong", str(context.exception))

    def test_cbm_pcs_proportion_abnormally_high_ratio(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('2'), "col_cbm": Decimal('1.5'), "_row_num": 16}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}

        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='cbm_proportion')
        self.assertIn("Row 16: High ratio (0.75 > 0.5)", str(context.exception))

        validate_data(data, "Table 1", column_mapping, phase='cbm_proportion', ignore_cbm_warning=True)

    def test_cbm_pcs_proportion_ratio_outliers(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('10.0'), "_row_num": 17},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('10.0'), "_row_num": 18},
            {"col_po": "PO1", "col_item": "ITEM3", "col_qty_pcs": Decimal('100'), "col_cbm": Decimal('10.0'), "_row_num": 19},
            {"col_po": "PO1", "col_item": "ITEM4", "col_qty_pcs": Decimal('10'), "col_cbm": Decimal('2.0'), "_row_num": 20},
            {"col_po": "PO1", "col_item": "ITEM5", "col_qty_pcs": Decimal('10'), "col_cbm": Decimal('0.05'), "_row_num": 21}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}

        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='cbm_proportion')
        self.assertIn("Row 21: Outlier (0.00 vs median 0.10)", str(context.exception))

        validate_data(data, "Table 1", column_mapping, phase='cbm_proportion', ignore_cbm_warning=True)

    def test_cbm_pcs_proportion_valid_table(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_qty_pcs": Decimal('10'), "col_cbm": Decimal('1.0')},
            {"col_po": "PO1", "col_item": "ITEM2", "col_qty_pcs": Decimal('20'), "col_cbm": Decimal('2.0')},
            {"col_po": "PO1", "col_item": "ITEM3", "col_qty_pcs": Decimal('30'), "col_cbm": Decimal('3.0')}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_cbm": "D"}
        validate_data(data, "Table 1", column_mapping, phase='cbm_proportion')

    def test_weight_integrity_unpaired_empty_string(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_net": Decimal('10.0'), "col_gross": "", "_row_num": 5}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_net": "C", "col_gross": "D"}
        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='integrity')
        self.assertIn("Row 5: Missing Gross Weight (Net = 10.00)", str(context.exception))

    def test_weight_integrity_both_empty(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_net": "", "col_gross": None, "_row_num": 5}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_net": "C", "col_gross": "D"}
        validate_data(data, "Table 1", column_mapping, phase='integrity')

    def test_weight_integrity_invalid_positivity(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_net": Decimal('10.0'), "col_gross": Decimal('9.0'), "_row_num": 6}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_net": "C", "col_gross": "D"}
        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='integrity')
        self.assertIn("Row 6: Gross <= Net (Net = 10.00, Gross = 9.00)", str(context.exception))

    def test_validate_data_runs_pallet_integrity(self):
        data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_net": Decimal('10.0'), "col_gross": Decimal('11.0'), "col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_po": "PO1", "col_item": "ITEM2", "col_net": Decimal('20.0'), "col_gross": Decimal('21.0'), "col_pallet_count": 1, "col_pallet_id": "01T26052608", "_row_num": 11},
            {"col_po": "PO1", "col_item": "ITEM3", "col_net": Decimal('30.0'), "col_gross": Decimal('31.0'), "col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 12},
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_net": "C", "col_gross": "D", "col_pallet_count": "E", "col_pallet_id": "F"}
        with self.assertRaises(DataValidationError) as context:
            validate_data(data, "Table 1", column_mapping, phase='integrity')
        self.assertIn("Pallet ID '01T26052605' reappeared after gap", str(context.exception))

    def test_verify_pallet_integrity_valid(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_pallet_count": 1, "col_pallet_id": "01T26052608", "_row_num": 11},
            {"col_pallet_count": 1, "col_pallet_id": "01T26052609", "_row_num": 12},
            {"col_pallet_count": 0, "col_pallet_id": "01T26052609", "_row_num": 13},
            {"col_pallet_count": 1, "col_pallet_id": "02T26052306", "_row_num": 14},
            {"col_pallet_count": 1, "col_pallet_id": "02T26052307", "_row_num": 15},
            {"col_pallet_count": 0, "col_pallet_id": "02T26052307", "_row_num": 16},
        ]
        data_processor.verify_pallet_integrity(data)

    def test_verify_pallet_integrity_invalid_transition(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_pallet_count": 0, "col_pallet_id": "01T26052608", "_row_num": 11},
        ]
        with self.assertRaises(DataValidationError) as context:
            data_processor.verify_pallet_integrity(data)
        self.assertIn("Row 11: Pallet ID changed to '01T26052608' but count is 0", str(context.exception))

    def test_verify_pallet_integrity_invalid_continuity(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052609", "_row_num": 12},
            {"col_pallet_count": 1, "col_pallet_id": "01T26052609", "_row_num": 13},
        ]
        with self.assertRaises(DataValidationError) as context:
            data_processor.verify_pallet_integrity(data)
        self.assertIn("Row 13: Pallet count is 1 but Pallet ID did not change ('01T26052609')", str(context.exception))

    def test_verify_pallet_integrity_missing_id(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "", "_row_num": 10},
        ]
        with self.assertRaises(DataValidationError) as context:
            data_processor.verify_pallet_integrity(data)
        self.assertIn("Row 10: Pallet ID missing (count = 1)", str(context.exception))

    def test_verify_pallet_integrity_recurrence(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_pallet_count": 1, "col_pallet_id": "01T26052608", "_row_num": 11},
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 12},
        ]
        with self.assertRaises(DataValidationError) as context:
            data_processor.verify_pallet_integrity(data)
        self.assertIn("Row 12: Pallet ID '01T26052605' reappeared after gap", str(context.exception))

    def test_verify_pallet_integrity_empty_rows(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_pallet_count": 0, "col_pallet_id": "", "_row_num": 11}, # completely empty
            {"col_pallet_count": 1, "col_pallet_id": "01T26052608", "_row_num": 12},
        ]
        data_processor.verify_pallet_integrity(data)

    def test_verify_pallet_integrity_spacer_row_continuation(self):
        data = [
            {"col_pallet_count": 1, "col_pallet_id": "01T26052605", "_row_num": 10},
            {"col_pallet_count": 0, "col_pallet_id": "", "_row_num": 11}, # spacer row
            {"col_pallet_count": 0, "col_pallet_id": "01T26052605", "_row_num": 12}, # continuation after spacer
        ]
        data_processor.verify_pallet_integrity(data)

    def test_validate_no_duplicate_amount_columns(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="P.O. No.")
        ws.cell(row=1, column=2, value="Amount")
        ws.cell(row=1, column=3, value="Quantity")
        
        validate_no_duplicate_amount_columns(ws, 1)
            
        ws.cell(row=1, column=4, value="Total value(USD)")
        
        with self.assertRaises(DataValidationError) as context:
            validate_no_duplicate_amount_columns(ws, 1)
            
        self.assertIn("Duplicate 'col_amount' columns detected", str(context.exception))

    def test_sheet_parser_detects_duplicate_col_amount(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="P.O. No.")
        ws.cell(row=1, column=2, value="Quantity")
        ws.cell(row=1, column=3, value="Amount")
        ws.cell(row=1, column=4, value="Total value(USD)")
        
        ws.cell(row=2, column=1, value="PO-100")
        ws.cell(row=2, column=2, value=10)
        ws.cell(row=2, column=3, value=100.0)
        ws.cell(row=2, column=4, value=100.0)
        
        with self.assertRaises(DataValidationError) as context:
            sheet_parser.find_and_map_smart_headers(ws)
            
        err_msg = str(context.exception)
        self.assertTrue(
            "Duplicate 'col_amount' columns detected" in err_msg or 
            "Duplicate mapping detected" in err_msg
        )

    def test_validate_no_duplicate_amount_columns_pattern_matching(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="Amount")
        ws.cell(row=2, column=1, value=123.45)
        
        ws.cell(row=1, column=2, value="Price")
        ws.cell(row=2, column=2, value=1.5)

        ws.cell(row=1, column=3, value="")
        ws.cell(row=2, column=3, value=456.78)
        
        with self.assertRaises(DataValidationError) as context:
            validate_no_duplicate_amount_columns(ws, 1)
            
        self.assertIn("Duplicate 'col_amount' columns detected", str(context.exception))

    def test_validate_no_duplicate_amount_columns_adjacent_heuristic(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="Amount")
        ws.cell(row=2, column=1, value=100.00)
        
        ws.cell(row=1, column=2, value="Price")
        ws.cell(row=2, column=2, value=1.5)
        
        ws.cell(row=1, column=3, value="")
        ws.cell(row=2, column=3, value=150.00)
        
        with self.assertRaises(DataValidationError) as context:
            validate_no_duplicate_amount_columns(ws, 1)
            
        self.assertIn("Duplicate 'col_amount' columns detected", str(context.exception))

    def test_validate_no_duplicate_amount_columns_unrecognized_header_with_round_floats(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="Amount")
        ws.cell(row=2, column=1, value=123.00)
        
        ws.cell(row=1, column=2, value="Price")
        ws.cell(row=2, column=2, value=1.5)

        ws.cell(row=1, column=3, value="")
        ws.cell(row=2, column=3, value=456.00)
        
        with self.assertRaises(DataValidationError) as context:
            validate_no_duplicate_amount_columns(ws, 1)
            
        self.assertIn("Duplicate 'col_amount' columns detected", str(context.exception))

    def test_validate_no_duplicate_amount_columns_ignores_mapped_columns(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="Amount")
        ws.cell(row=2, column=1, value=100.00)
        
        ws.cell(row=1, column=2, value="Net Weight")
        ws.cell(row=2, column=2, value=10.50)
        
        mapping = {
            'col_amount': 'A',
            'col_net': 'B'
        }
        
        validate_no_duplicate_amount_columns(ws, 1, column_mapping=mapping)

    def test_greedy_selection_mapping_precedence(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=1, column=1, value="P.O. No.")
        ws.cell(row=2, column=1, value="PO-100")
        
        ws.cell(row=1, column=2, value="Net Weight")
        ws.cell(row=2, column=2, value="ABC")
        
        ws.cell(row=1, column=3, value="Net Weight")
        ws.cell(row=2, column=3, value=10.5)
        
        mapping, score, header_text_matches = sheet_parser._process_row(ws, 1)
        
        self.assertEqual(mapping.get("col_net"), "C")

    def test_validate_no_duplicate_amount_columns_ignores_unrecognized_named_column_when_alias_exists(self):
        from openpyxl import Workbook
        
        wb = Workbook()
        ws = wb.active
        
        ws.cell(row=3, column=1, value="金额")
        ws.cell(row=4, column=1, value=1234.56)
        
        ws.cell(row=3, column=2, value="Price")
        ws.cell(row=4, column=2, value=1.15)
        
        ws.cell(row=3, column=3, value="报关金额")
        ws.cell(row=4, column=3, value=7890.12)
        
        validate_no_duplicate_amount_columns(ws, 3)

    def test_col_dc_aggregation(self):
        data = [
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_unit_price": Decimal("10.0"),
                "col_desc": "Desc A",
                "col_qty_sf": Decimal("50.0"),
                "col_amount": Decimal("500.0"),
                "col_net": Decimal("100.0"),
                "col_cbm": Decimal("1.5"),
                "col_dc": "DC-NORTH",
                "col_pallet_count": 1,
                "col_qty_pcs": 100,
                "col_gross": Decimal("110.0")
            },
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_unit_price": Decimal("10.0"),
                "col_desc": "Desc A",
                "col_qty_sf": Decimal("30.0"),
                "col_amount": Decimal("300.0"),
                "col_net": Decimal("60.0"),
                "col_cbm": Decimal("0.9"),
                "col_dc": "",
                "col_pallet_count": 1,
                "col_qty_pcs": 60,
                "col_gross": Decimal("66.0")
            }
        ]

        std_map = {}
        data_processor.aggregate_standard_by_po_item_price(data, std_map)
        key = ("PO1", "ITEM1", Decimal("10.0"), "Desc A")
        self.assertIn(key, std_map)
        self.assertEqual(std_map[key].get("col_dc"), "DC-NORTH")

        cust_map = {}
        data_processor.aggregate_custom_by_po_item(data, cust_map)
        cust_key = ("PO1", "ITEM1", None, "Desc A")
        self.assertIn(cust_key, cust_map)
        self.assertEqual(cust_map[cust_key].get("col_dc"), "DC-NORTH")

        pallet_agg = data_processor.aggregate_per_po_with_pallets(data)
        self.assertEqual(len(pallet_agg), 1)
        self.assertEqual(pallet_agg[0].get("col_dc"), "DC-NORTH")

        daf_res = data_processor.perform_DAF_compounding(data)
        self.assertEqual(len(daf_res), 1)
        self.assertEqual(daf_res[0].get("col_dc"), "DC-NORTH")
        self.assertEqual(daf_res[0].get("material_type"), "normal")

    def test_perform_daf_compounding_mixed_materials(self):
        data = [
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_desc": "BUFFALO LEATHER",
                "col_qty_sf": Decimal("50.0"),
                "col_amount": Decimal("500.0"),
                "col_net": Decimal("100.0"),
                "col_dc": "DC-NORTH"
            },
            {
                "col_po": "PO2",
                "col_item": "ITEM2",
                "col_desc": "COW LEATHER",
                "col_qty_sf": Decimal("30.0"),
                "col_amount": Decimal("300.0"),
                "col_net": Decimal("60.0"),
                "col_dc": "DC-SOUTH"
            }
        ]
        daf_res = data_processor.perform_DAF_compounding(data)
        self.assertEqual(len(daf_res), 2)
        self.assertEqual(daf_res[0]["material_type"], "buffalo")
        self.assertEqual(daf_res[1]["material_type"], "normal")

    def test_perform_daf_compounding_empty_rows(self):
        daf_res = data_processor.perform_DAF_compounding([])
        self.assertEqual(len(daf_res), 1)
        self.assertEqual(daf_res[0]["col_qty_sf"], Decimal(0))

    def test_validate_col_level_discount(self):
        from unittest.mock import MagicMock

        table_data = [
            {"col_po": "PO1", "col_level": "A", "_row_num": 1},
            {"col_po": "PO2", "col_level": "折扣", "_row_num": 2},
            {"col_po": "PO3", "col_level": "B折扣", "_row_num": 3},
        ]
        monitor = MagicMock()
        validate_col_level_discount(table_data, "Table 1", monitor=monitor)

        monitor.log_warning.assert_called_once()
        warning_msg = monitor.log_warning.call_args[0][0]
        self.assertIn("Warning: '折扣' (discount) detected in col_level on 2 row(s): row(s) 2, 3.", warning_msg)
        self.assertIn("[Table 1]", warning_msg)

    def test_validate_table_data_presence_po_item_zero(self):
        data = [
            {"col_po": "0", "col_item": "0", "col_qty_pcs": Decimal('10'), "col_net": Decimal('100'), "col_gross": Decimal('110'), "col_cbm": Decimal('1.0'), "_row_num": 1}
        ]
        column_mapping = {"col_po": "A", "col_item": "B", "col_qty_pcs": "C", "col_net": "D", "col_gross": "E", "col_cbm": "F"}
        validate_data(data, "Table 1", column_mapping, phase='presence')

    def test_validate_table_data_presence_all_rows_missing_val(self):
        column_mapping = {
            "col_po": "A", "col_item": "B", "col_qty_pcs": "C", 
            "col_net": "D", "col_gross": "E", "col_cbm": "F",
            "col_qty_sf": "G", "col_unit_price": "H", "col_amount": "I"
        }
        # Row 2 missing sqft and unit price
        data = [
            {"col_po": "PO1", "col_item": "Item1", "col_qty_pcs": Decimal('10'), "col_net": Decimal('100'), "col_gross": Decimal('110'), "col_cbm": Decimal('1.0'), "col_qty_sf": Decimal('50'), "col_unit_price": Decimal('5.5'), "col_amount": Decimal('275'), "_row_num": 1},
            {"col_po": "PO2", "col_item": "Item2", "col_qty_pcs": Decimal('20'), "col_net": Decimal('200'), "col_gross": Decimal('220'), "col_cbm": Decimal('2.0'), "col_qty_sf": "", "col_unit_price": "0", "col_amount": Decimal('550'), "_row_num": 2}
        ]
        with self.assertRaises(DataValidationError) as ctx:
            validate_data(data, "Table 1", column_mapping, phase='presence')
        self.assertIn("Row 2: Missing or zero value for: col_qty_sf, col_unit_price", str(ctx.exception))

    def test_validate_table_data_presence_all_rows_missing_po_item(self):
        column_mapping = {
            "col_po": "A", "col_item": "B", "col_qty_pcs": "C", 
            "col_net": "D", "col_gross": "E", "col_cbm": "F"
        }
        # Row 2 missing item
        data = [
            {"col_po": "PO1", "col_item": "Item1", "col_qty_pcs": Decimal('10'), "col_net": Decimal('100'), "col_gross": Decimal('110'), "col_cbm": Decimal('1.0'), "_row_num": 1},
            {"col_po": "PO2", "col_item": "", "col_qty_pcs": Decimal('20'), "col_net": Decimal('200'), "col_gross": Decimal('220'), "col_cbm": Decimal('2.0'), "_row_num": 2}
        ]
        with self.assertRaises(DataValidationError) as ctx:
            validate_data(data, "Table 1", column_mapping, phase='presence')
        self.assertIn("Row 2: Missing or zero value for: col_item", str(ctx.exception))

    def test_validate_table_data_presence_cbm_dimension_string(self):
        column_mapping = {
            "col_po": "A", "col_item": "B", "col_qty_pcs": "C", 
            "col_net": "D", "col_gross": "E", "col_cbm": "F"
        }
        data = [
            {
                "col_po": "PO1", "col_item": "Item1", "col_qty_pcs": Decimal('10'),
                "col_net": Decimal('100'), "col_gross": Decimal('110'),
                "col_cbm": "2.2*1.8*0.4", "_row_num": 1
            }
        ]
        validate_data(data, "Table 1", column_mapping, phase='presence')

    def test_verify_pallet_metric_alignment_warning(self):
        from unittest.mock import MagicMock
        data = [
            {
                "col_pallet_count": 1,
                "col_pallet_id": "P01",
                "col_net": Decimal('100.0'),
                "col_gross": Decimal('110.0'),
                "col_cbm": Decimal('1.5'),
                "_row_num": 1
            },
            {
                "col_pallet_count": 0,
                "col_pallet_id": "P01",
                "col_net": Decimal('50.0'),
                "col_gross": Decimal('55.0'),
                "col_cbm": Decimal('0.75'),
                "_row_num": 2
            },
            {
                "col_pallet_count": 1,
                "col_pallet_id": "P02",
                "col_net": None,
                "col_gross": None,
                "col_cbm": None,
                "_row_num": 3
            },
            {
                "col_pallet_count": 0,
                "col_pallet_id": "P02",
                "col_net": Decimal('80.0'),
                "col_gross": Decimal('88.0'),
                "col_cbm": "2.0",
                "_row_num": 4
            },
            {
                "col_pallet_count": 1,
                "col_pallet_id": "P03",
                "col_net": Decimal('120.0'),
                "col_gross": Decimal('130.0'),
                "col_cbm": Decimal('2.5'),
                "_row_num": 5
            },
            {
                "col_pallet_count": 0,
                "col_pallet_id": "P03",
                "col_net": None,
                "col_gross": "",
                "col_cbm": 0,
                "_row_num": 6
            }
        ]
        monitor = MagicMock()
        data_processor.verify_pallet_integrity(data, table_id_str="Table 1", monitor=monitor)

        # 3 misplaced metrics on row 2, and 3 misplaced metrics on row 4 = 6 warnings
        self.assertEqual(monitor.log_warning.call_count, 6)

        warning_calls = [call_args[0][0] for call_args in monitor.log_warning.call_args_list]

        self.assertTrue(any("Row 2 (Pallet 'P01'): Net Weight (50.0) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))
        self.assertTrue(any("Row 2 (Pallet 'P01'): Gross Weight (55.0) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))
        self.assertTrue(any("Row 2 (Pallet 'P01'): CBM (0.75) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))

        self.assertTrue(any("Row 4 (Pallet 'P02'): Net Weight (80.0) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))
        self.assertTrue(any("Row 4 (Pallet 'P02'): Gross Weight (88.0) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))
        self.assertTrue(any("Row 4 (Pallet 'P02'): CBM (2.0) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))

        for msg in warning_calls:
            self.assertIn("[Table 1]", msg)
            self.assertIn("is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1).", msg)

    def test_verify_pallet_metric_alignment_without_pallet_id(self):
        from unittest.mock import MagicMock
        data = [
            {
                "col_pallet_count": 1,
                "col_cbm": Decimal('1.5'),
                "_row_num": 1
            },
            {
                "col_pallet_count": 0,
                "col_cbm": Decimal('0.75'),
                "_row_num": 2
            }
        ]
        monitor = MagicMock()
        data_processor.verify_pallet_integrity(data, table_id_str="Table 1", monitor=monitor)

        # 1 skipping pallet ID validation warning + 1 alignment warning = 2 warnings
        self.assertEqual(monitor.log_warning.call_count, 2)
        warning_calls = [call_args[0][0] for call_args in monitor.log_warning.call_args_list]

        self.assertTrue(any("[Pallet Pairing] col_pallet_id column not found in Table 1. Skipping pallet ID validation." in msg for msg in warning_calls))
        self.assertTrue(any("[Table 1] Row 2: CBM (0.75) is on a sub-row (Pallet=0). Should be on main pallet row (Pallet=1)." in msg for msg in warning_calls))


class TestAggregationEntryPoints(unittest.TestCase):

    def setUp(self):
        self.raw_data = [
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_unit_price": Decimal("2.00"),
                "col_qty_sf": Decimal("100"),
                "col_amount": Decimal("200"),
                "col_desc": "BUFFALO LEATHER",
                "col_net": Decimal("50"),
                "col_cbm": Decimal("1.2"),
                "col_dc": "DC-NORTH"
            },
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_unit_price": Decimal("2.00"),
                "col_qty_sf": Decimal("50"),
                "col_amount": Decimal("100"),
                "col_desc": "BUFFALO LEATHER",
                "col_net": Decimal("25"),
                "col_cbm": Decimal("0.6"),
                "col_dc": "DC-NORTH"
            }
        ]

    def test_execute_strategy_standard_and_custom(self):
        # Standard: default list
        res_std = data_processor.execute_strategy('standard', self.raw_data)
        self.assertIsInstance(res_std, list)
        self.assertEqual(len(res_std), 1)
        self.assertEqual(res_std[0]['col_po'], "PO1")
        self.assertEqual(res_std[0]['col_item'], "ITEM1")
        self.assertEqual(res_std[0]['col_qty_sf'], Decimal("150.00"))
        self.assertEqual(res_std[0]['col_amount'], Decimal("300.00"))

        # Standard: return_raw=True
        res_std_raw = data_processor.execute_strategy('standard', self.raw_data, return_raw=True)
        self.assertIsInstance(res_std_raw, dict)
        std_key = ("PO1", "ITEM1", Decimal("2.00"), "BUFFALO LEATHER")
        self.assertIn(std_key, res_std_raw)
        self.assertEqual(res_std_raw[std_key]['col_qty_sf'], Decimal("150.00"))
        self.assertEqual(res_std_raw[std_key]['col_amount'], Decimal("300.00"))

        # Custom: default list
        res_cust = data_processor.execute_strategy('custom', self.raw_data)
        self.assertIsInstance(res_cust, list)
        self.assertEqual(len(res_cust), 1)
        self.assertEqual(res_cust[0]['col_po'], "PO1")
        self.assertEqual(res_cust[0]['col_item'], "ITEM1")
        self.assertEqual(res_cust[0]['col_qty_sf'], Decimal("150.00"))
        self.assertEqual(res_cust[0]['col_amount'], Decimal("300.00"))

        # Custom: return_raw=True
        res_cust_raw = data_processor.execute_strategy('custom', self.raw_data, return_raw=True)
        self.assertIsInstance(res_cust_raw, dict)
        cust_key = ("PO1", "ITEM1", None, "BUFFALO LEATHER")
        self.assertIn(cust_key, res_cust_raw)
        self.assertEqual(res_cust_raw[cust_key]['col_qty_sf'], Decimal("150.00"))
        self.assertEqual(res_cust_raw[cust_key]['col_amount'], Decimal("300.00"))

        # Item merged from raw data
        res_merged_raw = data_processor.execute_strategy('item_merged', self.raw_data)
        self.assertIsInstance(res_merged_raw, list)
        self.assertEqual(len(res_merged_raw), 1)
        self.assertEqual(res_merged_raw[0]['col_item'], "ITEM1")
        self.assertEqual(res_merged_raw[0]['col_po'], "PO1")
        self.assertEqual(res_merged_raw[0]['col_qty_sf'], Decimal("150.00"))

        # Item merged from pre-aggregated custom list
        res_merged_cust = data_processor.execute_strategy('item_merged', res_cust)
        self.assertIsInstance(res_merged_cust, list)
        self.assertEqual(len(res_merged_cust), 1)
        self.assertEqual(res_merged_cust[0]['col_item'], "ITEM1")
        self.assertEqual(res_merged_cust[0]['col_po'], "PO1")
        self.assertEqual(res_merged_cust[0]['col_qty_sf'], Decimal("150.00"))

    def test_execute_strategy_case_insensitivity_and_aliases(self):
        # 'STANDARD' uppercase
        res_std = data_processor.execute_strategy('STANDARD', self.raw_data)
        self.assertIsInstance(res_std, list)
        self.assertEqual(len(res_std), 1)

        # 'custom' lowercase
        res_cust = data_processor.execute_strategy('custom', self.raw_data)
        self.assertIsInstance(res_cust, list)
        self.assertEqual(len(res_cust), 1)

        # 'pallets' and 'manifest_by_pallet_per_po'
        res_pallets = data_processor.execute_strategy('pallets', self.raw_data)
        res_manifest = data_processor.execute_strategy('manifest_by_pallet_per_po', self.raw_data)
        self.assertIsInstance(res_pallets, list)
        self.assertIsInstance(res_manifest, list)
        self.assertEqual(res_pallets, res_manifest)

        # 'daf' (and case insensitive 'DAF')
        res_daf = data_processor.execute_strategy('daf', self.raw_data)
        res_daf_upper = data_processor.execute_strategy('DAF', self.raw_data)
        self.assertIsInstance(res_daf, list)
        self.assertEqual(res_daf, res_daf_upper)

        # 'leather_summary'
        res_leather = data_processor.execute_strategy('leather_summary', self.raw_data)
        self.assertIsInstance(res_leather, dict)
        self.assertIn('BUFFALO', res_leather)
        self.assertIn('COW', res_leather)
        self.assertEqual(res_leather['BUFFALO']['col_qty_sf'], Decimal('150.00'))

    def test_execute_strategy_invalid_name(self):
        with self.assertRaises(ValueError) as ctx:
            data_processor.execute_strategy('non_existent_strategy', self.raw_data)
        self.assertIn("Unknown aggregation strategy", str(ctx.exception))

    def test_run_aggregation_pipeline(self):
        pipeline_data = [
            {
                "col_po": "PO1",
                "col_item": "ITEM1",
                "col_unit_price": Decimal("2.00"),
                "col_qty_sf": Decimal("100"),
                "col_amount": Decimal("200"),
                "col_desc": "BUFFALO LEATHER",
                "col_net": Decimal("50"),
                "col_cbm": Decimal("1.2"),
                "col_dc": "DC1"
            },
            {
                "col_po": "PO2",
                "col_item": "ITEM2",
                "col_unit_price": Decimal("3.00"),
                "col_qty_sf": Decimal("50"),
                "col_amount": Decimal("150"),
                "col_desc": "COW LEATHER",
                "col_net": Decimal("25"),
                "col_cbm": Decimal("0.8"),
                "col_dc": "DC2"
            }
        ]
        res = data_processor.run_aggregation_pipeline(pipeline_data)
        expected_keys = {"standard", "custom", "manifest_by_pallet_per_po", "daf", "item_merged"}
        self.assertEqual(set(res.keys()), expected_keys)
        self.assertEqual(len(res["standard"]), 2)
        self.assertEqual(len(res["custom"]), 2)
        self.assertEqual(len(res["manifest_by_pallet_per_po"]), 2)
        self.assertEqual(len(res["item_merged"]), 2)
        self.assertGreaterEqual(len(res["daf"]), 1)

        # Empty input: []
        res_empty = data_processor.run_aggregation_pipeline([])
        self.assertEqual(set(res_empty.keys()), expected_keys)
        self.assertEqual(res_empty["standard"], [])
        self.assertEqual(res_empty["custom"], [])
        self.assertEqual(res_empty["manifest_by_pallet_per_po"], [])
        self.assertEqual(res_empty["item_merged"], [])
        self.assertEqual(len(res_empty["daf"]), 1)
        self.assertEqual(res_empty["daf"][0]["material_type"], "normal")
        self.assertEqual(res_empty["daf"][0]["col_qty_sf"], Decimal("0.00"))

    def test_daf_explicit_arguments(self):
        daf_data = [
            {"col_po": "PO1", "col_item": "ITEM1", "col_unit_price": Decimal("2.00"), "col_qty_sf": Decimal("10"), "col_amount": Decimal("20"), "col_desc": "BUFFALO 1", "col_net": Decimal("5"), "col_cbm": Decimal("0.1"), "col_dc": "DC1"},
            {"col_po": "PO2", "col_item": "ITEM2", "col_unit_price": Decimal("2.00"), "col_qty_sf": Decimal("10"), "col_amount": Decimal("20"), "col_desc": "BUFFALO 2", "col_net": Decimal("5"), "col_cbm": Decimal("0.1"), "col_dc": "DC2"},
            {"col_po": "PO3", "col_item": "ITEM3", "col_unit_price": Decimal("2.00"), "col_qty_sf": Decimal("10"), "col_amount": Decimal("20"), "col_desc": "BUFFALO 3", "col_net": Decimal("5"), "col_cbm": Decimal("0.1"), "col_dc": "DC3"},
        ]
        decimals_map = {"col_qty_sf": 3, "col_amount": 3, "col_net": 3}

        # 1. run_aggregation_pipeline with explicit arguments
        res1 = data_processor.run_aggregation_pipeline(
            daf_data,
            decimals_map=decimals_map,
            chunk_size=2,
            item_separator=" - ",
            line_separator=" | ",
        )
        self.assertIn(" - ", res1["daf"][0]["col_po"])
        self.assertIn(" | ", res1["daf"][0]["col_po"])
        self.assertEqual(res1["daf"][0]["col_po"], "PO1 - PO2 | PO3")
        self.assertEqual(res1["daf"][0]["col_item"], "ITEM1 - ITEM2 | ITEM3")
        self.assertEqual(res1["daf"][0]["col_qty_sf"], Decimal("30.000"))

        # 2. execute_strategy 'daf' with explicit arguments
        strat_res = data_processor.execute_strategy(
            'daf',
            daf_data,
            decimals_map=decimals_map,
            chunk_size=2,
            item_separator=" ; ",
            line_separator=" // ",
        )
        self.assertEqual(strat_res[0]["col_po"], "PO1 ; PO2 // PO3")
        self.assertEqual(strat_res[0]["col_qty_sf"], Decimal("30.000"))

        # 3. Direct perform_DAF_compounding with explicit arguments
        direct_res = data_processor.perform_DAF_compounding(
            daf_data,
            chunk_size=2,
            item_separator=" / ",
            line_separator="\n",
            decimals_map=decimals_map,
        )
        self.assertEqual(direct_res[0]["col_po"], "PO1 / PO2\nPO3")
        self.assertEqual(direct_res[0]["col_qty_sf"], Decimal("30.000"))

        # 4. Empty string separators
        direct_empty = data_processor.perform_DAF_compounding(
            daf_data,
            chunk_size=2,
            item_separator="",
            line_separator="",
        )
        self.assertEqual(direct_empty[0]["col_po"], "PO1PO2PO3")
        self.assertEqual(direct_empty[0]["col_item"], "ITEM1ITEM2ITEM3")

        # 5. Invalid chunk_size fallback to 2
        direct_zero = data_processor.perform_DAF_compounding(
            daf_data,
            chunk_size=0,
            item_separator="/",
            line_separator="\n",
        )
        self.assertEqual(direct_zero[0]["col_po"], "PO1/PO2\nPO3")

        direct_neg = data_processor.perform_DAF_compounding(
            daf_data,
            chunk_size=-1,
            item_separator="/",
            line_separator="\n",
        )
        self.assertEqual(direct_neg[0]["col_po"], "PO1/PO2\nPO3")


if __name__ == '__main__':
    unittest.main()

