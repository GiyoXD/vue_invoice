"""
Workbook Utilities

Generic helper functions for workbook manipulation during invoice generation.
Extracted from generate_invoice.py to keep the orchestrator lean.

These functions operate on GeneratorContext or raw workbook objects and handle:
- Worksheet deep copying
- Unknown sheet injection
- Print area configuration
- Output filename generation
- Layout column counting
"""

import logging
import re
from pathlib import Path
from typing import Optional
from copy import copy

import openpyxl
from openpyxl.utils import get_column_letter

from .print_area_config import configure_print_area
from ...utils.file_lock import ensure_file_unlocked

logger = logging.getLogger(__name__)


# ---------------------------------------------------------------------------
# Worksheet Copying
# ---------------------------------------------------------------------------

def deep_copy_worksheet(source_ws, target_ws):
    """
    Deep copy a worksheet's content from one workbook to another.

    Copies cell values, styles, merged cells, column dimensions,
    row dimensions, and sheet visibility state.

    Args:
        source_ws: Source worksheet (from bundled .xlsx)
        target_ws: Target worksheet (in output workbook)
    """
    # 1. Copy cell values and styles
    for row in source_ws.iter_rows():
        for cell in row:
            target_cell = target_ws.cell(row=cell.row, column=cell.column, value=cell.value)

            # Copy style attributes
            if cell.has_style:
                target_cell.font = copy(cell.font)
                target_cell.fill = copy(cell.fill)
                target_cell.border = copy(cell.border)
                target_cell.alignment = copy(cell.alignment)
                target_cell.number_format = cell.number_format
                target_cell.protection = copy(cell.protection)

    # 2. Copy merged cell ranges
    for merged_range in source_ws.merged_cells.ranges:
        target_ws.merge_cells(str(merged_range))

    # 3. Copy column dimensions (widths)
    for col_letter, col_dim in source_ws.column_dimensions.items():
        target_ws.column_dimensions[col_letter].width = col_dim.width
        target_ws.column_dimensions[col_letter].hidden = col_dim.hidden

    # 4. Copy row dimensions (heights)
    for row_num, row_dim in source_ws.row_dimensions.items():
        target_ws.row_dimensions[row_num].height = row_dim.height
        target_ws.row_dimensions[row_num].hidden = row_dim.hidden

    # 5. Copy sheet visibility state (visible/hidden/veryHidden)
    target_ws.sheet_state = source_ws.sheet_state

    # 6. Copy print area, page setup, margins, and header/footer metadata
    target_ws.print_area = source_ws.print_area
    target_ws.print_title_rows = source_ws.print_title_rows
    target_ws.print_title_cols = source_ws.print_title_cols

    if source_ws.page_setup:
        target_ws.page_setup.orientation = source_ws.page_setup.orientation
        target_ws.page_setup.paperSize = source_ws.page_setup.paperSize
        target_ws.page_setup.fitToWidth = source_ws.page_setup.fitToWidth
        target_ws.page_setup.fitToHeight = source_ws.page_setup.fitToHeight

    if source_ws.page_margins:
        target_ws.page_margins.left = source_ws.page_margins.left
        target_ws.page_margins.right = source_ws.page_margins.right
        target_ws.page_margins.top = source_ws.page_margins.top
        target_ws.page_margins.bottom = source_ws.page_margins.bottom
        target_ws.page_margins.header = source_ws.page_margins.header
        target_ws.page_margins.footer = source_ws.page_margins.footer

    if source_ws.sheet_properties and source_ws.sheet_properties.pageSetUpPr:
        target_ws.sheet_properties.pageSetUpPr.fitToPage = source_ws.sheet_properties.pageSetUpPr.fitToPage

    if source_ws.HeaderFooter:
        target_ws.HeaderFooter = copy(source_ws.HeaderFooter)



# ---------------------------------------------------------------------------
# Unknown Sheet Injection
# ---------------------------------------------------------------------------

def inject_unknown_sheets(ctx):
    """
    Copy unknown sheets from bundled .xlsx into output workbook.

    Unknown sheets are those present in the original bundled .xlsx template
    but NOT defined in the JSON template config. These are preserved as-is,
    maintaining their original visibility state (visible/hidden/veryHidden),
    cell values, styles, merged cells, and dimensions.
    """
    # OPT-OUT: Skip if config explicitly says no static sheets
    if ctx.config_loader and not ctx.config_loader.has_static_sheets():
        logger.info("[Unknown Sheets] Static sheet injection skipped (has_static_sheets=false)")
        return

    # Check for in-memory xlsx bytes first (DB blueprints)
    from io import BytesIO
    template_xlsx_bytes = getattr(ctx, 'template_xlsx_bytes', None)
    if template_xlsx_bytes:
        try:
            source_wb = openpyxl.load_workbook(BytesIO(template_xlsx_bytes))
        except Exception as e:
            logger.warning(f"[Unknown Sheets] Failed to load in-memory xlsx: {e}")
            return
    else:
        # Derive the bundle directory from the config path
        config_path = Path(ctx.paths.get('config', ''))
        if not config_path.exists():
            logger.debug("No config path found, skipping unknown sheet injection.")
            return

        bundle_dir = config_path.parent

        # Derive prefix from config filename (e.g. "TEST_VN_config.json" -> "TEST_VN")
        config_stem = config_path.stem  # "TEST_VN_config"
        prefix = config_stem.replace("_config", "")  # "TEST_VN"

        # Find matching .xlsx by prefix first, fallback to any .xlsx
        source_xlsx_path = bundle_dir / f"{prefix}.xlsx"
        if not source_xlsx_path.exists():
            xlsx_candidates = list(bundle_dir.glob("*.xlsx"))
            if not xlsx_candidates:
                logger.debug(f"No .xlsx files found in bundle dir: {bundle_dir}")
                return
            source_xlsx_path = xlsx_candidates[0]
        logger.info(f"[Unknown Sheets] Loading source template: {source_xlsx_path.name}")

        try:
            source_wb = openpyxl.load_workbook(source_xlsx_path)
        except Exception as e:
            logger.warning(f"[Unknown Sheets] Failed to load source xlsx: {e}")
            return

    # Determine which sheets are "configured" (already in output)
    configured_sheets = set(ctx.output_workbook.sheetnames)
    unknown_sheets = [s for s in source_wb.sheetnames if s not in configured_sheets]

    if not unknown_sheets:
        logger.info("[Unknown Sheets] No unknown sheets to inject.")
        source_wb.close()
        return

    logger.info(f"[Unknown Sheets] Found {len(unknown_sheets)} unknown sheet(s): {unknown_sheets}")

    for sheet_name in unknown_sheets:
        try:
            source_ws = source_wb[sheet_name]
            target_ws = ctx.output_workbook.create_sheet(sheet_name)
            deep_copy_worksheet(source_ws, target_ws)
            logger.info(f"  ✅ Injected '{sheet_name}' (state={source_ws.sheet_state})")
        except Exception as e:
            logger.warning(f"  ⚠ Failed to inject '{sheet_name}': {e}")

    source_wb.close()


# ---------------------------------------------------------------------------
# Print Area & Layout
# ---------------------------------------------------------------------------

def count_layout_columns(config_loader, sheet_name: str) -> Optional[int]:
    """
    Count the actual number of Excel columns from the layout config's structure.columns.

    Accounts for parent columns with children (e.g. col_qty_header with
    children col_qty_pcs + col_qty_sf = 2 actual columns, not 3).

    Args:
        config_loader: The ConfigStore instance.
        sheet_name: Name of the sheet to count columns for.

    Returns:
        int column count, or None if no layout structure is defined.
    """
    layout = config_loader.get_layout_config(sheet_name)
    columns = layout.get('structure', {}).get('columns', [])
    if not columns:
        return None

    count = 0
    for col in columns:
        children = col.get('children', [])
        if children:
            count += len(children)
        else:
            count += 1

    logger.debug(f"[PrintArea] Layout column count for '{sheet_name}': {count}")
    return count


def apply_print_settings(ctx):
    """Apply print area and page setup to all configured sheets."""
    logger.info("Applying Print Area & Page Setup...")
    configured_sheets = set(ctx.config_loader.get_sheets_to_process())

    for sheet in ctx.output_workbook.sheetnames:
        if sheet not in configured_sheets:
            logger.info(f"Skipping print setup for unconfigured sheet '{sheet}'")
            continue

        try:
            ws = ctx.output_workbook[sheet]
            if ws is None:
                logger.warning(f"Sheet '{sheet}' is in sheetnames but returned None - skipping print setup")
                continue

            max_col = ctx.config_loader.get_max_columns(sheet)
            configure_print_area(ws, max_col_override=max_col)
        except Exception as e:
            logger.error(f"Print setup failed for '{sheet}': {e}")


# ---------------------------------------------------------------------------
# Output Filename
# ---------------------------------------------------------------------------

def get_mode_suffix(ctx) -> str:
    """
    Returns a filename suffix based on the active generation mode.

    Returns:
        ' DAF' if daf_mode is active,
        ' Custom' if custom_mode is active,
        '' for standard mode.
    """
    if ctx.daf_mode:
        return " DAF"
    elif ctx.custom_mode:
        return " Custom"
    return ""


def build_output_filename(ctx):
    """
    Builds a dynamic output filename based on what sheets are present,
    the invoice ID, and the active generation mode.

    Maps sheet names to abbreviations:
        - "Contract" -> "CT"
        - "Invoice"  -> "INV"
        - "Packing list" -> "PL"

    Result examples:
        - Standard: "CT&INV&PL MT2-26007E.xlsx"
        - Custom:   "CT&INV&PL MT2-26007E Custom.xlsx"
        - DAF:      "CT&INV&PL MT2-26007E DAF.xlsx"
    """
    # Sheet name -> abbreviation mapping (order matters for the prefix)
    SHEET_ABBREVS = [
        ("Contract",     "CT"),
        ("Invoice",      "INV"),
        ("Packing list", "PL"),
    ]

    # Build prefix from sheets present in the output workbook
    present_abbrevs = []
    for sheet_name, abbrev in SHEET_ABBREVS:
        if sheet_name in ctx.output_workbook.sheetnames:
            present_abbrevs.append(abbrev)

    if not present_abbrevs:
        logger.warning("[Filename] No recognizable sheets found. Using default filename.")
        return

    prefix = "&".join(present_abbrevs)

    # Extract invoice_id from invoice_data
    inv_no = ""
    if 'invoice_info' in ctx.invoice_data:
        inv_no = ctx.invoice_data['invoice_info'].get('col_inv_no', "") or \
                 ctx.invoice_data['invoice_info'].get('inv_no', "")

    if not inv_no:
        # Fallback: try processed_tables_multi
        tables = ctx.invoice_data.get('processed_tables_multi', {})
        table_1 = tables.get('1', {})
        vals = table_1.get('col_inv_no', [])
        if isinstance(vals, list):
            for v in vals:
                if v:
                    inv_no = str(v)
                    break

    # Get mode suffix (e.g. " DAF", " Custom", or "")
    mode_suffix = get_mode_suffix(ctx)

    # Compose filename
    if inv_no:
        new_name = f"{prefix} {inv_no}{mode_suffix}.xlsx"
    else:
        new_name = f"{prefix}{mode_suffix}.xlsx"

    # Sanitize: remove characters illegal in Windows filenames
    new_name = re.sub(r'[<>:"/\\|?*]', '_', new_name)

    new_output_path = ctx.output_path.parent / new_name
    logger.info(f"[Filename] Dynamic output: {new_name}")
    ctx.output_path = new_output_path


# ---------------------------------------------------------------------------
# Sheet Splitting
# ---------------------------------------------------------------------------

def _build_split_filename(base_stem: str, sheet_name: str, suffix: str) -> str:
    """
    Build a filename for a split sheet.
    Replaces 'Invoice' in the stem with the sheet name to stay tied
    to the original input file identity.

    Examples:
        ("TH26001_Invoice_KH", "Packing list", ".xlsx") -> "TH26001_Packing list_KH.xlsx"
        ("TH26001_Invoice",    "Contract",     ".xlsx") -> "TH26001_Contract.xlsx"
        ("SomeFile",           "Invoice",      ".xlsx") -> "SomeFile_Invoice.xlsx"
    """
    if "Invoice" in base_stem:
        new_base = base_stem.replace("Invoice", sheet_name)
    else:
        new_base = f"{base_stem}_{sheet_name}"
    return f"{new_base}{suffix}"


def split_workbook_to_buffers(workbook, output_path) -> list:
    """
    Split a workbook into per-sheet in-memory byte buffers.

    Skips hidden sheets (openpyxl cannot save a workbook where the only
    sheet is hidden). Preserves all print areas, dimensions, and styles.

    Args:
        workbook:    The openpyxl Workbook to split.
        output_path: A Path used to derive filenames for each split file.

    Returns:
        List of (filename: str, file_bytes: bytes) tuples.
    """
    import io

    logger.info("Splitting workbook into individual sheet buffers...")

    # Snapshot the full workbook once
    master_buffer = io.BytesIO()
    workbook.save(master_buffer)

    results = []
    for sheet_name in workbook.sheetnames:
        if workbook[sheet_name].sheet_state != 'visible':
            logger.info(f"Skipping hidden sheet '{sheet_name}' during split.")
            continue

        master_buffer.seek(0)
        wb = openpyxl.load_workbook(master_buffer)

        # Remove every sheet except the target and metadata sheets (e.g. DeepSheet)
        for sn in wb.sheetnames:
            if sn != sheet_name and sn != "DeepSheet":
                wb.remove(wb[sn])

        sheet_buffer = io.BytesIO()
        wb.save(sheet_buffer)
        wb.close()

        filename = _build_split_filename(
            output_path.stem, sheet_name, output_path.suffix
        )
        results.append((filename, sheet_buffer.getvalue()))
        logger.info(f"  ✅ Split sheet '{sheet_name}' -> {filename}")

    return results


# ---------------------------------------------------------------------------
# Finalization
# ---------------------------------------------------------------------------

def finalize(ctx):
    """Apply print settings, optionally split sheets, and save the workbook to disk."""
    apply_print_settings(ctx)

    # Split sheets to individual files if requested
    if getattr(ctx, 'split_sheets', False):
        for filename, fbytes in split_workbook_to_buffers(ctx.output_workbook, ctx.output_path):
            split_path = ctx.output_path.parent / filename
            try:
                ensure_file_unlocked(split_path)
            except Exception as e:
                logger.error(f"File Lock Error on split file {split_path}: {e}")
                continue
            split_path.write_bytes(fbytes)
            logger.info(f"Saved split file: {split_path}")

    logger.info(f"Saving workbook to {ctx.output_path}")

    # Check for file locks and attempting to kill Excel if needed
    try:
        ensure_file_unlocked(ctx.output_path)
    except Exception as e:
        logger.error(f"File Lock Error: {e}")
        # We raise here because if we can't write, we can't save.
        raise e

    ctx.output_workbook.save(ctx.output_path)

    # Cleanup
    if ctx.output_workbook: ctx.output_workbook.close()
