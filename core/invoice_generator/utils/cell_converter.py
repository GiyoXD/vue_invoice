from typing import List, Dict, Any, Optional
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell import Cell
from core.models.cell import (
    UnitRow,
    UnitCell,
    CellStyle,
    FontStyle,
    AlignmentStyle,
    FillStyle,
    BorderStyle,
    TemplateMerge
)

def extract_color(color_obj: Any) -> Optional[str]:
    """Helper to extract RGB color from openpyxl color object safely."""
    if not color_obj:
        return None
    if isinstance(color_obj, str):
        return color_obj
    if getattr(color_obj, 'type', None) != 'rgb':
        return None
    if hasattr(color_obj, 'rgb') and color_obj.rgb:
        # openpyxl colors are usually ARGB (e.g. 'FFFF0000'). We might just store the raw string.
        return str(color_obj.rgb)
    return None

def convert_registry_style_to_cell_style(style_dict: Dict[str, Any]) -> CellStyle:
    """Helper to convert StyleRegistry merged style dict to CellStyle model."""
    if not style_dict:
        return CellStyle()
    
    font_style = FontStyle(
        name=style_dict.get('font_name'),
        size=style_dict.get('font_size'),
        bold=bool(style_dict.get('bold')),
        italic=bool(style_dict.get('italic')),
        color=style_dict.get('font_color')
    )
    
    # alignment
    align_val = style_dict.get('alignment')
    horizontal = align_val.get('horizontal') if isinstance(align_val, dict) else align_val
    vertical = align_val.get('vertical') if isinstance(align_val, dict) else style_dict.get('vertical_alignment', 'center')
    wrap_text = bool(align_val.get('wrap_text', False) if isinstance(align_val, dict) else style_dict.get('wrap_text', False))
    
    alignment_style = AlignmentStyle(
        horizontal=horizontal,
        vertical=vertical,
        wrap_text=wrap_text
    )
    
    # fill
    fill_style = None
    if style_dict.get('fill_color'):
        fill_color = style_dict['fill_color']
        if fill_color.startswith('#'):
            fill_color = fill_color[1:]
        if len(fill_color) == 6:
            fill_color = 'FF' + fill_color
        fill_style = FillStyle(
            fill_type='solid',
            color=fill_color
        )
        
    # border: NOT handled here — BorderResolver stamps borders directly onto cells
    # after all sections are built. See border_resolver.py.
            
    return CellStyle(
        font=font_style,
        alignment=alignment_style,
        fill=fill_style,
        border=None,
        number_format=style_dict.get('format', 'General')
    )

def convert_worksheet_to_models(ws: Worksheet, start_row: int = 1, end_row: Optional[int] = None) -> List[UnitRow]:
    """
    Converts a populated openpyxl Worksheet into a list of UnitRow models.
    """
    unit_rows = []
    max_row = end_row if end_row is not None else ws.max_row
    
    # Pre-process merged cells to attach to top-left cell
    # Key: (row, col)
    merges_dict: Dict[tuple, TemplateMerge] = {}
    for merged_range in ws.merged_cells.ranges:
        # We only assign the merge to the top-left cell
        min_r, min_c, max_r, max_c = merged_range.min_row, merged_range.min_col, merged_range.max_row, merged_range.max_col
        # openpyxl ranges are 1-indexed
        value = ws.cell(row=min_r, column=min_c).value
        merges_dict[(min_r, min_c)] = TemplateMerge(
            min_col=min_c,
            max_col=max_c,
            row_span=max_r - min_r + 1,
            value=str(value) if value is not None else ""
        )

    for row_idx in range(start_row, max_row + 1):
        # We will use relative index based on start_row to be 0-indexed or 1-indexed?
        # The builders usually treat relative_index from 0 within their own block.
        # But here we might just map directly. Let's use the actual row index for now, 
        # or we can pass an offset. Let's make relative_index = row_idx - start_row
        row_height = ws.row_dimensions[row_idx].height
        
        cells = []
        # get max column for this row specifically, or just overall max_col
        for col_idx in range(1, ws.max_column + 1):
            cell_obj = ws.cell(row=row_idx, column=col_idx)
            
            # Skip empty cells unless they are the top-left of a merge
            is_top_left_merge = (row_idx, col_idx) in merges_dict
            
            # Check if cell has value or styling (we might want to skip completely blank unstyled cells)
            # A completely default cell has no value and 'General' number format and no borders/fills
            has_value = cell_obj.value is not None
            has_style = cell_obj.has_style  # openpyxl property
            
            if not has_value and not has_style and not is_top_left_merge:
                continue

            # Convert style
            cell_style = None
            if has_style:
                # Font
                font_style = None
                if cell_obj.font:
                    font_style = FontStyle(
                        name=cell_obj.font.name,
                        size=cell_obj.font.size,
                        bold=cell_obj.font.b,
                        italic=cell_obj.font.i,
                        color=extract_color(cell_obj.font.color)
                    )
                
                # Alignment
                alignment_style = None
                if cell_obj.alignment:
                    alignment_style = AlignmentStyle(
                        horizontal=cell_obj.alignment.horizontal,
                        vertical=cell_obj.alignment.vertical,
                        wrap_text=cell_obj.alignment.wrap_text
                    )

                # Fill
                fill_style = None
                if cell_obj.fill:
                    # solid fills use fgColor for the actual color
                    color = extract_color(cell_obj.fill.fgColor) if cell_obj.fill.fill_type == "solid" else None
                    fill_style = FillStyle(
                        fill_type=cell_obj.fill.fill_type,
                        color=color
                    )

                # Border
                border_style = None
                if cell_obj.border:
                    border_style = BorderStyle(
                        left=cell_obj.border.left.style if cell_obj.border.left else None,
                        right=cell_obj.border.right.style if cell_obj.border.right else None,
                        top=cell_obj.border.top.style if cell_obj.border.top else None,
                        bottom=cell_obj.border.bottom.style if cell_obj.border.bottom else None
                    )

                # Assemble CellStyle
                if font_style or alignment_style or fill_style or border_style or cell_obj.number_format != 'General':
                    cell_style = CellStyle(
                        font=font_style,
                        alignment=alignment_style,
                        fill=fill_style,
                        border=border_style,
                        number_format=cell_obj.number_format
                    )
            
            unit_cell = UnitCell(
                col_index=col_idx,
                value=cell_obj.value,
                style=cell_style,
                merge=merges_dict.get((row_idx, col_idx))
            )
            cells.append(unit_cell)
            
        # Add the row even if empty, as height might be important or merges span across it
        unit_rows.append(UnitRow(
            relative_index=row_idx - start_row,
            height=row_height,
            cells=cells
        ))

    return unit_rows

_font_cache = {}
_alignment_cache = {}
_fill_cache = {}
_border_cache = {}

def write_models_to_worksheet(ws: Worksheet, rows: List[UnitRow], start_row: int = 1) -> None:
    """
    Writes a list of UnitRow models to an openpyxl Worksheet starting at start_row.
    """
    from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

    for row_model in rows:
        target_row_idx = start_row + row_model.relative_index
        
        if row_model.height is not None:
            ws.row_dimensions[target_row_idx].height = row_model.height
            
        for cell_model in row_model.cells:
            target_col_idx = cell_model.col_index
            cell_obj = ws.cell(row=target_row_idx, column=target_col_idx)
            
            # Set value (MergedCell value is read-only)
            if cell_obj.__class__.__name__ != 'MergedCell':
                cell_obj.value = cell_model.value
            
            # Set styles
            if cell_model.style:
                if cell_model.style.font:
                    font = cell_model.style.font
                    key = (font.name, font.size, font.bold, font.italic, font.color)
                    if key not in _font_cache:
                        kwargs = {}
                        if font.name: kwargs['name'] = font.name
                        if font.size: kwargs['size'] = font.size
                        if font.bold: kwargs['bold'] = font.bold
                        if font.italic: kwargs['italic'] = font.italic
                        if font.color: kwargs['color'] = font.color
                        _font_cache[key] = Font(**kwargs)
                    cell_obj.font = _font_cache[key]
                    
                if cell_model.style.alignment:
                    align = cell_model.style.alignment
                    key = (align.horizontal, align.vertical, align.wrap_text)
                    if key not in _alignment_cache:
                        kwargs = {}
                        if align.horizontal: kwargs['horizontal'] = align.horizontal
                        if align.vertical: kwargs['vertical'] = align.vertical
                        if align.wrap_text: kwargs['wrap_text'] = align.wrap_text
                        _alignment_cache[key] = Alignment(**kwargs)
                    cell_obj.alignment = _alignment_cache[key]
                    
                if cell_model.style.fill:
                    fill = cell_model.style.fill
                    key = (fill.fill_type, fill.color)
                    if key not in _fill_cache:
                        if fill.fill_type == "solid" and fill.color:
                            _fill_cache[key] = PatternFill(fill_type="solid", fgColor=fill.color)
                        elif fill.fill_type:
                            _fill_cache[key] = PatternFill(fill_type=fill.fill_type)
                        else:
                            _fill_cache[key] = None
                    if _fill_cache[key] is not None:
                        cell_obj.fill = _fill_cache[key]
                        
                if cell_model.style.border:
                    border = cell_model.style.border
                    key = (border.left, border.right, border.top, border.bottom)
                    if key not in _border_cache:
                        kwargs = {}
                        if border.left: kwargs['left'] = Side(style=border.left)
                        if border.right: kwargs['right'] = Side(style=border.right)
                        if border.top: kwargs['top'] = Side(style=border.top)
                        if border.bottom: kwargs['bottom'] = Side(style=border.bottom)
                        _border_cache[key] = Border(**kwargs)
                    cell_obj.border = _border_cache[key]
                    
                if cell_model.style.number_format:
                    cell_obj.number_format = cell_model.style.number_format
                    
            # Apply merges
            if cell_model.merge:
                merge = cell_model.merge
                # Calculate actual row/col span target
                merge_min_row = target_row_idx
                merge_max_row = target_row_idx + merge.row_span - 1
                merge_min_col = merge.min_col
                merge_max_col = merge.max_col
                
                ws.merge_cells(
                    start_row=merge_min_row, start_column=merge_min_col,
                    end_row=merge_max_row, end_column=merge_max_col
                )
