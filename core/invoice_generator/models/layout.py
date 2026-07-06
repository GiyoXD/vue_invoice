from dataclasses import dataclass, field
from typing import List, Tuple, Optional, Any, Dict, Set, Union
import logging
from copy import deepcopy

from openpyxl.utils import get_column_letter
from core.models.cell import UnitRow, UnitCell
from ..utils.cell_converter import convert_registry_style_to_cell_style, write_models_to_worksheet

logger = logging.getLogger(__name__)

@dataclass
class TableZoneBoundary:
    """Represents the row boundaries for a single built table on the sheet."""
    table_key: str
    header_range: Tuple[int, int]
    data_range: Tuple[int, int]
    footer_range: Tuple[int, int]

@dataclass
class SheetLayoutState:
    """
    Tracks the layout of a sheet as it is built sequentially.
    Manages occupied row ranges to determine the next free row.
    Also coordinates coordinate mapping, cell writing, and merge safety.
    """
    table_zones: List[TableZoneBoundary] = field(default_factory=list)
    
    # Template zones
    template_header_range: Optional[Tuple[int, int]] = None
    template_footer_range: Optional[Tuple[int, int]] = None
    
    # Internal tracker for the next available row on the sheet
    _next_free_row: int = 1
    
    # --- Grid Tracking state (Set on binding) ---
    occupied_cells: Set[Tuple[int, int]] = field(default_factory=set)  # Set of (row, col) coordinates swallowed by merges (blocked from writing)
    merged_cells: Set[Tuple[int, int]] = field(default_factory=set)    # Set of all coordinates inside any merge (blocked from further merges)
    column_mapping: Dict[str, int] = field(default_factory=dict)       # logical_id -> physical_index (1-based)
    _rows_with_height_applied: Set[int] = field(default_factory=set)
    style_registry: Any = None
    dimension_registry: Any = None
    ws: Any = None
    
    # Track the data range of the current table for formula generation
    current_data_range: Optional[Tuple[int, int]] = None

    @property
    def next_free_row(self) -> int:
        """Returns the next available row on the sheet that has not been built on."""
        return self._next_free_row
        
    def advance_to(self, row: int):
        """Advances the free row pointer."""
        self._next_free_row = max(self._next_free_row, row)

    def add_table_zone(self, table_key: str, header: Tuple[int, int], data: Tuple[int, int], footer: Tuple[int, int]):
        """Records a built table zone and advances the free row pointer."""
        self.table_zones.append(TableZoneBoundary(
            table_key=table_key,
            header_range=header,
            data_range=data,
            footer_range=footer
        ))
        
        # Advance the free row past the footer
        if footer and footer[1] >= footer[0]:
            self.advance_to(footer[1] + 1)
        elif data and data[1] >= data[0]:
            self.advance_to(data[1] + 1)
        elif header and header[1] >= header[0]:
            self.advance_to(header[1] + 1)

    # --- Binding & Coordination Methods ---

    def bind(self, worksheet: Any, column_mapping: Dict[str, int], style_registry: Any, dimension_registry: Any = None):
        """Bind layout session context to the state."""
        self.ws = worksheet
        self.column_mapping = column_mapping
        self.style_registry = style_registry
        self.dimension_registry = dimension_registry
        logger.info(f"SheetLayoutState bound to sheet '{worksheet.title if worksheet else 'None'}' with {len(column_mapping)} column mappings")

    def resolve_column(self, col_id_or_idx: Union[str, int]) -> Optional[int]:
        """Resolves a column ID or direct index to a physical 1-based index."""
        if isinstance(col_id_or_idx, int):
            return col_id_or_idx
        return self.column_mapping.get(col_id_or_idx)

    def record_data_range(self, start_row: int, end_row: int):
        """Record the data range of the current table block."""
        self.current_data_range = (start_row, end_row)
        logger.debug(f"Recorded data range: {start_row} to {end_row}")

    def get_sum_formula_for_column(self, col_id_or_idx: Union[str, int]) -> str:
        """Generates a sum formula for the current data range on the target column."""
        col_idx = self.resolve_column(col_id_or_idx)
        if not col_idx or not self.current_data_range:
            return ""
        col_letter = get_column_letter(col_idx)
        start, end = self.current_data_range
        return f"=SUM({col_letter}{start}:{col_letter}{end})"

    def write_cell(self, row: int, col_id_or_idx: Union[str, int], value: Any, context: str = 'data', apply_border: bool = True):
        """Writes a value and applies StyleRegistry styling to a cell safely."""
        col_idx = self.resolve_column(col_id_or_idx)
        if not col_idx or not self.ws:
            return

        # Skip writing if the cell is swallowed/hidden by a vertical or horizontal merge
        if (row, col_idx) in self.occupied_cells:
            logger.debug(f"Skipped write to occupied cell ({row}, {col_idx})")
            return

        # Apply styles from StyleRegistry if context and IDs are available
        style_written = False
        if self.style_registry and isinstance(col_id_or_idx, str):
            style_dict = self.style_registry.get_style(col_id_or_idx, context=context)
            
            # Optionally strip borders (used for grand_total footer sections)
            if not apply_border and style_dict:
                style_dict = deepcopy(style_dict)
                style_dict['border_style'] = None

            if style_dict:
                cell_style = convert_registry_style_to_cell_style(style_dict)
                # Apply CellStyle to openpyxl cell (reusing formatting logic or custom writers)
                unit_row = UnitRow(relative_index=0, height=None, cells=[UnitCell(col_index=col_idx, value=value, style=cell_style)])
                write_models_to_worksheet(self.ws, [unit_row], start_row=row)
                style_written = True

        if not style_written:
            self.ws.cell(row=row, column=col_idx, value=value)

        # Set row height once per row context
        if row not in self._rows_with_height_applied and self.dimension_registry:
            height = self.dimension_registry.get_row_height(context)
            if height:
                self.ws.row_dimensions[row].height = height
            self._rows_with_height_applied.add(row)

    def merge_cells(self, start_row: int, start_col_id_or_idx: Union[str, int], end_row: int, end_col_id_or_idx: Union[str, int]) -> bool:
        """
        Merges cells in the worksheet and marks swallowed cells as occupied.
        Validates that the range does not conflict/overlap with any existing merges.
        """
        start_col = self.resolve_column(start_col_id_or_idx)
        end_col = self.resolve_column(end_col_id_or_idx)
        if not start_col or not end_col or not self.ws:
            return False

        # --- CONFLICT CHECK ---
        # If any cell in the proposed range (including the anchor) is already in a merge, block it
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                if (r, c) in self.merged_cells:
                    logger.warning(f"Merge Conflict blocked: ({r}, {c}) is already part of a merged range.")
                    return False

        # Register merge
        merge_range = f"{get_column_letter(start_col)}{start_row}:{get_column_letter(end_col)}{end_row}"
        try:
            self.ws.merge_cells(merge_range)
        except ValueError as e:
            logger.warning(f"Excel merge failed: {e}")
            return False

        # Register coordinates in layout sets
        for r in range(start_row, end_row + 1):
            for c in range(start_col, end_col + 1):
                self.merged_cells.add((r, c))
                if not (r == start_row and c == start_col):
                    self.occupied_cells.add((r, c))

        logger.debug(f"Registered merge range: {merge_range}")
        return True

    def write_row_models(self, rows: List[Any], start_row: int = 1):
        """
        Writes a list of UnitRow models to the sheet.
        Automatically digests cell merges into coordinator sets to keep track of occupancy.
        """
        if not self.ws:
            return

        # 1. Write the row models using the existing helper
        write_models_to_worksheet(self.ws, rows, start_row)

        # 2. Track merges in occupied_cells and merged_cells
        for row_model in rows:
            target_row_idx = start_row + row_model.relative_index
            self._rows_with_height_applied.add(target_row_idx)

            for cell_model in row_model.cells:
                target_col_idx = cell_model.col_index

                if cell_model.merge:
                    m = cell_model.merge
                    min_r = target_row_idx
                    max_r = target_row_idx + m.row_span - 1
                    min_c = m.min_col
                    max_c = m.max_col

                    for r in range(min_r, max_r + 1):
                        for c in range(min_c, max_c + 1):
                            self.merged_cells.add((r, c))
                            if not (r == min_r and c == min_c):
                                self.occupied_cells.add((r, c))

@dataclass
class TableData:
    """Holds only the content to be rendered in the table grid."""
    table_key: Optional[str]
    columns: List[Dict[str, Any]]
    data_rows: List[Dict[str, Any]]
    local_totals: Dict[str, Any] = field(default_factory=dict)

@dataclass
class AddonSummaryRow:
    """A generic key-value/summary row to append under the table (e.g., Grand Totals)."""
    label: str
    values: Dict[str, Any]
    style_context: str = "grand_total"

@dataclass
class TableLayoutConfig:
    """Explicit styling and layout commands for building a table layout block."""
    is_first_table: bool = True
    is_last_table: bool = True
    skip_template_footer: bool = False
    total_net_weight: Optional[float] = None
    total_gross_weight: Optional[float] = None
    addons: List[AddonSummaryRow] = field(default_factory=list)
