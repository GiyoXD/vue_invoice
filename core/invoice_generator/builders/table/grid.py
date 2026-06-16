import logging
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.utils import get_column_letter

from ...styling.style_registry import StyleRegistry
from ...utils.cell_converter import convert_registry_style_to_cell_style
from core.models.cell import UnitRow, UnitCell, TemplateMerge

logger = logging.getLogger(__name__)

class Grid:
    """
    Virtual translation layer for painting table sections.
    Translates logical IDs to physical indices and applies styles based on context.
    """
    def __init__(self, column_mapping: Dict[str, int], style_registry: StyleRegistry, column_colspan: Optional[Dict[str, int]] = None):
        self.column_mapping = column_mapping
        self.style_registry = style_registry
        self.column_colspan = column_colspan or {}
        self._grid: Dict[int, Dict[int, UnitCell]] = {}
        self._row_heights: Dict[int, Optional[float]] = {}
        self.start_row_index = 0
        self._cursor_row = 0
        self._sections: Dict[str, Tuple[int, int]] = {}

    @property
    def num_columns(self) -> int:
        """Returns the maximum column index in the column mapping."""
        return max(self.column_mapping.values()) if self.column_mapping else 0

    def advance_row(self, count: int = 1):
        """Advances the internal row cursor for subsequent sections."""
        self._cursor_row += count

    def set_start_row(self, row: int):
        """Sets the absolute start row for this grid context, used for formulas."""
        self.start_row_index = row

    def get_column_letter(self, col_id: str) -> str:
        """Resolves logical column ID to Excel column letter."""
        idx = self._resolve_column(col_id)
        if not idx:
            raise ValueError(f"Column ID '{col_id}' not found in grid mapping.")
        return get_column_letter(idx)

    def get_column_index(self, col_id: str) -> int:
        """Resolves logical column ID to physical 1-based column index."""
        idx = self._resolve_column(col_id)
        if not idx:
            raise ValueError(f"Column ID '{col_id}' not found in grid mapping.")
        return idx

    def mark_section_start(self, name: str):
        """Marks the starting relative cursor row of a named section."""
        self._sections[name] = (self._cursor_row, -1)

    def mark_section_end(self, name: str):
        """Marks the ending relative cursor row of a named section."""
        start = self._sections.get(name, (self._cursor_row, -1))[0]
        end = self._cursor_row - 1
        self._sections[name] = (start, end)

    def get_section_range(self, name: str) -> Tuple[int, int]:
        """Returns the physical start and end row indices for the named section."""
        if name not in self._sections:
            # Fallback to current cursor position
            phys = self.start_row_index + self._cursor_row
            return phys, phys
        start, end = self._sections[name]
        phys_start = self.start_row_index + start
        phys_end = self.start_row_index + end
        return phys_start, phys_end

    def _resolve_column(self, col_id: str) -> Optional[int]:
        """Resolves a logical column ID to a physical 1-based index."""
        if isinstance(col_id, int):
            return col_id
        
        try:
            val = int(col_id)
            return val
        except ValueError:
            pass
            
        return self.column_mapping.get(col_id)

    def _get_or_create_cell(self, row: int, col: int) -> UnitCell:
        if row not in self._grid:
            self._grid[row] = {}
        if col not in self._grid[row]:
            self._grid[row][col] = UnitCell(col_index=col, value=None)
        return self._grid[row][col]

    def write(self, row: int, col_id: str, value: Any, context: str = 'data'):
        """
        Writes a value to the grid at the given relative row and logical column ID.
        Applies styling automatically based on the given context.
        """
        col_idx = self._resolve_column(col_id)
        if not col_idx:
            logger.debug(f"Grid: Column ID '{col_id}' not found in mapping.")
            return

        actual_row = row + self._cursor_row
        cell = self._get_or_create_cell(actual_row, col_idx)
        cell.value = value

        # Set row height if not already set
        if actual_row not in self._row_heights:
            self._row_heights[actual_row] = self.style_registry.get_row_height(context)

        # Apply style
        if self.style_registry:
            # We skip 'col_id' existence check strictly here to allow generic styling if provided
            # But normally we look up by col_id. If col_id is an int, it might fail in registry,
            # so we only apply if it's a valid ID string.
            if isinstance(col_id, str):
                style_dict = self.style_registry.get_style(col_id, context=context)
                if style_dict:
                    # Special override rules (e.g. static col in footer) can be passed as kwargs in future if needed
                    # but for now we trust the registry.
                    cell.style = convert_registry_style_to_cell_style(style_dict)

    def write_formula(self, row: int, col_id: str, template: str, inputs: List[str], context: str = 'data'):
        """
        Translates a logical formula into a physical Excel formula and writes it to the grid.
        E.g., template="={col_ref_0}*{col_ref_1}", inputs=["col_qty", "col_price"]
        """
        col_idx = self._resolve_column(col_id)
        if not col_idx:
            return

        formula = template
        for i, input_id in enumerate(inputs):
            input_col_idx = self._resolve_column(input_id)
            if input_col_idx:
                col_letter = get_column_letter(input_col_idx)
                formula = formula.replace(f'{{col_ref_{i}}}', f'{col_letter}{{row}}')

        # Calculate actual absolute row
        actual_row = row + self._cursor_row
        absolute_row = self.start_row_index + actual_row
        formula = formula.replace('{row}', str(absolute_row))

        if not formula.startswith('='):
            formula = '=' + formula

        self.write(row, col_id, formula, context)

    def merge(self, row: int, col_id: str, rowspan: int = 1, colspan: int = 1):
        """
        Merges cells starting from (row, col_id).
        """
        if rowspan <= 1 and colspan <= 1:
            return

        col_idx = self._resolve_column(col_id)
        if not col_idx:
            return

        actual_row = row + self._cursor_row
        cell = self._get_or_create_cell(actual_row, col_idx)
        # Note: TemplateMerge uses absolute coordinates for Excel. We store relative or absolute?
        # The models usually expect physical column indices and physical row span.
        cell.merge = TemplateMerge(
            min_col=col_idx,
            max_col=col_idx + colspan - 1,
            row_span=rowspan,
            value=str(cell.value) if cell.value is not None else ""
        )
        
        # Clear out values in merged span
        for r in range(actual_row, actual_row + rowspan):
            for c in range(col_idx, col_idx + colspan):
                if r == actual_row and c == col_idx:
                    continue
                clear_cell = self._get_or_create_cell(r, c)
                clear_cell.value = None



    def get_cell(self, row: int, col_id: Any) -> UnitCell:
        """Gets or creates a cell at a relative row and column ID/index."""
        col_idx = self._resolve_column(col_id)
        actual_row = row + self._cursor_row
        return self._get_or_create_cell(actual_row, col_idx)

    def get_row_models(self) -> List[UnitRow]:
        """
        Exports the internal grid matrix to a list of UnitRow objects.
        """
        models = []
        if not self._grid:
            return models

        max_row = max(self._grid.keys())
        for r in range(max_row + 1):
            cells = []
            if r in self._grid:
                cells = sorted(list(self._grid[r].values()), key=lambda c: c.col_index)
            
            models.append(UnitRow(
                relative_index=r,
                height=self._row_heights.get(r),
                cells=cells
            ))
        return models
