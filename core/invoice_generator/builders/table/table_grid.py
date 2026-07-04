import logging
from typing import Any, Dict, List, Optional, Tuple

from openpyxl.utils import get_column_letter

from ...styling.style_registry import StyleRegistry
from ...styling.dimension_registry import DimensionRegistry
from ...utils.cell_converter import convert_registry_style_to_cell_style
from core.models.cell import UnitRow, UnitCell, TemplateMerge

from core.models.grid import Grid as CoreGrid

logger = logging.getLogger(__name__)

class TableGrid(CoreGrid):
    """
    Virtual translation layer for painting table sections.
    Translates logical IDs to physical indices and applies styles based on context.
    """
    def __init__(self, column_mapping: Dict[str, int], style_registry: StyleRegistry, column_colspan: Optional[Dict[str, int]] = None, dimension_registry: Optional[DimensionRegistry] = None):
        super().__init__()
        self.column_mapping = column_mapping
        self.style_registry = style_registry
        self.dimension_registry = dimension_registry
        self.column_colspan = column_colspan or {}
        self.start_row_index = 0
        self._cursor_row = 0
        self._sections: Dict[str, Tuple[int, int]] = {}

    @property
    def num_columns(self) -> int:
        """Returns the maximum column index in the column mapping."""
        return max(self.column_mapping.values()) if self.column_mapping else 0

    def advance_row(self, count: int = 1):
        """Advances the internal row cursor for subsequent sections."""
        self._cursor_row += count

    def set_start_row(self, row: int):
        """Sets the absolute start row for this grid context, used for formulas."""
        self.start_row_index = row

    def get_column_letter(self, col_id: str) -> str:
        """Resolves logical column ID to Excel column letter."""
        idx = self._resolve_column(col_id)
        if not idx:
            raise ValueError(f"Column ID '{col_id}' not found in grid mapping.")
        return get_column_letter(idx)

    def get_column_index(self, col_id: str) -> int:
        """Resolves logical column ID to physical 1-based column index."""
        idx = self._resolve_column(col_id)
        if not idx:
            raise ValueError(f"Column ID '{col_id}' not found in grid mapping.")
        return idx

    def mark_section_start(self, name: str):
        """Marks the starting relative cursor row of a named section."""
        self._sections[name] = (self._cursor_row, -1)

    def mark_section_end(self, name: str):
        """Marks the ending relative cursor row of a named section."""
        start = self._sections.get(name, (self._cursor_row, -1))[0]
        end = self._cursor_row - 1
        self._sections[name] = (start, end)

    def set_section_bounds(self, name: str, start: int, end: int):
        """Explicitly sets the start and end rows (relative to grid) for a section."""
        self._sections[name] = (start, end)

    def get_section_range(self, name: str) -> Tuple[int, int]:
        """Returns the physical start and end row indices for the named section."""
        if name not in self._sections:
            # Fallback to current cursor position
            phys = self.start_row_index + self._cursor_row
            return phys, phys
        start, end = self._sections[name]
        phys_start = self.start_row_index + start
        phys_end = self.start_row_index + end
        return phys_start, phys_end

    def _resolve_column(self, col_id: str) -> Optional[int]:
        """Resolves a logical column ID to a physical 1-based index."""
        if isinstance(col_id, int):
            return col_id
        
        try:
            val = int(col_id)
            return val
        except ValueError:
            pass
            
        return self.column_mapping.get(col_id)

    def write(self, row: int, col_id: str, value: Any, context: str = 'data'):
        """
        Writes a value to the grid at the given relative row and logical column ID.
        Applies styling (font, format, alignment, fill) based on the given context.
        Borders are applied separately by BorderResolver after all sections are built.
        """
        col_idx = self._resolve_column(col_id)
        if not col_idx:
            logger.debug(f"Grid: Column ID '{col_id}' not found in mapping.")
            return

        actual_row = row + self._cursor_row
        
        # Set row height if not already set
        if actual_row not in self._row_heights:
            if self.dimension_registry:
                self._row_heights[actual_row] = self.dimension_registry.get_row_height(context)

        # Apply style (font, format, alignment, fill — no borders)
        cell_style = None
        if self.style_registry:
            if isinstance(col_id, str):
                style_dict = self.style_registry.get_style(col_id, context=context)
                if style_dict:
                    cell_style = convert_registry_style_to_cell_style(style_dict)

        super().write(actual_row, col_idx, value, style=cell_style)

    def write_formula(self, row: int, col_id: str, template: str, inputs: List[str], context: str = 'data'):
        """
        Translates a logical formula into a physical Excel formula and writes it to the grid.
        E.g., template="={col_ref_0}*{col_ref_1}", inputs=["col_qty", "col_price"]
        """
        col_idx = self._resolve_column(col_id)
        if not col_idx:
            return

        formula = template
        for i, input_id in enumerate(inputs):
            input_col_idx = self._resolve_column(input_id)
            if input_col_idx:
                col_letter = get_column_letter(input_col_idx)
                formula = formula.replace(f'{{col_ref_{i}}}', f'{col_letter}{{row}}')

        # Calculate actual absolute row
        actual_row = row + self._cursor_row
        absolute_row = self.start_row_index + actual_row
        formula = formula.replace('{row}', str(absolute_row))

        if not formula.startswith('='):
            formula = '=' + formula

        self.write(row, col_id, formula, context)

    def write_section_aggregate(self, row: int, col_id: str, function: str = "SUM", section: str = "data", context: str = 'footer'):
        """
        Automatically calculates the absolute coordinates for a section range on a column 
        and writes the formula to the grid. E.g. write_section_aggregate(row, 'col_qty', 'SUM', 'data')
        """
        col_idx = self._resolve_column(col_id)
        if not col_idx:
            return

        start, end = self.get_section_range(section)
        if start > 0 and end >= start:
            col_letter = get_column_letter(col_idx)
            formula = f"={function}({col_letter}{start}:{col_letter}{end})"
            self.write(row, col_id, formula, context)

    def merge(self, row: int, col_id: str, rowspan: int = 1, colspan: int = 1):
        """
        Merges cells starting from (row, col_id).
        """
        if rowspan <= 1 and colspan <= 1:
            return

        col_idx = self._resolve_column(col_id)
        if not col_idx:
            return

        actual_row = row + self._cursor_row
        super().merge(actual_row, col_idx, rowspan, colspan)

    def get_cell(self, row: int, col_id: Any, resolve_merge: bool = False) -> UnitCell:
        """Gets or creates a cell at a relative row and column ID/index."""
        col_idx = self._resolve_column(col_id)
        actual_row = row + self._cursor_row
        return super().get_cell(actual_row, col_idx, resolve_merge=resolve_merge)

    def to_dict(self) -> Dict[str, Any]:
        """
        Serializes the TableGrid subclass and base Grid components to a dictionary.
        """
        d = super().to_dict()
        d.update({
            "column_mapping": self.column_mapping,
            "column_colspan": self.column_colspan,
            "start_row_index": self.start_row_index,
            "cursor_row": self._cursor_row,
            "sections": self._sections
        })
        return d

    @classmethod
    def from_dict(cls, d: Dict[str, Any], style_registry: Any = None, dimension_registry: Any = None) -> 'TableGrid':
        """
        Deserializes a dictionary into a TableGrid instance with an optional style_registry context.
        """
        grid_obj = cls(
            column_mapping=d.get("column_mapping", {}),
            style_registry=style_registry,
            column_colspan=d.get("column_colspan", {}),
            dimension_registry=dimension_registry
        )
        grid_obj.start_row_index = d.get("start_row_index", 0)
        grid_obj._cursor_row = d.get("cursor_row", 0)
        grid_obj._sections = {k: tuple(v) for k, v in d.get("sections", {}).items()}
        grid_obj._row_heights = {int(k): v for k, v in d.get("row_heights", {}).items()}
        
        # Re-populate physical merges
        for k, v in d.get("merge_map", {}).items():
            r, c = map(int, k.split(','))
            grid_obj._merge_map[(r, c)] = (v[0], v[1])
            
        # Re-populate physical cells
        grid_raw = d.get("grid", {})
        for r_str, row_data in grid_raw.items():
            r = int(r_str)
            grid_obj._grid[r] = {}
            for c_str, cell_data in row_data.items():
                c = int(c_str)
                grid_obj._grid[r][c] = UnitCell.from_dict(cell_data)
                
        return grid_obj

# Alias for backward compatibility
Grid = TableGrid
