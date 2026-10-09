import logging
import decimal
import re
from typing import List, Dict, Any, Optional
from openpyxl.utils import get_column_letter
from ..config import HEADER_SEARCH_COL_RANGE, HEADERLESS_COLUMN_PATTERNS
from .exceptions import DataValidationError
from .helpers import _find_first_data_row, _is_valid_cell_value


def validate_no_duplicate_amount_columns(sheet, header_row: int, column_mapping: Optional[Dict[str, str]] = None):
    """
    Scans the header row and raises DataValidationError if multiple columns
    are mapped to 'col_amount' (via header aliases or data pattern matching).
    """
    from ..sheet_parser import _ALIAS_REVERSE_LOOKUP

    # Find the first row below header that actually has data
    data_row = _find_first_data_row(sheet, header_row)

    # Collect columns that have already been mapped to other headers (not col_amount)
    mapped_cols = set()
    if column_mapping:
        for canonical, col_letter in column_mapping.items():
            if canonical != 'col_amount' and col_letter:
                mapped_cols.add(col_letter.upper())

    amount_cols = []
    alias_matched_cols = []
    
    # 1. First pass: check explicit header text alias mapping
    for col_num in range(HEADER_SEARCH_COL_RANGE[0], HEADER_SEARCH_COL_RANGE[1] + 1):
        cell = sheet.cell(row=header_row, column=col_num)
        val = str(cell.value or '').strip().upper()
        if val:
            col_letter = get_column_letter(col_num)
            if col_letter.upper() not in mapped_cols:
                candidates = _ALIAS_REVERSE_LOOKUP.get(val, [])
                if "col_amount" in candidates:
                    alias_matched_cols.append(col_num)
                    amount_cols.append(f"{get_column_letter(col_num)} ('{cell.value or '<empty>'}')")

    # 2. Second pass: check unrecognized/unmapped columns using pattern + left adjacent cell < 2 heuristic
    has_explicit_amount = len(alias_matched_cols) > 0
    for col_num in range(HEADER_SEARCH_COL_RANGE[0], HEADER_SEARCH_COL_RANGE[1] + 1):
        if col_num in alias_matched_cols:
            continue

        cell = sheet.cell(row=header_row, column=col_num)
        val = str(cell.value or '').strip().upper()

        # If we have an explicit alias match elsewhere, only perform the heuristic check
        # on truly headerless/empty columns. If no explicit alias match exists, we can
        # check any unrecognized column.
        if has_explicit_amount and val != "":
            continue

        col_letter = get_column_letter(col_num)
        if col_letter.upper() not in mapped_cols:
            amount_patterns = HEADERLESS_COLUMN_PATTERNS.get("col_amount", [])
            if amount_patterns:
                data_cell = sheet.cell(row=data_row, column=col_num)
                data_val = data_cell.value
                
                matches_pattern = False
                if data_val is not None:
                    # Robust handling of float stringification for pattern match
                    val_strs = [str(data_val).strip()]
                    try:
                        # Only format as 2-decimal if int/float/Decimal or contains '.'
                        if isinstance(data_val, (int, float, decimal.Decimal)) or '.' in str(data_val):
                            clean_str = str(data_val).replace(',', '').strip()
                            if re.match(r'^-?\d+(\.\d+)?$', clean_str):
                                num_val = float(clean_str)
                                val_strs.append(f"{num_val:.2f}")
                    except (ValueError, TypeError):
                        pass

                    for pattern in amount_patterns:
                        for val_str in val_strs:
                            try:
                                if re.match(pattern, val_str):
                                    matches_pattern = True
                                    break
                            except re.error:
                                continue
                        if matches_pattern:
                            break
                            
                # Left adjacent cell heuristic check (value < 2)
                has_left_unit_price = False
                if col_num > 1:
                    left_cell = sheet.cell(row=data_row, column=col_num - 1)
                    left_val = left_cell.value
                    left_num = None
                    if left_val is not None:
                        try:
                            left_num = float(str(left_val).replace(',', '').strip())
                        except ValueError:
                            pass
                    
                    if left_num is not None and 0 < left_num < 2:
                        # Ensure current cell also has a valid numeric value
                        if data_val is not None:
                            try:
                                float(str(data_val).replace(',', '').strip())
                                has_left_unit_price = True
                            except ValueError:
                                pass
                                
                if matches_pattern and has_left_unit_price:
                    amount_cols.append(f"{get_column_letter(col_num)} ('{cell.value or '<empty>'}')")
                
    if len(amount_cols) > 1:
        raise DataValidationError(
            f"Data Validation Error: Duplicate 'col_amount' columns detected on header row {header_row}: {', '.join(amount_cols)}. "
            f"Please ensure only one column is mapped to 'Amount'."
        )


def validate_col_level_discount(
    table_data: List[Dict[str, Any]], 
    table_id_str: Optional[str] = None, 
    monitor: Optional[Any] = None
) -> None:
    """
    Scans table_data for any rows where 'col_level' contains '折扣'.
    Logs a warning if discount rows are found.
    """
    if not table_data:
        return

    discount_rows = []
    for idx, row in enumerate(table_data):
        val = row.get('col_level')
        if val is not None:
            try:
                if '折扣' in str(val):
                    discount_rows.append(row.get('_row_num', idx + 1))
            except Exception:
                continue

    if discount_rows:
        row_list = ', '.join(str(r) for r in discount_rows)
        msg = f"[{table_id_str or 'Table'}] Warning: '折扣' (discount) detected in col_level on {len(discount_rows)} row(s): row(s) {row_list}."
        logging.warning(msg)
        if monitor:
            monitor.log_warning(msg)


check_discount_in_col_level = validate_col_level_discount


def validate_table_data_presence(
    current_table_data: List[Dict[str, Any]], 
    table_id_str: str, 
    column_mapping: Dict[str, str], 
    monitor: Optional[Any] = None,
    parser_mode: str = "standard"
):
    """
    Validates that required columns exist in headers and contain valid values across rows.
    """
    validate_col_level_discount(current_table_data, table_id_str, monitor=monitor)

    if parser_mode == "summary":
        ALWAYS_REQUIRED = ['col_po', 'col_item', 'col_net']
    else:
        ALWAYS_REQUIRED = ['col_po', 'col_item', 'col_qty_pcs', 'col_net', 'col_gross', 'col_cbm']
    STRICT_ALL_ROWS_COLS = ['col_po', 'col_item', 'col_qty_pcs', 'col_qty_sf', 'col_unit_price', 'col_amount']
    PRICING_COLUMNS = ['col_amount', 'col_unit_price', 'col_qty_sf']

    if not current_table_data:
        logging.warning(f"Validation: {table_id_str} has no data rows to validate.")
        return

    first_row = current_table_data[0]
    required_cols = list(ALWAYS_REQUIRED) + [c for c in PRICING_COLUMNS if c in column_mapping]
    missing_data_cols = [f"{c} (Missing Header)" for c in required_cols if c not in column_mapping]

    for col in required_cols:
        if col in column_mapping:
            is_num = col not in ('col_po', 'col_item')
            if not _is_valid_cell_value(first_row.get(col), is_numeric=is_num):
                missing_data_cols.append(col)

    if missing_data_cols:
        row_num = first_row.get('_row_num', '?')
        err_msg = f"Row {row_num}: Missing columns: {', '.join(missing_data_cols)}"
        if monitor:
            monitor.log_process_item(f"{table_id_str} First-Row Validation", status="error", error=err_msg)
        raise DataValidationError(err_msg)

    # Validate all rows for critical identifiers, qty & pricing columns
    check_all_cols = [c for c in STRICT_ALL_ROWS_COLS if c in column_mapping]

    for row in current_table_data:
        missing_in_row = [
            c for c in check_all_cols
            if not _is_valid_cell_value(row.get(c), is_numeric=(c not in ('col_po', 'col_item')))
        ]
        if missing_in_row:
            row_num = row.get('_row_num', '?')
            err_msg = f"Row {row_num}: Missing or zero value for: {', '.join(missing_in_row)}"
            if monitor:
                monitor.log_process_item(f"{table_id_str} Row Validation", status="error", error=err_msg)
            raise DataValidationError(err_msg)
