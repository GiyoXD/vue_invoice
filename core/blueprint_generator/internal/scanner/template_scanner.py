import re
import logging
import hashlib
import json
from typing import Dict, List, Any, Optional, Tuple
from openpyxl.worksheet.worksheet import Worksheet
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils import get_column_letter

from core.utils.loop_profiler import tick
from .models import (
    ZoneBoundaries, ColumnInfo, FooterInfo, TemplateLayout,
    UnitRow, UnitCell, TemplateMerge, CellStyle,
    FontStyle, AlignmentStyle, FillStyle, BorderStyle
)

logger = logging.getLogger(__name__)

class TemplateScanner:
    """Scans static template zones (Zone 1 & 3): content outside the table area."""

    def __init__(self):
        self.logger = logging.getLogger(self.__class__.__name__)
        # Excel default dimensions (used to filter out empty cells with default size)
        self.DEFAULT_ROW_HEIGHT = 15.0  # Excel default row height in points
        self.DEFAULT_COL_WIDTH = 8.43   # Excel default column width in characters

    def scan_static_content(self, worksheet: Worksheet, boundaries: ZoneBoundaries, columns: List[ColumnInfo], sheet_name: str) -> TemplateLayout:
        """
        Scan static content zones outside the table area (Zone 1 & 3).
        Returns the layout metadata.
        """
        self.logger.info(f"  Scanning static template layout: {sheet_name}")
        
        layout = TemplateLayout()
        safe_max_column = boundaries.max_col
        
        layout.header_rows = self._capture_header_rows(worksheet, boundaries, safe_max_column)
        layout.footer_rows = self._capture_footer_rows(worksheet, boundaries, safe_max_column, sheet_name)
        
        return layout

    def _capture_header_rows(self, ws: Worksheet, boundaries: ZoneBoundaries, safe_max_column: int) -> List[UnitRow]:
        header_rows = []
        if boundaries.header_row <= 1:
            return header_rows

        # Capture merges in header
        merges_map = {}
        for merged_range in ws.merged_cells.ranges:
            if merged_range.max_row < boundaries.header_row:
                 top_left_cell = ws.cell(row=merged_range.min_row, column=merged_range.min_col)
                 val = str(top_left_cell.value) if top_left_cell.value is not None else ""
                 merges_map[(merged_range.min_row, merged_range.min_col)] = TemplateMerge(
                     min_col=merged_range.min_col,
                     max_col=merged_range.max_col,
                     row_span=merged_range.max_row - merged_range.min_row + 1,
                     value=val.strip()
                 )

        for r in boundaries.template_header_range:
            row_cells = []
            row_height = None
            
            if r in ws.row_dimensions:
                h = ws.row_dimensions[r].height
                if h is not None:
                     row_height = h
            
            has_content_or_style = False
            for c in range(1, safe_max_column + 1):
                cell = ws.cell(row=r, column=c)
                is_empty = (cell.value is None)
                
                style_obj = self._capture_cell_style(cell, is_empty=is_empty)
                merge_obj = merges_map.get((r, c))
                
                if is_empty and not style_obj and not merge_obj:
                    continue
                    
                has_content_or_style = True
                cell_obj = UnitCell(col_index=c, style=style_obj, merge=merge_obj)
                
                if not is_empty:
                     val_str = str(cell.value)
                     if val_str.startswith('='):
                         val_str = re.sub(r'\[\d+\]', '', val_str)
                     cell_obj.value = val_str
                    
                row_cells.append(cell_obj)
                
            if row_height is not None or has_content_or_style:
                header_rows.append(UnitRow(
                    relative_index=r - 1,
                    height=row_height,
                    cells=row_cells
                ))
        return header_rows

    def _capture_footer_rows(self, ws: Worksheet, boundaries: ZoneBoundaries, safe_max_column: int, sheet_name: str) -> List[UnitRow]:
        if boundaries.footer_row is None:
            self.logger.warning(
                f"    [SKIP] Sheet '{sheet_name}': table footer (TOTAL row) not found "
                f"(scanned from row {boundaries.header_row + 1} to end-of-sheet). "
                f"Treating as Form/Static sheet."
            )
            return []

        end_delete = boundaries.footer_row
        start_delete = boundaries.header_row
        
        self.logger.info(f"    Capturing footer data (Rows {end_delete + 1} to EOF)")
        
        footer_merges_map = {}
        for merged_range in list(ws.merged_cells.ranges):
            m_min_row, m_min_col, m_max_row, m_max_col = merged_range.min_row, merged_range.min_col, merged_range.max_row, merged_range.max_col
            
            if m_min_row > start_delete:
                top_left_cell = ws.cell(row=m_min_row, column=m_min_col)
                val = str(top_left_cell.value) if top_left_cell.value is not None else ""
                footer_merges_map[(m_min_row, m_min_col)] = TemplateMerge(
                    min_col=m_min_col,
                    max_col=m_max_col,
                    row_span=m_max_row - m_min_row + 1,
                    value=val.strip()
                )

        template_footer_rows = []
        
        for r in boundaries.template_footer_range(ws.max_row):
            tick("scanner._capture_template_footer", sub="rows_processed")
            rel_r = r - (end_delete + 1)
            
            row_height = None
            if r in ws.row_dimensions:
                h = ws.row_dimensions[r].height
                if h is not None:
                    row_height = h
                    
            row_cells = []
            has_content_or_style = False
            for c in range(1, safe_max_column + 1):
                cell = ws.cell(row=r, column=c)
                is_empty = (cell.value is None)
                
                style_obj = self._capture_cell_style(cell, is_empty=is_empty)
                merge_obj = footer_merges_map.get((r, c))
                
                if is_empty and not style_obj and not merge_obj:
                    continue
                    
                cell_obj = UnitCell(col_index=c, style=style_obj, merge=merge_obj)
                has_content_or_style = True
                
                if not is_empty:
                    val_str = str(cell.value)
                    if val_str.startswith('='):
                        val_str = re.sub(r'\[\d+\]', '', val_str)
                    cell_obj.value = val_str
                    
                row_cells.append(cell_obj)
                
            if row_height is not None or has_content_or_style:
                template_footer_rows.append(UnitRow(
                    relative_index=rel_r,
                    height=row_height,
                    cells=row_cells
                ))
                
        return template_footer_rows



    def _safe_str(self, val) -> Optional[str]:
        return val if isinstance(val, str) else None

    def _safe_num(self, val) -> Optional[float]:
        return val if isinstance(val, (int, float)) and not isinstance(val, bool) else None

    def _safe_bool(self, val) -> bool:
        return val if isinstance(val, bool) else False

    def _capture_cell_style(self, cell: Cell, is_empty: bool = False) -> Optional[CellStyle]:
        font_style = None
        alignment_style = None
        fill_style = None
        border_style = None
        number_format = "General"
        has_significant_style = False
        
        # 1. Font
        if cell.font:
            # Handle MagicMocks
            if not (hasattr(cell.font, "_mock_return_value") or hasattr(cell.font, "mock_add_spec")):
                font_data = {}
                name = self._safe_str(cell.font.name)
                size = self._safe_num(cell.font.size)
                bold = self._safe_bool(cell.font.bold)
                italic = self._safe_bool(cell.font.italic)
                
                if not is_empty:
                    if name and name not in ["Calibri", "Arial"]:
                        font_data["name"] = name
                    if size is not None:
                        is_calibri = (name == "Calibri")
                        is_size_11 = (size in [11.0, 11])
                        if not (is_calibri and is_size_11):
                            font_data["size"] = size
                if bold: font_data["bold"] = True
                if italic: font_data["italic"] = True
                if cell.font.color and hasattr(cell.font.color, "rgb"):
                     color_val = self._serialize_color(cell.font.color)
                     if color_val and color_val not in ("00000000", "FF000000"):
                          if not is_empty or (color_val != "FFFFFFFF" and not color_val.startswith("theme-")):
                              font_data["color"] = color_val

                if font_data:
                    font_style = FontStyle(
                        name=font_data.get("name"),
                        size=font_data.get("size"),
                        bold=font_data.get("bold", False),
                        italic=font_data.get("italic", False),
                        color=font_data.get("color")
                    )
                    if not is_empty:
                        has_significant_style = True
            
        # 2. Alignment
        if cell.alignment:
            # Handle MagicMocks
            if not (hasattr(cell.alignment, "_mock_return_value") or hasattr(cell.alignment, "mock_add_spec")):
                align_data = {}
                horizontal = self._safe_str(cell.alignment.horizontal)
                vertical = self._safe_str(cell.alignment.vertical)
                wrap_text = self._safe_bool(cell.alignment.wrap_text)
                
                if horizontal and horizontal != 'general':
                    align_data["horizontal"] = horizontal
                if vertical and vertical != 'bottom':
                    align_data["vertical"] = vertical
                if wrap_text:
                    align_data["wrap_text"] = True
                if align_data:
                    alignment_style = AlignmentStyle(
                        horizontal=align_data.get("horizontal"),
                        vertical=align_data.get("vertical"),
                        wrap_text=align_data.get("wrap_text", False)
                    )
                    if not is_empty:
                        has_significant_style = True
            
        # 3. Fill
        if cell.fill:
            # Handle MagicMocks
            if not (hasattr(cell.fill, "_mock_return_value") or hasattr(cell.fill, "mock_add_spec")):
                fill_type = self._safe_str(cell.fill.fill_type)
                if fill_type and fill_type != "none":
                    if hasattr(cell.fill, "start_color"):
                         color_val = self._serialize_color(cell.fill.start_color)
                         if color_val and color_val not in ["00000000", "FFFFFFFF"]:
                             fill_style = FillStyle(
                                 fill_type=fill_type,
                                 color=color_val
                             )
                             has_significant_style = True
             
        # 4. Border
        if cell.border:
            # Handle MagicMocks
            if not (hasattr(cell.border, "_mock_return_value") or hasattr(cell.border, "mock_add_spec")):
                 border_data = {}
                 left_style = self._safe_str(cell.border.left.style) if cell.border.left else None
                 right_style = self._safe_str(cell.border.right.style) if cell.border.right else None
                 top_style = self._safe_str(cell.border.top.style) if cell.border.top else None
                 bottom_style = self._safe_str(cell.border.bottom.style) if cell.border.bottom else None
                 
                 if left_style: border_data["left"] = left_style
                 if right_style: border_data["right"] = right_style
                 if top_style: border_data["top"] = top_style
                 if bottom_style: border_data["bottom"] = bottom_style
                 if border_data:
                     border_style = BorderStyle(
                         left=border_data.get("left"),
                         right=border_data.get("right"),
                         top=border_data.get("top"),
                         bottom=border_data.get("bottom")
                     )
                     has_significant_style = True
             
        # 5. Number Format
        number_fmt = self._safe_str(cell.number_format)
        if number_fmt and number_fmt != "General":
            number_format = number_fmt
            if not is_empty:
                has_significant_style = True
        
        if has_significant_style:
            return CellStyle(
                font=font_style,
                alignment=alignment_style,
                fill=fill_style,
                border=border_style,
                number_format=number_format
            )
        return None

    def _serialize_color(self, color) -> Optional[str]:
        if color is None: return None
        if hasattr(color, "_mock_return_value") or hasattr(color, "mock_add_spec"):
            return None
        if hasattr(color, "rgb") and color.rgb:
            rgb_val = color.rgb
            if isinstance(rgb_val, str):
                return rgb_val
        if hasattr(color, "theme") and color.theme is not None:
            theme_val = color.theme
            if isinstance(theme_val, (int, str)):
                return f"theme-{theme_val}"
        return None

