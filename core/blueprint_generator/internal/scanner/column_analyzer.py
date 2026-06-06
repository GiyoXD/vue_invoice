import logging
from typing import Dict, List, Any, Optional, Tuple
from collections import Counter
from openpyxl.worksheet.worksheet import Worksheet

from core.blueprint_generator.schema import BlueprintSchema
from core.utils.loop_profiler import loop_profiler, tick
from core.blueprint_generator.utils.openpyxl_utils import get_actual_column_width
from .models import ColumnInfo
from .header_detector import get_cell_value

logger = logging.getLogger(__name__)


class ColumnAnalyzer:
    """Analyzes table columns, types, names, and formats."""

    def __init__(self):
        self.logger = logger

    @loop_profiler.watch("scanner._analyze_columns")
    def analyze_columns(self, worksheet: Worksheet, header_row: int, 
                        header_cells: List[Tuple[int, str]],
                        data_start_row: int,
                        mapping_config: Optional[Dict[str, Any]] = None) -> List[ColumnInfo]:
        """Analyze columns from header row."""
        columns = []
        processed_cols = set()
        
        # Get merged cell ranges that START at the header row.
        # Used for colspan/rowspan detection of actual table header columns.
        merged_ranges = []
        for merged in worksheet.merged_cells.ranges:
            tick("scanner._analyze_columns", sub="merge_scan_header")
            if merged.min_row == header_row:
                merged_ranges.append(merged)
        
        # Also collect ALL merges that overlap the header row (including from above).
        # Used for resolving cell values within merged regions.
        all_merges_at_header = []
        for merged in worksheet.merged_cells.ranges:
            tick("scanner._analyze_columns", sub="merge_scan_overlap")
            if merged.min_row <= header_row <= merged.max_row:
                all_merges_at_header.append(merged)
        
        # Cap column analysis at 25 (Col Y)
        safe_max = min(worksheet.max_column + 1, BlueprintSchema.MAX_SCAN_COLUMN)
        
        for col in range(1, safe_max):
            if col in processed_cols:
                continue
            
            cell = worksheet.cell(row=header_row, column=col)
            value = get_cell_value(cell)
            
            if not value:
                # Check if this is part of a merged cell
                for merged in all_merges_at_header:
                    if merged.min_col <= col <= merged.max_col:
                        # Get value from top-left of merged range
                        value = get_cell_value(
                            worksheet.cell(row=merged.min_row, column=merged.min_col)
                        )
                        break
            
            if not value:
                continue
            
            # Determine column ID
            col_id = self.determine_column_id(value, col, mapping_config)
            
            # [Smart Feature] Leak Filter: Ignore long, unmapped template headers sitting on the same row.
            if col_id.startswith("col_unknown_") and len(value) > 35:
                self.logger.info(f"    [Leak Filter] Ignored long unmapped text as column header: '{value[:30]}...'")
                continue
                
            # Check for merged cells (colspan/rowspan) BEFORE calculating width
            colspan = 1
            rowspan = 1
            for merged in merged_ranges:
                if merged.min_col == col and merged.min_row == header_row:
                    colspan = merged.max_col - merged.min_col + 1
                    rowspan = merged.max_row - merged.min_row + 1
                    # Mark these columns as processed
                    for c in range(merged.min_col, merged.max_col + 1):
                        processed_cols.add(c)
                    break
                    
            # Get total column width spanning all merged cells
            width = get_actual_column_width(worksheet, col, colspan)
            
            # [Smart Feature] Determine format: Check data first, then Rules
            # We use the calculated data_start_row which accounts for multi-row headers.
            format_str = self.sample_column_format(worksheet, col, data_start_row)
            if not format_str or format_str == "General":
                # Fallback to rules if no data or general format
                format_str = self.determine_format(col_id, value)
                
            # FORCE text format for identifiers to prevent scientific notation or leading zero loss
            if col_id in ["col_po", "col_item", "col_no", "col_container_no", "col_hs_code", "col_pallet_count", "col_dc"]:
                format_str = "@"
            
            # Check alignment
            alignment = "center"
            if cell.alignment:
                alignment = cell.alignment.horizontal or "center"
            
            # Check wrap text
            wrap_text = cell.alignment.wrap_text if cell.alignment else False
            
            column = ColumnInfo(
                id=col_id,
                header=value,
                col_index=col,
                width=width,
                format=format_str,
                alignment=alignment,
                rowspan=rowspan,
                colspan=colspan,
                wrap_text=wrap_text
            )
            
            # Check for child columns (multi-row headers)
            if rowspan == 1 and colspan > 1:
                # To be a true parent, the row below MUST be split into multiple smaller columns (or cells)
                # and at least one of those cells should have text that corresponds to a mapping.
                # If the row below is merged across the EXACT SAME columns, it's just a wide data column.
                is_true_parent = False
                
                # Check how the row below is structured within this parent's colspan
                child_cells_with_data = 0
                for c in range(col, col + colspan):
                    child_cell = worksheet.cell(row=header_row + 1, column=c)
                    
                    # Look if this cell is the start of a merge covering the whole parent area
                    is_full_width_merge = False
                    for merged in worksheet.merged_cells.ranges:
                        tick("scanner._analyze_columns", sub="child_merge_scan")
                        if merged.min_row == header_row + 1 and merged.min_col == col and merged.max_col == col + colspan - 1:
                            is_full_width_merge = True
                            break
                            
                    if is_full_width_merge:
                        break # It's exactly the same width as the parent. Not a parent header.
                        
                    val = get_cell_value(child_cell)
                    if val:
                        # Check if it matches a known mapping to be safe
                        mapped_id = self.determine_column_id(val, c, mapping_config)
                        if mapped_id and not mapped_id.startswith("col_unknown"):
                            # Found a valid mapped child! This is a real parent header.
                            is_true_parent = True
                            break
                        child_cells_with_data += 1
                
                # If we found multiple independent cells with data under it, it's a parent
                if child_cells_with_data > 1:
                    is_true_parent = True
                    
                if is_true_parent:
                    children = self.find_child_columns(worksheet, header_row + 1, col, colspan, mapping_config)
                    column.children = children
            
            columns.append(column)
            processed_cols.add(col)
        
        return columns

    @loop_profiler.watch("scanner._sample_column_format")
    def sample_column_format(self, worksheet: Worksheet, col: int, start_row: int, max_rows: int = 15) -> Optional[str]:
        """
        Sample data rows to find the most common number format.
        
        1. Start from start_row (after header).
        2. Scan until we find the first row that actually contains numeric data.
        3. Extract the format from that row and subsequent data rows.
        """
        formats = []
        rows_to_scan = min(start_row + 500, worksheet.max_row + 1) # Scan up to 500 rows to find start of data
        
        data_found_at = None
        
        # Step 1: Find the first row with a numeric value in this column
        for r in range(start_row, rows_to_scan):
            tick("scanner._sample_column_format", sub="find_numeric_rows")
            cell = worksheet.cell(row=r, column=col)
            if cell.value is not None:
                # Check if it's numeric (the real indicator of a data row for amounts/qtys)
                # Exclude booleans which are technically ints in Python
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    data_found_at = r
                    break

        if not data_found_at:
             # Fallback: if no numeric data found, try sampling empty cells for their pre-set format
             # This handles cases where a template is blank but the cells are formatted.
             for r in range(start_row, min(start_row + 500, worksheet.max_row + 1)):
                 tick("scanner._sample_column_format", sub="fallback_format_rows")
                 cell = worksheet.cell(row=r, column=col)
                 if cell.number_format and cell.number_format != 'General':
                     formats.append(cell.number_format)
        else:
            # Step 2: Sample up to max_rows starting from the first data row
            for r in range(data_found_at, min(data_found_at + max_rows, worksheet.max_row + 1)):
                tick("scanner._sample_column_format", sub="sample_rows")
                cell = worksheet.cell(row=r, column=col)
                if cell.value is not None and cell.number_format and cell.number_format != 'General':
                    if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                        formats.append(cell.number_format)
        
        if not formats:
            return None
            
        # Find most common
        most_common = Counter(formats).most_common(1)
        return most_common[0][0] if most_common else None

    def find_child_columns(self, worksheet: Worksheet, row: int, 
                            start_col: int, span: int,
                            mapping_config: Optional[Dict[str, Any]] = None) -> List[ColumnInfo]:
        """Find child columns under a parent header."""
        children = []
        for col in range(start_col, start_col + span):
            cell = worksheet.cell(row=row, column=col)
            value = get_cell_value(cell)
            if value:
                # Value is the Child Header text (e.g. "BUFFALO LEATHER")
                col_id = self.determine_column_id(value, col, mapping_config)
                
                # Basic logging to debug mapping failures
                if not col_id or col_id.startswith("col_unknown"):
                    self.logger.debug(f"Child Column '{value}' at index {col} mapped to unknown.")

                format_str = self.determine_format(col_id, value)
                
                # FORCE text format for identifiers to prevent scientific notation or leading zero loss
                if col_id in ["col_po", "col_item", "col_no", "col_container_no", "col_hs_code", "col_pallet_count", "col_dc"]:
                    format_str = "@"
                    
                # Child columns might also be merged (though rare), but usually colspan=1
                # We calculate standard width to be consistent
                child_colspan = 1
                for merged in worksheet.merged_cells.ranges:
                    if merged.min_col == col and merged.min_row == row:
                        child_colspan = merged.max_col - merged.min_col + 1
                        break
                        
                width = get_actual_column_width(worksheet, col, child_colspan)
                
                children.append(ColumnInfo(
                    id=col_id,
                    header=value,
                    col_index=col,
                    width=width,
                    format=format_str
                ))
        return children

    def determine_column_id(self, header_text: str, col_index: int,
                             mapping_config: Optional[Dict[str, Any]] = None) -> str:
        """Determine column ID from header text using Config first, then Rules."""
        header_text_stripped = header_text.strip()
        
        # 1. Check User Mapping Config
        if mapping_config:
            mappings = mapping_config.get('header_text_mappings', {}).get('mappings', {})
            
            # Exact match (stripped)
            if header_text_stripped in mappings:
                return mappings[header_text_stripped]
                
            # Case-insensitive match (and strip keys in mapping)
            for mapped_header, mapped_id in mappings.items():
                if mapped_header.strip().lower() == header_text_stripped.lower():
                    return mapped_id

        # 2. Use simple rule-based matching (System Defaults)
        col_def = BlueprintSchema.get_column_by_keyword(header_text)
        if col_def:
            return col_def.id
        
        # Fallback: Unknown Column
        return f"col_unknown_{col_index}"
    
    def determine_format(self, col_id: str, header_text: str) -> str:
        """Determine number format for column using Rules."""
        return BlueprintSchema.get_format_for_id(col_id)
