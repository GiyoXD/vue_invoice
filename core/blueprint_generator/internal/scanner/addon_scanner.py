import logging
from typing import List, Dict, Any, Optional
from openpyxl.worksheet.worksheet import Worksheet

from .models.boundaries import ZoneBoundaries
from .models.table import ColumnInfo
from .models.addons import BaseAddonFact, LeatherSummaryFact

logger = logging.getLogger(__name__)


class AddonScanner:
    """Scans worksheet to detect facts/addons outside standard tabular data (Zone 2)."""

    def __init__(self) -> None:
        self.logger = logger

    def scan_addons(self, worksheet: Worksheet, boundaries: ZoneBoundaries,
                    columns: List[ColumnInfo], data_source: str) -> List[BaseAddonFact]:
        """
        Detect and extract addon facts from the worksheet.
        """
        addon_facts: List[BaseAddonFact] = []
        if not boundaries.footer_row or data_source != "processed_tables_multi":
            return addon_facts

        end_scan_row = boundaries.footer_end_row if boundaries.footer_end_row else boundaries.footer_row
        end_scan_row = max(end_scan_row, boundaries.footer_row + 5)
        end_scan_row = min(end_scan_row, worksheet.max_row)

        for r in range(boundaries.footer_row + 1, end_scan_row + 1):
            for c in range(1, min(worksheet.max_column + 1, 40)):
                cell_val = str(worksheet.cell(row=r, column=c).value or "").strip()
                if "total of:" in cell_val.lower():
                    next_c = c + 1
                    next_val = ""
                    while next_c <= min(worksheet.max_column, 40):
                        temp_val = str(worksheet.cell(row=r, column=next_c).value or "").strip()
                        if temp_val:
                            next_val = temp_val
                            break
                        next_c += 1

                    if "leather" in next_val.lower():
                        next_val_lower = next_val.lower()
                        leather_key = "BUFFALO" if "buffalo" in next_val_lower else "COW"

                        total_col_id = self._find_col_id(columns, c) or "col_po"
                        label_col_id = self._find_col_id(columns, next_c) or "col_item"

                        fact = LeatherSummaryFact(
                            leather_key=leather_key,
                            total_col_id=total_col_id,
                            label_col_id=label_col_id,
                            total_value=cell_val,
                            label_value=next_val
                        )
                        addon_facts.append(fact)

        return addon_facts

    def _find_col_id(self, columns: List[ColumnInfo], col_idx: int) -> Optional[str]:
        """Find col_id from columns list or children by col_index."""
        for col in columns:
            if col.col_index == col_idx:
                return col.id
            for child in col.children:
                if child.col_index == col_idx:
                    return child.id
        return None
