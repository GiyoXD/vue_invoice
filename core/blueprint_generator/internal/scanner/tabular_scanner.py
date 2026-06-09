import logging
from typing import Dict, List, Any, Optional, Tuple
from collections import Counter
from openpyxl.worksheet.worksheet import Worksheet

from core.blueprint_generator.schema import BlueprintSchema
from core.utils.loop_profiler import loop_profiler, tick
from core.blueprint_generator.utils.openpyxl_utils import get_actual_column_width
from .models import ColumnInfo, TableLayout, ZoneBoundaries
from .header_detector import get_cell_value
from core.blueprint_generator.utils.footer_scanner import scan_footer
from core.blueprint_generator.utils.content_extractor import (
    detect_static_description_label,
    extract_static_column_values
)

logger = logging.getLogger(__name__)


class _ColumnScanner:
    """Short-lived class that scans columns for a single worksheet, preventing stale state issues."""
    
    def __init__(self, worksheet: Worksheet, header_row: int, data_start_row: int, 
                 mapping_config: Optional[Dict[str, Any]] = None):
        self.worksheet = worksheet
        self.header_row = header_row
        self.data_start_row = data_start_row
        self.mapping_config = mapping_config
        self.logger = logger

    def _get_merged_header_ranges(self) -> Tuple[List[Any], List[Any]]:
        """Get merged ranges starting at or overlapping the header row."""
        merged_ranges = []
        all_merges_at_header = []
        for merged in self.worksheet.merged_cells.ranges:
            tick("scanner._scan_columns", sub="merge_scan_header")
            if merged.min_row == self.header_row:
                merged_ranges.append(merged)
            tick("scanner._scan_columns", sub="merge_scan_overlap")
            if merged.min_row <= self.header_row <= merged.max_row:
                all_merges_at_header.append(merged)
        return merged_ranges, all_merges_at_header

    def _get_merged_cell_value(self, cell: Any, all_merges_at_header: List[Any]) -> Optional[str]:
        """Get cell value, checking merged ranges if top-left is empty."""
        value = get_cell_value(cell)
        if not value:
            for merged in all_merges_at_header:
                if merged.min_col <= cell.column <= merged.max_col:
                    value = get_cell_value(
                        self.worksheet.cell(row=merged.min_row, column=merged.min_col)
                    )
                    break
        return value

    def _is_true_parent(self, cell: Any, colspan: int) -> bool:
        """Determine if a multi-row header cell is a true parent split into children."""
        is_full_width_merge = False
        for merged in self.worksheet.merged_cells.ranges:
            tick("scanner._scan_columns", sub="child_merge_scan")
            if merged.min_row == cell.row + 1 and merged.min_col == cell.column and merged.max_col == cell.column + colspan - 1:
                is_full_width_merge = True
                break
                
        if is_full_width_merge:
            return False

        child_cells_with_data = 0
        for col_offset in range(colspan):
            child_cell = cell.offset(row=1, column=col_offset)
            val = get_cell_value(child_cell)
            if val:
                mapped_id = self.determine_column_id(val, child_cell.column)
                if mapped_id and not mapped_id.startswith("col_unknown"):
                    return True
                child_cells_with_data += 1
                
        return child_cells_with_data > 1

    def _get_column_format(self, col: int, col_id: str) -> str:
        """Determine and return the formatting for the column."""
        # [Smart Feature] Determine format: Check data first, then Rules
        format_str = self.sample_column_format(col)
        if not format_str or format_str == "General":
            format_str = self.determine_format(col_id)
            
        # FORCE text format for identifiers to prevent scientific notation or leading zero loss
        if col_id in ["col_po", "col_item", "col_no", "col_container_no", "col_hs_code", "col_pallet_count", "col_dc"]:
            format_str = "@"
            
        return format_str

    def sample_column_format(self, col: int, max_rows: int = 15) -> Optional[str]:
        """Sample data rows to find the most common number format."""
        formats = []
        rows_to_scan = min(self.data_start_row + 500, self.worksheet.max_row + 1)
        
        data_found_at = None
        
        # Step 1: Find the first row with a numeric value in this column
        for r in range(self.data_start_row, rows_to_scan):
            tick("scanner._sample_column_format", sub="find_numeric_rows")
            cell = self.worksheet.cell(row=r, column=col)
            if cell.value is not None:
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    data_found_at = r
                    break

        if not data_found_at:
             # Fallback
             for r in range(self.data_start_row, min(self.data_start_row + 500, self.worksheet.max_row + 1)):
                 tick("scanner._sample_column_format", sub="fallback_format_rows")
                 cell = self.worksheet.cell(row=r, column=col)
                 if cell.number_format and cell.number_format != 'General':
                     formats.append(cell.number_format)
        else:
             # Step 2
             for r in range(data_found_at, min(data_found_at + max_rows, self.worksheet.max_row + 1)):
                 tick("scanner._sample_column_format", sub="sample_rows")
                 cell = self.worksheet.cell(row=r, column=col)
                 if cell.value is not None and cell.number_format and cell.number_format != 'General':
                     if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                         formats.append(cell.number_format)
        
        if not formats:
            return None
            
        most_common = Counter(formats).most_common(1)
        return most_common[0][0] if most_common else None

    def find_child_columns(self, row: int, start_col: int, span: int) -> List[ColumnInfo]:
        """Find child columns under a parent header."""
        children = []
        for col in range(start_col, start_col + span):
            cell = self.worksheet.cell(row=row, column=col)
            value = get_cell_value(cell)
            if value:
                col_id = self.determine_column_id(value, col)
                
                if not col_id or col_id.startswith("col_unknown"):
                    self.logger.debug(f"Child Column '{value}' at index {col} mapped to unknown.")

                format_str = self._get_column_format(col, col_id)
                    
                child_colspan = 1
                for merged in self.worksheet.merged_cells.ranges:
                    if merged.min_col == col and merged.min_row == row:
                        child_colspan = merged.max_col - merged.min_col + 1
                        break
                        
                width = get_actual_column_width(self.worksheet, col, child_colspan)
                
                children.append(ColumnInfo(
                    id=col_id,
                    header=value,
                    col_index=col,
                    width=width,
                    format=format_str
                ))
        return children

    def determine_column_id(self, header_text: str, col_index: int) -> str:
        """Determine column ID from header text using Config first, then Rules."""
        header_text_stripped = header_text.strip()
        
        if self.mapping_config:
            mappings = self.mapping_config.get('header_text_mappings', {}).get('mappings', {})
            if header_text_stripped in mappings:
                return mappings[header_text_stripped]
                
            for mapped_header, mapped_id in mappings.items():
                if mapped_header.strip().lower() == header_text_stripped.lower():
                    return mapped_id

        col_def = BlueprintSchema.get_column_by_keyword(header_text)
        if col_def:
            return col_def.id
        
        return f"col_unknown_{col_index}"
    
    def determine_format(self, col_id: str) -> str:
        """Determine number format for column using Rules."""
        return BlueprintSchema.get_format_for_id(col_id)


    def _get_cell_span(self, col: int, merged_ranges: List[Any]) -> Tuple[int, int]:
        """Determine colspan and rowspan for the cell at the header row."""
        for merged in merged_ranges:
            if merged.min_col == col and merged.min_row == self.header_row:
                colspan = merged.max_col - merged.min_col + 1
                rowspan = merged.max_row - merged.min_row + 1
                return colspan, rowspan
        return 1, 1

    def _create_column_info(self, cell: Any, value: str, col_id: str, 
                            merged_ranges: List[Any]) -> ColumnInfo:
        """Create and populate a ColumnInfo object for a scanned column."""
        col = cell.column
        colspan, rowspan = self._get_cell_span(col, merged_ranges)
        
        width = get_actual_column_width(self.worksheet, col, colspan)
        format_str = self._get_column_format(col, col_id)
        
        alignment = "center"
        if cell.alignment:
            alignment = cell.alignment.horizontal or "center"
        
        wrap_text = cell.alignment.wrap_text if cell.alignment else False
        
        column = ColumnInfo(
            id=col_id,
            header=value,
            col_index=col,
            width=width,
            format=format_str,
            alignment=alignment,
            rowspan=rowspan,
            colspan=colspan,
            wrap_text=wrap_text
        )
        
        if rowspan == 1 and colspan > 1:
            if self._is_true_parent(cell, colspan):
                column.children = self.find_child_columns(self.header_row + 1, col, colspan)
                
        return column

    @loop_profiler.watch("scanner._scan_columns")
    def scan(self) -> List[ColumnInfo]:
        """Orchestrate the scanning of columns from the header row."""
        columns = []
        
        merged_ranges, all_merges_at_header = self._get_merged_header_ranges()
        safe_max = min(self.worksheet.max_column + 1, BlueprintSchema.MAX_SCAN_COLUMN)
        
        col_iter = iter(range(1, safe_max))
        for col in col_iter:
            cell = self.worksheet.cell(row=self.header_row, column=col)
            value = self._get_merged_cell_value(cell, all_merges_at_header)
            if not value:
                continue
            
            col_id = self.determine_column_id(value, col)
            if col_id.startswith("col_unknown_") and len(value) > 35:
                self.logger.info(f"    [Leak Filter] Ignored long unmapped text as column header: '{value[:30]}...'")
                continue
                
            column = self._create_column_info(cell, value, col_id, merged_ranges)
            columns.append(column)
            
            if column.colspan > 1:
                for _ in range(column.colspan - 1):
                    next(col_iter, None)
        
        return columns


def scan_columns(worksheet: Worksheet, 
                 header_row: int, 
                 data_start_row: int,
                 mapping_config: Optional[Dict[str, Any]] = None) -> List[ColumnInfo]:
    """Scan and analyze columns for a worksheet using a short-lived stateful scanner."""
    scanner = _ColumnScanner(worksheet, header_row, data_start_row, mapping_config)
    return scanner.scan()


class TabularScanner:
    """Scans the entire table zone (Zone 2): columns, styling, and footer."""

    def scan_table(self, worksheet: Worksheet, boundaries: ZoneBoundaries, 
                   mapping_config: Optional[Dict[str, Any]] = None,
                   skip_desc_scan: bool = False,
                   skip_hs_scan: bool = False, sheet_name: str = "Unknown") -> TableLayout:
        # 1. Scan columns
        columns = scan_columns(worksheet, boundaries.header_row, boundaries.data_start_row, mapping_config)
        
        # 2. Extract font info
        header_font = self.extract_font_info(worksheet, boundaries.header_row, 1)
        data_font = self.extract_font_info(worksheet, boundaries.data_start_row, 1)
        
        # 3. Extract row heights
        row_heights = self.extract_row_heights(worksheet, boundaries.header_row, "dataset_default", boundaries.data_start_row)
        
        # 4. Scan footer
        footer_info = scan_footer(
            worksheet, boundaries.header_row, columns, logger,
            sheet_name=sheet_name, mapping_config=mapping_config, skip_hs_scan=skip_hs_scan,
            footer_row=boundaries.footer_row
        )
        
        # 5. Detect static content hints (description fallback, static lines) from the table
        footer_row = footer_info.row_num if footer_info else None
        static_hints = self.detect_static_content(worksheet, boundaries.header_row, columns, skip_desc_scan, footer_row=footer_row)
        
        has_multi_row = boundaries.data_start_row > boundaries.header_row + 1
        
        return TableLayout(
            header_row=boundaries.header_row,
            data_start_row=boundaries.data_start_row,
            columns=columns,
            header_font=header_font,
            data_font=data_font,
            row_heights=row_heights,
            has_multi_row_header=has_multi_row,
            footer_info=footer_info,
            static_content_hints=static_hints
        )

    @loop_profiler.watch("scanner._detect_static_content")
    def detect_static_content(self, worksheet: Worksheet, header_row: int, 
                              columns: List[ColumnInfo], skip_desc_scan: bool = False,
                              footer_row: Optional[int] = None) -> Dict[str, List[str]]:
        """Detect static content patterns in the data area."""
        hints = {}
        
        if not skip_desc_scan:
            desc_fallback = detect_static_description_label(worksheet, header_row, columns)
            if desc_fallback:
                hints["description_fallback"] = desc_fallback
        
        static_lines = extract_static_column_values(worksheet, header_row, columns, footer_row=footer_row)
        if static_lines:
            hints["static_lines"] = static_lines
        
        return hints

    def extract_font_info(self, worksheet: Worksheet, row: int, col: int) -> Dict[str, Any]:
        """Extract font information from a cell."""
        cell = worksheet.cell(row=row, column=col)
        font = cell.font

        if not font:
            raise ValueError(f"No font detected at Row {row}, Col {col}. Please ensure the template cell has explicit styling.")

        name = font.name
        size = font.size

        if name is None: name = "Calibri"
        if size is None: size = 11.0

        return {
            "name": name,
            "size": size,
            "bold": font.bold or False,
            "italic": font.italic or False
        }

    def extract_row_heights(self, worksheet: Worksheet, header_row: int, 
                            data_source: str = "dataset_default", 
                            data_start_row: int = None) -> Dict[str, float]:
        """Extract row heights using 3-step fallback strategy."""
        if data_start_row is None:
            data_start_row = header_row + 1
        
        def get_height(row_idx: int, standard_key: str) -> float:
            # 1. Explicit
            if row_idx in worksheet.row_dimensions:
                h = worksheet.row_dimensions[row_idx].height
                if h is not None:
                    return float(h)
            
            # 2. Sheet Default (Strict: If 15.0 or None, fail)
            default_h = worksheet.sheet_format.defaultRowHeight
            if default_h is not None and default_h != 15.0:
                 return float(default_h)
                 
            # 3. Failure
            logger.warning(f"Row {row_idx} has no explicit height. Defaulting to 20.0")
            return 20.0

        try:
            header_height = get_height(header_row, "header")
        except ValueError as e:
             logger.warning(f"Header Row {header_row} height detection failed: {e}")
             header_height = 20.0
            
        # Scan next 10 rows for data height (Median)
        heights = []
        for r in range(data_start_row, min(data_start_row + 10, worksheet.max_row + 1)):
            heights.append(get_height(r, "data"))
        
        if heights:
            heights.sort()
            mid = len(heights) // 2
            data_height = heights[mid]
        else:
            data_height = get_height(data_start_row, "data")
            
        return {
            "header": header_height,
            "data": data_height,
            "footer": header_height  # Usually same as header
        }
