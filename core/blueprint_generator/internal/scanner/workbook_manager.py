import logging
from pathlib import Path
from typing import Dict, List, Any, Optional
import openpyxl

from core.blueprint_generator.schema import BlueprintSchema
from core.utils.snitch import snitch
from .models import ColumnInfo, SheetAnalysis, TemplateAnalysisResult
from .header_detector import BoundaryDetector
from .sheet_manager import SheetManager

logger = logging.getLogger(__name__)


class WorkbookManager:
    """Orchestrates sheet-level analysis across an entire Excel workbook."""
    
    def __init__(self):
        self.logger = logging.getLogger(self.__class__.__name__)
        self.boundary_detector = BoundaryDetector()
        self.sheet_manager = SheetManager(self.boundary_detector)

    @snitch
    def scan_template(self, template_path: str, mapping_config: Optional[Dict[str, Any]] = None, 
                      workbook: Optional[openpyxl.Workbook] = None) -> TemplateAnalysisResult:
        """
        Analyze an Excel template file and extract structure.
        
        Args:
            template_path: Path to the Excel template file
            mapping_config: Optional configuration for header mappings
            workbook: Optional pre-loaded openpyxl Workbook object (for performance)
            
        Returns:
            TemplateAnalysisResult with complete analysis
        """
        path = Path(template_path)
        if not path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")
        
        # Extract customer code from filename (e.g., "CLW.xlsx" -> "CLW")
        customer_code = path.stem.upper()
        
        self.logger.info(f"Scanning template: {path.name} (customer: {customer_code})")
        
        if mapping_config:
            BlueprintSchema.load_dynamic_columns(mapping_config)
        
        if workbook is None:
             self.logger.debug("Loading workbook from disk...")
             workbook = openpyxl.load_workbook(template_path, data_only=False)
        else:
             self.logger.debug("Using pre-loaded workbook.")
        
        sheets = []
        warnings = []
        
        global_desc = None
        global_hs_code = None
        global_hs_colspan = 1
        global_hs_col_id = None
        
        supported_sheet_names = set()
        for sheet_name in workbook.sheetnames:
            if not self._is_sheet_supported(sheet_name, mapping_config):
                self.logger.info(f"  Skipping sheet '{sheet_name}': Not in allowed search list.")
                continue
            worksheet = workbook[sheet_name]
            analysis = self.sheet_manager.analyze_sheet(
                worksheet, 
                sheet_name, 
                mapping_config,
                skip_desc_scan=bool(global_desc),
                skip_hs_scan=False
            )
            if analysis:
                sheets.append(analysis)
                supported_sheet_names.add(sheet_name)
                # Collect proactive warnings
                if hasattr(analysis, "_temp_warning"):
                    warnings.append(getattr(analysis, "_temp_warning"))
                    
                # Harvest globals
                if analysis.static_content_hints and analysis.static_content_hints.get("description_fallback"):
                    global_desc = analysis.static_content_hints.get("description_fallback")
                
                if analysis.footer_info and analysis.footer_info.has_hs_code:
                    if not global_hs_code:
                        global_hs_code = analysis.footer_info.hs_code_text
                        global_hs_colspan = analysis.footer_info.hs_code_colspan
                        global_hs_col_id = analysis.footer_info.hs_code_col_id
 
        # Check mapping config options to ignore missing description fallback
        ignore_missing_desc = False
        if mapping_config:
            ignore_missing_desc = (
                mapping_config.get("ignore_missing_description", False) or
                mapping_config.get("fallback_strategies", {}).get("ignore_missing_description", False)
            )

        # Strict Validation: Both values are required for the blueprint to be complete.
        if not global_desc:
            if ignore_missing_desc:
                warnings.append(
                    f"Missing Description Fallback! No 'DES: ...' label found in col_static. "
                    f"Bypassed because 'ignore_missing_description' is enabled."
                )
            else:
                raise ValueError(
                    f"Missing Description Fallback! No 'DES: ...' label found in col_static. "
                    f"Ensure the template has product descriptions in the static column: {path.name}"
                )

        if not global_hs_code:
            if ignore_missing_desc:
                warnings.append(
                    f"Missing HS Code! No 'HS.CODE' row found in any sheet footer. "
                    f"Bypassed because 'ignore_missing_description' is enabled."
                )
            else:
                raise ValueError(
                    f"Missing HS Code! Ensure at least one sheet footer has an 'HS.CODE' row: {path.name}"
                )

        # Detect static sheets (sheets in workbook but not supported/analyzed)
        has_static = any(s not in supported_sheet_names for s in workbook.sheetnames)

        # Apply globals to all sheets
        for sheet in sheets:
            if global_desc:
                if not sheet.static_content_hints:
                    sheet.static_content_hints = {}
                if "description_fallback" not in sheet.static_content_hints:
                    sheet.static_content_hints["description_fallback"] = global_desc
            
            if global_hs_code:
                if sheet.footer_info and not sheet.footer_info.has_hs_code:
                    sheet.footer_info.has_hs_code = True
                    sheet.footer_info.hs_code_text = global_hs_code
                    sheet.footer_info.hs_code_colspan = global_hs_colspan
                    sheet.footer_info.hs_code_col_id = global_hs_col_id

        if not sheets:
             self.logger.warning(f"No valid sheets found in {template_path}. Ensure the file contains recognizable headers.")
             raise ValueError("No valid invoice structure detected. Please check your file content or Mapping Config.")
             
        return TemplateAnalysisResult(
            file_path=str(path.absolute()),
            customer_code=customer_code,
            sheets=sheets,
            warnings=warnings,
            has_static_sheets=has_static
        )

    def _is_sheet_supported(self, sheet_name: str, mapping_config: Optional[Dict[str, Any]] = None) -> bool:
        """Check if a sheet name is in the allowed search list."""
        normalized_name = sheet_name.lower().strip()
        
        # Fast mapping resolution using nested structure
        if mapping_config and isinstance(mapping_config, dict):
            sheet_mappings = mapping_config.get('sheet_name_mappings', {}).get('mappings', {})
            if isinstance(sheet_mappings, dict):
                # Fast case-insensitive exact matching
                lower_mappings = {k.lower().strip(): v for k, v in sheet_mappings.items()}
                if normalized_name in lower_mappings:
                    normalized_name = lower_mappings[normalized_name].lower().strip()

        is_supported = False
        
        # Create a set of variants for matching (with/without underscores/spaces)
        variants_to_check = {
            normalized_name,
            normalized_name.replace(' ', '_'),
            normalized_name.replace('_', ' ')
        }
        
        # Exact match check
        for variant in variants_to_check:
            if variant in BlueprintSchema.ALLOWED_SEARCH_SHEETS:
                is_supported = True
                break
                
        return is_supported


if __name__ == "__main__":
    import sys
    from core.logger_config import setup_logging
    from core.system_config import sys_config
    setup_logging(log_dir=sys_config.run_log_dir)
    
    if len(sys.argv) < 2:
        print("Usage: python workbook_manager.py <template.xlsx>")
        sys.exit(1)
    
    analyzer = WorkbookManager()
    result = analyzer.scan_template(sys.argv[1])
    
    print(f"\nTemplate: {result.customer_code}")
    print(f"Sheets: {len(result.sheets)}")
    for sheet in result.sheets:
        print(f"\n  {sheet.name}:")
        print(f"    Header row: {sheet.header_row}")
        print(f"    Data source: {sheet.data_source}")
        print(f"    Columns: {len(sheet.columns)}")
        for col in sheet.columns:
            print(f"      - {col.id}: '{col.header}' (width={col.width:.1f})")
